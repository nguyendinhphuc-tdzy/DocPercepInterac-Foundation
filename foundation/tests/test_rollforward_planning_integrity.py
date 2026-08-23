"""
Roll-Forward Planning Integrity Tests (Phase C2 / D3.2)
=======================================================
Location: foundation/tests/test_rollforward_planning_integrity.py

Each test closes one Phase D3.1 finding and is written to FAIL if the defect
returns. None of these tests may be weakened to obtain a green run.

  P0-1  no positional table identity in planning
  P0-2  no Ground Truth dependency in planning
  P0-3  semantic placement violations block before D1
  P0-4  VERIFIED lineage requires a source location
  P0-5  a benchmarking target without a benchmarking source blocks
  P0-6  real postcondition hashes; no "dummy"

Synthetic fixtures are labelled SYNTHETIC in their docstring and never stand
in for real-fixture evidence.
"""
from pathlib import Path
import sys

from docx import Document
import pytest

REPO_ROOT = Path(__file__).resolve().parents[2]
for _p in (str(REPO_ROOT), str(REPO_ROOT / "foundation")):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from foundation.applications.rollforward.data_reconciliation import (
    CellReconciliationRecord,
    LineageAddressabilityGate,
    ReconciliationStatus,
    SourceAddressability,
    SourceBindingLineageStatus,
    SourceCellReference,
    TargetCellReference,
)
from foundation.applications.rollforward.mutation_precondition import (
    FORBIDDEN_POSTCONDITION_VALUES,
    MutationPreconditionValidator,
    PostconditionHasher,
    PreconditionViolationCode,
    TableAnatomyProfiler,
)
from foundation.applications.rollforward.semantic_binding import (
    BindingVerdict,
    FieldSemanticRole,
    RowSemantics,
    SemanticBindingValidator,
    TargetSchemaDeriver,
    classify_value_shape,
)
from foundation.applications.rollforward.source_capability import (
    DatasetRole,
    SourceCapabilityProfiler,
)
from foundation.applications.rollforward.structural_writeback import (
    CellMutationSpec,
    RowMutationSpec,
    TableMutationSpec,
)
from foundation.applications.rollforward.table_identity import (
    CorrespondenceConfidence,
    SemanticLabel,
    TableCorrespondenceResolver,
    TableIdentityProfiler,
)
from foundation.tests.evaluation.rollforward_clean_planner_c2 import (
    PATH_APP1,
    PATH_FARPT,
    PATH_GROUND_TRUTH_FORBIDDEN,
    PATH_HIST,
    PATH_TMPL,
    CleanRollForwardPlanner,
    GroundTruthContaminationError,
    GroundTruthQuarantine,
    Readiness,
    RecommendedStrategy,
    run as run_clean_plan,
)

# Golden-four template ordinals, used ONLY to address tables inside the
# template document itself -- never to correlate across documents.
T_ARMS_LENGTH, T_SCREENING, T_VN_COMPARABLES, T_COMPARABLES = 10, 13, 14, 15


@pytest.fixture(scope="module")
def clean_plan():
    result, benchmark = run_clean_plan(write=False)
    return result, benchmark


@pytest.fixture(scope="module")
def workbooks():
    return [SourceCapabilityProfiler.profile_workbook(p) for p in (PATH_FARPT, PATH_APP1)]


@pytest.fixture(scope="module")
def template_signatures():
    return {s.ordinal: s for s in TableIdentityProfiler.profile_document(PATH_TMPL, "TEMPLATE")}


# ============================================================================
# P0-1 — NO POSITIONAL TABLE IDENTITY
# ============================================================================

def test_correspondence_resolver_cannot_read_ordinal():
    """Two signatures differing only in ordinal must score identically."""
    assert TableCorrespondenceResolver.assert_no_positional_input() is True


def test_shifting_every_ordinal_does_not_change_correspondence():
    """Renumbering tables must not change a single correspondence verdict."""
    import dataclasses

    tmpl = TableIdentityProfiler.profile_document(PATH_TMPL, "TEMPLATE")
    hist = TableIdentityProfiler.profile_document(PATH_HIST, "HIST_FY2023")

    baseline = TableCorrespondenceResolver.correspond(tmpl, hist)
    shifted_tmpl = [dataclasses.replace(s, ordinal=s.ordinal + 1000) for s in tmpl]
    shifted_hist = [dataclasses.replace(s, ordinal=999 - s.ordinal) for s in reversed(hist)]
    shifted = TableCorrespondenceResolver.correspond(shifted_tmpl, shifted_hist)

    assert [c.confidence for c in baseline] == [c.confidence for c in shifted]
    assert [round(c.score, 6) for c in baseline] == [round(c.score, 6) for c in shifted]
    assert [c.right.identity_key if c.right else None for c in baseline] == \
           [c.right.identity_key if c.right else None for c in shifted]


def test_identity_key_excludes_ordinal():
    """The position-free identity digest must not encode the ordinal."""
    import dataclasses

    sigs = TableIdentityProfiler.profile_document(PATH_TMPL, "TEMPLATE")
    for s in sigs:
        assert dataclasses.replace(s, ordinal=s.ordinal + 500).identity_key == s.identity_key


def test_positional_pairing_reproduces_the_contaminated_deltas_and_is_rejected():
    """The old phantom deltas are reproducible by index -- and must not be used.

    This is the regression lock for P0-1: it demonstrates the exact arithmetic
    the contaminated planner used, then asserts the position-free resolver
    reaches a different (UNCORRELATED) verdict for those same tables.
    """
    hist = Document(str(PATH_HIST))
    gt = Document(str(PATH_GROUND_TRUTH_FORBIDDEN))
    contaminated = {10: (2, 11), 13: (4, 6), 14: (6, 10), 15: (10, 16)}
    for idx, (h_rows, g_rows) in contaminated.items():
        assert len(hist.tables[idx].rows) == h_rows
        assert len(gt.tables[idx].rows) == g_rows

    tmpl_sigs = TableIdentityProfiler.profile_document(PATH_TMPL, "TEMPLATE")
    hist_sigs = TableIdentityProfiler.profile_document(PATH_HIST, "HIST_FY2023")
    by_ordinal = {c.left.ordinal: c for c in
                  TableCorrespondenceResolver.correspond(tmpl_sigs, hist_sigs)}

    for idx in contaminated:
        corr = by_ordinal[idx]
        assert corr.confidence == CorrespondenceConfidence.UNCORRELATED, (
            f"Template table {idx} must not be correlated to any FY2023 table; the "
            f"contaminated planner paired it with FY2023 table {idx} by index alone."
        )


def test_clean_benchmark_declares_no_positional_identity(clean_plan):
    _result, benchmark = clean_plan
    assert benchmark["identity_policy"]["positional_table_identity_used"] is False
    assert benchmark["identity_policy"]["ordinal_role"] == "display and intra-document locator only"


# ============================================================================
# P0-2 — NO GROUND TRUTH IN PLANNING
# ============================================================================

def test_ground_truth_quarantine_raises_when_planning_opens_it():
    """The evaluation-only rule is enforced, not merely documented."""
    GroundTruthQuarantine.arm([PATH_GROUND_TRUTH_FORBIDDEN])
    try:
        with pytest.raises(GroundTruthContaminationError, match="Ground Truth"):
            GroundTruthQuarantine.open_document(PATH_GROUND_TRUTH_FORBIDDEN)
        # Permitted planning inputs stay open.
        assert GroundTruthQuarantine.open_document(PATH_TMPL) is not None
    finally:
        GroundTruthQuarantine.disarm()


def test_clean_planner_runs_with_quarantine_armed(clean_plan):
    """A full clean plan completes without ever touching the Ground Truth."""
    result, benchmark = clean_plan
    assert result.ground_truth_opened is False
    assert benchmark["ground_truth"]["used_in_planning"] is False
    assert result.cases, "the clean planner must produce cases"


def test_no_case_derives_its_target_from_ground_truth(clean_plan):
    result, _ = clean_plan
    for case in result.cases:
        assert case.target_state["derived_from_ground_truth"] is False, case.case_id


def test_clean_target_rows_never_equal_the_positional_ground_truth_artifact(clean_plan):
    """The clean plan must not reproduce GT.tables[i].rows for the golden four."""
    result, _ = clean_plan
    gt = Document(str(PATH_GROUND_TRUTH_FORBIDDEN))
    by_ordinal = {c.template_ordinal_display_only: c for c in result.cases}
    for idx in (T_ARMS_LENGTH, T_SCREENING, T_VN_COMPARABLES, T_COMPARABLES):
        contaminated_target = len(gt.tables[idx].rows)
        clean_target = by_ordinal[idx].target_state["rows"]
        if clean_target is not None:
            assert clean_target != contaminated_target or clean_target == len(
                Document(str(PATH_TMPL)).tables[idx].rows), (
                f"table {idx}: clean target {clean_target} coincides with the "
                f"positional Ground-Truth figure {contaminated_target}"
            )


def test_contaminated_planner_modules_are_marked_not_deleted():
    """History is preserved and labelled, per the phase's hard rule."""
    from foundation.tests.evaluation import rollforward_profiler, rollforward_source_binding

    for mod in (rollforward_source_binding, rollforward_profiler):
        assert mod.PLANNING_INTEGRITY_STATUS == "INVALIDATED_FOR_PLANNING_CONTAMINATION"
    assert rollforward_source_binding.SUPERSEDED_BY.endswith("rollforward_clean_planner_c2.py")
    # The contaminated planner still exists and still runs.
    assert hasattr(rollforward_source_binding, "run_rollforward_source_binding_pipeline")


# ============================================================================
# P0-5 — SOURCE SEMANTIC COMPATIBILITY
# ============================================================================

def test_source_capability_is_content_derived_not_name_derived(workbooks):
    """A sheet earns a role from its cells, never from its title."""
    farpt, app1 = workbooks
    assert farpt.has_role(DatasetRole.FINANCIAL_ANALYSIS)
    for sheet in farpt.sheets_with_role(DatasetRole.FINANCIAL_ANALYSIS):
        assert sheet.role_evidence.get(DatasetRole.FINANCIAL_ANALYSIS.value), (
            "a role must carry the content tokens that earned it")
    # "Interest expenses" is a sheet NAME in Appendix I; the role must be earned.
    named = [s for s in app1.sheets if s.sheet_name == "Interest expenses"]
    assert named, "fixture changed: expected an 'Interest expenses' sheet"
    if DatasetRole.INTEREST_EXPENSE in named[0].roles:
        assert named[0].role_evidence.get(DatasetRole.INTEREST_EXPENSE.value)


def test_neither_workbook_contains_any_benchmarking_dataset(workbooks):
    """P0-5: the bound sources carry no benchmarking data of any kind."""
    benchmarking = (DatasetRole.BENCHMARKING_DATA, DatasetRole.COMPARABLE_COMPANIES,
                    DatasetRole.IQR_RESULTS, DatasetRole.SCREENING_RESULTS)
    for wb in workbooks:
        for role in benchmarking:
            assert not wb.has_role(role), (
                f"{wb.document_name} unexpectedly reports {role.value}; "
                f"if the fixture genuinely gained benchmarking data this test must be "
                f"re-derived, not deleted."
            )


def test_benchmarking_target_without_benchmarking_source_is_blocked(template_signatures, workbooks):
    """A comparables table cannot bind to financial-analysis data."""
    for ordinal in (T_ARMS_LENGTH, T_SCREENING, T_VN_COMPARABLES, T_COMPARABLES):
        schema = TargetSchemaDeriver.derive(f"rft-{ordinal}", template_signatures[ordinal])
        verdict = SemanticBindingValidator.validate(schema, workbooks)
        assert verdict.verdict == BindingVerdict.MISSING_CURRENT_SOURCE, (
            f"template table {ordinal} ({schema.expected_domain.value}) must not be VERIFIED "
            f"against workbooks with no matching dataset; got {verdict.verdict.value}"
        )
        assert verdict.is_executable is False
        assert verdict.violations


def test_binding_verdict_names_the_required_roles_it_could_not_find(template_signatures, workbooks):
    schema = TargetSchemaDeriver.derive("rft-comparables", template_signatures[T_COMPARABLES])
    verdict = SemanticBindingValidator.validate(schema, workbooks)
    assert DatasetRole.COMPARABLE_COMPANIES in verdict.required_roles
    assert verdict.roles_found == []
    assert verdict.violations, "a blocked binding must say why"
    # The violation must name every role it looked for and failed to find.
    joined = " ".join(verdict.violations)
    for role in verdict.required_roles:
        assert role.value in joined, f"violation does not name the missing role {role.value}"
    # And it must evidence what the workbooks DO contain, so the gap is auditable.
    assert any("roles present" in e for e in verdict.evidence)


# ============================================================================
# P0-3 — SEMANTIC PLACEMENT
# ============================================================================

def test_value_shape_classification_keys_on_formatting_not_magnitude():
    assert classify_value_shape("0201234500") == "code"        # tax code, leading zero
    assert classify_value_shape("14100") == "integer"          # SIC code
    assert classify_value_shape("194,469,728,040") == "monetary"
    assert classify_value_shape("4.20%") == "percentage"
    assert classify_value_shape("TC-100") == "code"
    assert classify_value_shape("Vietnam") == "text"


def test_percentage_into_description_column_is_rejected(template_signatures):
    """The exact P0-3 defect: '4.20%' written into 'Business description'."""
    schema = TargetSchemaDeriver.derive("rft-comparables", template_signatures[T_COMPARABLES])
    desc = [c for c in schema.columns if c.role == FieldSemanticRole.DESCRIPTION]
    assert desc, "expected a DESCRIPTION column in the comparables table"
    ok, why = SemanticBindingValidator.validate_value_against_column(schema, desc[0].index, "4.20%")
    assert ok is False
    assert "percentage" in why and "DESCRIPTION" in why


def test_footer_band_is_detected_in_the_real_template():
    """Template table 14 ends in a full-width footer row."""
    doc = Document(str(PATH_TMPL))
    anatomy = TableAnatomyProfiler.profile(doc.tables[T_VN_COMPARABLES], T_VN_COMPARABLES)
    assert anatomy.footer_row_idxs == (7,)
    assert anatomy.data_region == (1, 6)
    assert anatomy.append_position == 6


def test_precondition_gate_blocks_every_phase_d3_semantic_defect(template_signatures):
    """Replays the published D3 plan and requires it to be refused."""
    schemas = {}
    specs = []
    plan_rows = {
        T_ARMS_LENGTH: (6, [["Net Sales", "194,469,728,040", "194,469,728,040"]]),
        T_VN_COMPARABLES: (8, [["1", "Peer 1 JSC", "Hai Phong", "0201234500", "14100", "Gloves"]]),
        T_COMPARABLES: (7, [["1", "Peer Company 1", "Vietnam", "TC-100", "4.20%"]]),
    }
    for ordinal, (start, rows) in plan_rows.items():
        rid = f"rft-{ordinal}"
        schemas[rid] = TargetSchemaDeriver.derive(rid, template_signatures[ordinal])
        specs.append(TableMutationSpec(
            target_region_id=rid, table_index=ordinal, table_hash="h",
            initial_row_count=start, target_row_count=start + len(rows), insert_count=len(rows),
            expected_precondition_hash="pre", expected_postcondition_hash="dummy",
            row_mutations=[
                RowMutationSpec(row_idx=start + i, cells=[
                    CellMutationSpec(col_idx=c, source_doc_name="HMV-FA&RPT FY2024.xlsx", value=v)
                    for c, v in enumerate(vals)])
                for i, vals in enumerate(rows)],
        ))

    report = MutationPreconditionValidator.validate(PATH_TMPL, specs, schemas)
    assert report.is_executable is False
    codes = {v.code for v in report.violations}

    assert PreconditionViolationCode.ROW_SEMANTICS_MISMATCH in codes, \
        "P&L rows appended to the arm's-length range table must be refused"
    assert PreconditionViolationCode.INSERTION_BELOW_FOOTER in codes, \
        "a row landing below the table footer must be refused"
    assert PreconditionViolationCode.COLUMN_DOMAIN_VIOLATION in codes, \
        "a percentage in a Business description column must be refused"
    assert PreconditionViolationCode.DUMMY_POSTCONDITION_HASH in codes, \
        "a placeholder postcondition hash must be refused"
    assert PreconditionViolationCode.PLACEHOLDER_ROWS_NOT_CONSUMED in codes


def test_precondition_gate_reports_the_exact_offending_location(template_signatures):
    rid = f"rft-{T_COMPARABLES}"
    schema = TargetSchemaDeriver.derive(rid, template_signatures[T_COMPARABLES])
    spec = TableMutationSpec(
        target_region_id=rid, table_index=T_COMPARABLES, table_hash="h",
        initial_row_count=7, target_row_count=8, insert_count=1,
        expected_precondition_hash="pre",
        expected_postcondition_hash=PostconditionHasher.project_from_path(
            PATH_TMPL, T_COMPARABLES,
            [RowMutationSpec(row_idx=7, cells=[CellMutationSpec(col_idx=4, source_doc_name="s", value="4.20%")])]),
        row_mutations=[RowMutationSpec(row_idx=7, cells=[
            CellMutationSpec(col_idx=4, source_doc_name="s.xlsx", value="4.20%")])],
    )
    report = MutationPreconditionValidator.validate(PATH_TMPL, [spec], {rid: schema})
    offending = [v for v in report.violations
                 if v.code == PreconditionViolationCode.COLUMN_DOMAIN_VIOLATION]
    assert offending
    assert offending[0].location == f"table {T_COMPARABLES} row 7 col 4"


def test_table_identity_mismatch_blocks_a_plan_aimed_at_the_wrong_table(template_signatures):
    """A plan whose ordinal points at a different table than its region declares."""
    rid = "rft-comparables"
    schema = TargetSchemaDeriver.derive(rid, template_signatures[T_COMPARABLES])
    spec = TableMutationSpec(
        target_region_id=rid, table_index=T_ARMS_LENGTH, table_hash="h",
        initial_row_count=6, target_row_count=7, insert_count=1,
        expected_precondition_hash="pre", expected_postcondition_hash="a" * 32,
        row_mutations=[],
    )
    report = MutationPreconditionValidator.validate(PATH_TMPL, [spec], {rid: schema})
    assert PreconditionViolationCode.TABLE_IDENTITY_MISMATCH in {v.code for v in report.violations}


# ============================================================================
# P0-4 — LINEAGE ADDRESSABILITY HARD GATE
# ============================================================================

def _record(source):
    return CellReconciliationRecord(
        reconciliation_id="rec-1", manifest_id="rfm-1", manifest_version=1, mutation_id="mut-1",
        target=TargetCellReference(region_id="r", table_index=0, table_hash="h", row_idx=1, col_idx=0),
        source=source, semantic_match=True, display_match=True,
        status=ReconciliationStatus.MATCH,
    )


def test_cell_without_any_source_cannot_be_match():
    rec = LineageAddressabilityGate.apply(_record(None))
    assert rec.source_addressability == SourceAddressability.UNADDRESSED
    assert rec.binding_status == SourceBindingLineageStatus.UNVERIFIED
    assert rec.status == ReconciliationStatus.BLOCKED
    assert rec.value_semantic_status == ReconciliationStatus.MATCH  # value verdict preserved


def test_cell_with_document_but_no_cell_address_is_partial_and_blocked():
    """The exact P0-4 shape: workbook named, no cell address, reported as MATCH."""
    rec = LineageAddressabilityGate.apply(_record(SourceCellReference(
        document_id="doc-farpt", document_name="HMV-FA&RPT FY2024.xlsx", sheet_name="FS")))
    assert rec.source_addressability == SourceAddressability.PARTIAL
    assert rec.status == ReconciliationStatus.BLOCKED
    assert "SOURCE_NOT_ADDRESSABLE" in rec.discrepancy_reason


def test_fully_addressed_cell_keeps_match_and_becomes_verified():
    rec = LineageAddressabilityGate.apply(_record(SourceCellReference(
        document_id="doc-farpt", document_name="HMV-FA&RPT FY2024.xlsx",
        sheet_name="Financial Analysis", cell_address="D34")))
    assert rec.source_addressability == SourceAddressability.ADDRESSED
    assert rec.binding_status == SourceBindingLineageStatus.VERIFIED
    assert rec.status == ReconciliationStatus.MATCH


def test_authoritative_element_id_satisfies_addressability():
    rec = LineageAddressabilityGate.apply(_record(SourceCellReference(
        document_id="doc-hist", document_name="FY2023.docx",
        element_id="948bbe17-7991-53a7-ba07-01458403670c")))
    assert rec.source_addressability == SourceAddressability.ADDRESSED
    assert rec.status == ReconciliationStatus.MATCH


def test_the_72_of_72_result_does_not_survive_the_gate():
    """Every cell of the published D3 lineage was source-unaddressable."""
    import json

    report_path = REPO_ROOT / "docs/evaluation/LocalFile_RollForward_D3_Execution_Report.json"
    if not report_path.exists():
        pytest.skip("Phase D3 execution report not present")
    payload = json.loads(report_path.read_text(encoding="utf-8"))
    cells = [n for n in payload["lineage"]["nodes"] if n["type"] == "TARGET_CELL"]
    assert cells, "expected TARGET_CELL lineage nodes"
    addressed = [n for n in cells if n["metadata"].get("source_cell")]
    assert addressed == [], (
        f"{len(addressed)}/{len(cells)} cells carry a source cell address; the historical "
        f"record is 0/72 and this test locks that finding in place."
    )


# ============================================================================
# P0-6 — REAL POSTCONDITION HASH
# ============================================================================

def test_placeholder_postcondition_values_are_rejected():
    assert "dummy" in FORBIDDEN_POSTCONDITION_VALUES
    assert "post-dummy" in FORBIDDEN_POSTCONDITION_VALUES
    assert "" in FORBIDDEN_POSTCONDITION_VALUES


def test_projected_postcondition_matches_the_executed_result(tmp_path):
    """SYNTHETIC fixture. A projected hash must equal the hash after mutation."""
    doc_path = tmp_path / "proj.docx"
    doc = Document()
    t = doc.add_table(rows=2, cols=3)
    for c in range(3):
        t.rows[0].cells[c].text = f"H{c}"
    for c in range(3):
        t.rows[1].cells[c].text = f"proto{c}"
    doc.save(str(doc_path))

    mutations = [RowMutationSpec(row_idx=2, cells=[
        CellMutationSpec(col_idx=0, source_doc_name="s", value="A"),
        CellMutationSpec(col_idx=1, source_doc_name="s", value="B"),
        CellMutationSpec(col_idx=2, source_doc_name="s", value="C"),
    ])]
    projected = PostconditionHasher.project_from_path(doc_path, 0, mutations, prototype_row_idx=1)

    # Apply the same mutation for real, using the production cloner.
    from foundation.applications.rollforward.structural_writeback import OxmlRowCloner
    live = Document(str(doc_path))
    OxmlRowCloner.clone_and_populate_row(
        table=live.tables[0], prototype_row_idx=1, insert_after_row_idx=1,
        cell_specs=mutations[0].cells)
    out = tmp_path / "applied.docx"
    live.save(str(out))

    actual = PostconditionHasher.from_table(Document(str(out)).tables[0])
    assert projected == actual, "projected postcondition must equal the executed postcondition"


def test_postcondition_hash_is_deterministic_and_sensitive():
    doc = Document(str(PATH_TMPL))
    table = doc.tables[T_COMPARABLES]
    a = PostconditionHasher.from_table(table)
    assert a == PostconditionHasher.from_table(table)
    projected = PostconditionHasher.project(table, [RowMutationSpec(row_idx=7, cells=[
        CellMutationSpec(col_idx=1, source_doc_name="s", value="ACME JSC")])])
    assert projected != a
    assert len(a) == 32


# ============================================================================
# CLEAN BENCHMARK — HONEST CLASSIFICATION
# ============================================================================

def test_every_benchmark_case_is_classified_and_justified(clean_plan):
    result, benchmark = clean_plan
    valid = {r.value for r in Readiness}
    for case in result.cases:
        assert case.readiness in valid
        if case.readiness == Readiness.BLOCKED.value:
            assert case.blocking_reasons, f"{case.case_id} blocked without a reason"
        if case.readiness == Readiness.MANUAL_REVIEW.value:
            assert case.manual_review_reasons, f"{case.case_id} needs review without a reason"
        if case.readiness == Readiness.READY.value:
            assert not case.blocking_reasons and not case.manual_review_reasons
            assert case.strategy_is_implemented, (
                f"{case.case_id} is READY but its strategy "
                f"{case.recommended_strategy} is not implemented")
    assert sum(benchmark["readiness_summary"].values()) == len(result.cases)


def test_no_case_is_ready_on_an_unimplemented_strategy(clean_plan):
    result, _ = clean_plan
    for case in result.cases:
        if not case.strategy_is_implemented:
            assert case.readiness != Readiness.READY.value


def test_the_four_contaminated_golden_tables_are_not_ready(clean_plan):
    """The regions the contaminated plan called READY are now honestly classified."""
    result, _ = clean_plan
    by_ordinal = {c.template_ordinal_display_only: c for c in result.cases}
    for ordinal in (T_ARMS_LENGTH, T_SCREENING, T_VN_COMPARABLES, T_COMPARABLES):
        case = by_ordinal[ordinal]
        assert case.readiness == Readiness.BLOCKED.value, (
            f"template table {ordinal} is {case.readiness}; it has no current source and "
            f"needs {case.recommended_strategy}, so it must be BLOCKED"
        )
        assert any("MISSING_CURRENT_SOURCE" in r or "STRATEGY_NOT_IMPLEMENTED" in r
                   for r in case.blocking_reasons)


def test_benchmark_records_real_postcondition_hashes_for_every_case(clean_plan):
    result, _ = clean_plan
    for case in result.cases:
        h = case.real_postcondition_hash_current
        assert h and h.lower() not in FORBIDDEN_POSTCONDITION_VALUES
        assert len(h) == 32


def test_arms_length_table_is_recognised_as_fixed_shape(clean_plan):
    """Table 10 is a statistic summary; its row count must not change."""
    result, _ = clean_plan
    case = next(c for c in result.cases if c.template_ordinal_display_only == T_ARMS_LENGTH)
    assert case.template_state["row_semantics"] == RowSemantics.STATISTIC.value
    assert case.recommended_strategy == RecommendedStrategy.ACTIVATE_PLACEHOLDER.value
    assert case.target_state["rows"] == case.template_state["rows"]


def test_screening_matrix_needs_placeholder_activation_not_insertion(clean_plan):
    result, _ = clean_plan
    case = next(c for c in result.cases if c.template_ordinal_display_only == T_SCREENING)
    assert case.semantic_domain == SemanticLabel.SCREENING_STRATEGY.value
    assert case.recommended_strategy == RecommendedStrategy.ACTIVATE_PLACEHOLDER.value
    assert case.strategy_is_implemented is False


def test_delete_rows_is_not_implemented_anywhere():
    """This phase must not add DELETE_ROWS; only recommend it."""
    from foundation.applications.rollforward import structural_writeback

    src = Path(structural_writeback.__file__).read_text(encoding="utf-8")
    assert "DELETE_ROWS" not in src
    assert RecommendedStrategy.DELETE_ROWS.value == "DELETE_ROWS"  # recommendation vocabulary only
