"""
Source Completeness & Readiness Audit Tests (Phase E)
=====================================================
Location: foundation/tests/test_source_completeness_audit.py

Locks the Phase E audit's honesty properties:

    every workbook sheet is profiled
    dataset roles derive from content, never from sheet names
    benchmarking absence is proven, not assumed
    narrative-source gaps are detected
    every figure is assessed individually
    historical text is never auto-promoted to current-year verified
    blocked regions stay blocked
    Ground Truth cannot affect readiness
    the requirement matrix is deterministic and reproducible
    no fabricated source artifacts, ranges or element ids

These tests must not be weakened to obtain a green run.
"""
from pathlib import Path
import json
import re
import sys

import openpyxl
import pytest

REPO_ROOT = Path(__file__).resolve().parents[2]
for _p in (str(REPO_ROOT), str(REPO_ROOT / "foundation")):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from foundation.tests.evaluation.rollforward_clean_planner_c2 import (
    PATH_APP1,
    PATH_FARPT,
    PATH_GROUND_TRUTH_FORBIDDEN,
    PATH_TMPL,
    GroundTruthContaminationError,
    GroundTruthQuarantine,
)
from foundation.tests.evaluation.source_completeness_audit import (
    DOMAIN_RULES,
    ArtifactType,
    ExtendedSourceProfiler,
    FigureInventory,
    FigureKind,
    FigureReadiness,
    GapEvaluation,
    HistoricalReuse,
    Priority,
    Readiness,
    RegionDomain,
    SourceCompletenessAudit,
    SourceCoverage,
    SourceDatasetRole,
    build_matrix_artifact,
    build_readiness_artifact,
    run as run_audit,
)

BENCHMARKING_ROLES = (
    SourceDatasetRole.BENCHMARKING_DATA,
    SourceDatasetRole.COMPARABLE_COMPANIES,
    SourceDatasetRole.IQR_RESULTS,
    SourceDatasetRole.SCREENING_RESULTS,
)


@pytest.fixture(scope="module")
def audit():
    return SourceCompletenessAudit.freeze()


@pytest.fixture(scope="module")
def artifacts(audit):
    return build_readiness_artifact(audit), build_matrix_artifact(audit)


# ============================================================================
# SOURCE PROFILING
# ============================================================================

def test_every_workbook_sheet_is_profiled(audit):
    """No sheet may be skipped: absence must be proven over the whole workbook."""
    expected = 0
    for path in (PATH_FARPT, PATH_APP1):
        wb = openpyxl.load_workbook(str(path), read_only=True)
        try:
            expected += len(wb.sheetnames)
        finally:
            wb.close()
    assert len(audit.sheet_profiles) == expected
    assert expected == sum(w.sheet_count for w in audit.workbooks)

    for path in (PATH_FARPT, PATH_APP1):
        wb = openpyxl.load_workbook(str(path), read_only=True)
        try:
            names = set(wb.sheetnames)
        finally:
            wb.close()
        profiled = {s["sheet_name"] for s in audit.sheet_profiles if s["workbook"] == path.name}
        assert profiled == names, f"unprofiled sheets in {path.name}: {names - profiled}"


def test_every_sheet_profile_carries_content_evidence(audit):
    """A profile must record what it actually saw, not just a verdict."""
    for s in audit.sheet_profiles:
        assert s["row_count"] >= 0 and s["column_count"] >= 0
        assert s["record_pattern"] in (
            "EMPTY", "UNSTRUCTURED", "TABULAR_NUMERIC", "TABULAR_TEXT", "TABULAR_WITH_NARRATIVE")
        assert s["dataset_roles"], f"{s['sheet_name']} has no role verdict"
        assert isinstance(s["formula_presence"], bool)
        for role in s["dataset_roles"]:
            if role != SourceDatasetRole.UNKNOWN.value:
                assert role in s["role_evidence"] or role in s["c2_roles"], (
                    f"{s['sheet_name']} claims {role} without recorded evidence")


def test_dataset_roles_derive_from_content_not_sheet_name(audit):
    """A sheet named for a role must still earn it from its cells."""
    named = [s for s in audit.sheet_profiles if s["sheet_name"] == "Interest expenses"]
    assert named, "fixture changed: expected an 'Interest expenses' sheet"
    sheet = named[0]
    if SourceDatasetRole.INTEREST_EXPENSE.value in sheet["dataset_roles"]:
        assert sheet["role_evidence"].get(SourceDatasetRole.INTEREST_EXPENSE.value) or \
            "INTEREST_EXPENSE" in sheet["c2_roles"]

    # The inverse: a sheet whose name says nothing may still earn roles.
    fs = [s for s in audit.sheet_profiles if s["sheet_name"] == "FS"]
    assert fs and any(SourceDatasetRole.FINANCIAL_STATEMENTS.value in s["dataset_roles"] for s in fs)


def test_role_detection_ignores_the_sheet_title_entirely():
    """Renaming a sheet must not change its detected roles."""
    import tempfile

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Company name", "Ticker", "Tax code", "Business description"])
    ws.append(["ACME JSC", "ACM", "0101234567", "Manufacturer of gloves"])
    d = Path(tempfile.mkdtemp())
    a = d / "a.xlsx"
    wb.save(str(a))

    wb2 = openpyxl.load_workbook(str(a))
    wb2["Sheet1"].title = "TOTALLY UNRELATED NAME"
    b = d / "b.xlsx"
    wb2.save(str(b))
    wb.close()
    wb2.close()

    _wbs_a, sheets_a = ExtendedSourceProfiler.profile([a])
    _wbs_b, sheets_b = ExtendedSourceProfiler.profile([b])
    assert sheets_a[0]["dataset_roles"] == sheets_b[0]["dataset_roles"]
    assert SourceDatasetRole.COMPARABLE_COMPANIES.value in sheets_a[0]["dataset_roles"]


# ============================================================================
# BENCHMARKING ABSENCE (§9)
# ============================================================================

def test_benchmarking_absence_is_proven_over_every_sheet(audit):
    bm = audit.benchmarking_audit
    assert bm["sheets_scanned"] == len(audit.sheet_profiles)
    assert bm["any_benchmarking_source_present"] is False
    for label, finding in bm["findings"].items():
        assert finding["present"] is False, f"{label} unexpectedly present"
        assert finding["sheets"] == []
    assert bm["independence_codes_present"] is False
    assert bm["accepted_rejected_comparables_present"] is False


def test_no_sheet_anywhere_carries_a_benchmarking_role(audit):
    present = ExtendedSourceProfiler.roles_present(audit.sheet_profiles)
    for role in BENCHMARKING_ROLES:
        assert role.value not in present, (
            f"{role.value} appeared in the current source set; if the fixture genuinely gained "
            f"benchmarking data this test must be re-derived from evidence, not deleted."
        )


def test_every_benchmarking_region_is_blocked_missing_source(audit):
    regions = [e for e in audit.entries if e.domain == RegionDomain.BENCHMARKING.value]
    assert regions, "expected at least one benchmarking region in the template"
    for e in regions:
        assert e.readiness == Readiness.BLOCKED_MISSING_SOURCE.value
        assert e.source_coverage == SourceCoverage.MISSING_SOURCE.value
        assert e.blocking_reasons
        assert e.required_artifact_type == ArtifactType.BENCHMARKING_REPORT.value


def test_benchmarking_gap_names_the_minimum_required_artifact(audit):
    minimum = audit.benchmarking_audit["minimum_required_artifact"]
    assert minimum["type"] == ArtifactType.BENCHMARKING_REPORT.value
    content = " ".join(minimum["content"]).lower()
    for required in ("comparable", "screening", "independence", "quartile", "pli"):
        assert required in content, f"minimum artifact does not name {required}"


# ============================================================================
# NARRATIVE GAPS (§10)
# ============================================================================

def test_narrative_source_gap_is_detected(audit):
    """Financial workbooks are not a source for FAR / business narrative."""
    present = ExtendedSourceProfiler.roles_present(audit.sheet_profiles)
    assert SourceDatasetRole.NARRATIVE_DATA.value not in present
    assert SourceDatasetRole.ORGANIZATIONAL_DATA.value not in present
    assert SourceDatasetRole.CONTRACTUAL_DATA.value not in present

    far = [e for e in audit.entries if e.domain == RegionDomain.FUNCTIONAL_ANALYSIS.value]
    assert far, "expected functional-analysis regions"
    for e in far:
        assert e.readiness == Readiness.BLOCKED_MISSING_SOURCE.value
        assert e.required_artifact_type == ArtifactType.MANAGEMENT_INFORMATION.value


def test_narrative_gaps_state_the_artifact_kind_not_just_more_data(audit):
    generic = {"", "more data", "additional data", "data"}
    for e in audit.entries:
        if not e.readiness.startswith("BLOCKED"):
            continue
        assert e.required_artifact_type not in generic
        assert e.required_artifact_type in {a.value for a in ArtifactType}
        assert e.required_artifact_type != ArtifactType.NONE_REQUIRED.value


# ============================================================================
# FIGURES (§11)
# ============================================================================

def test_every_drawing_construct_is_assessed_individually(audit):
    assert audit.figures, "expected drawing constructs in the template"
    ids = [f.figure_id for f in audit.figures]
    assert len(ids) == len(set(ids)), "figure ids must be unique"
    for f in audit.figures:
        assert f.kind in FigureKind
        assert f.readiness in FigureReadiness
        assert f.rationale, f"{f.figure_id} has no rationale"
        assert f.paragraph_index > 0


def test_figures_are_not_assumed_uniformly_updateable(audit):
    kinds = {f.kind for f in audit.figures}
    readiness = {f.readiness for f in audit.figures}
    assert len(kinds) > 1, "figures must be distinguished by kind"
    assert len(readiness) > 1, "figures must be distinguished by readiness"
    assert FigureReadiness.UPDATEABLE not in readiness or all(
        f.required_artifact != ArtifactType.NONE_REQUIRED
        for f in audit.figures if f.readiness == FigureReadiness.UPDATEABLE)


def test_layout_text_boxes_are_not_treated_as_data_figures(audit):
    boxes = [f for f in audit.figures if f.kind == FigureKind.LAYOUT_TEXT_BOX]
    assert boxes, "expected layout text boxes in this template"
    for f in boxes:
        assert f.readiness == FigureReadiness.STATIC_PRESERVE
        assert f.required_artifact == ArtifactType.NONE_REQUIRED


def test_template_has_no_raster_media_and_that_is_recorded(audit):
    """The package genuinely contains no images or charts."""
    import zipfile

    with zipfile.ZipFile(str(PATH_TMPL)) as z:
        names = z.namelist()
    assert not [n for n in names if n.startswith("word/media/")]
    assert not [n for n in names if "chart" in n.lower()]
    assert any(n.startswith("word/diagrams/") for n in names), "expected a SmartArt part"

    assert all(f.has_raster_image is False for f in audit.figures)
    assert any(f.is_smartart for f in audit.figures), "the SmartArt diagram must be identified"
    assert audit.scorecards["figure_totals"]["data_figures_automatable"] == 0


def test_data_figures_require_a_named_artifact_kind(audit):
    data_figures = [f for f in audit.figures if f.kind != FigureKind.LAYOUT_TEXT_BOX]
    assert data_figures
    for f in data_figures:
        assert f.readiness != FigureReadiness.UPDATEABLE
        assert f.required_artifact != ArtifactType.NONE_REQUIRED


# ============================================================================
# HISTORICAL TEXT (§12)
# ============================================================================

def test_historical_text_is_never_automatically_current_verified(audit):
    valid = {h.value for h in HistoricalReuse}
    for e in audit.entries:
        assert e.historical_reuse in valid
        # There is no "VERIFIED_CURRENT" class, and no region may claim one.
        assert "VERIFIED_CURRENT" not in e.historical_reuse


def test_year_specific_domains_require_current_year_source(audit):
    year_specific = {
        RegionDomain.RELATED_PARTY_TRANSACTIONS.value,
        RegionDomain.FINANCIAL_INFORMATION.value,
        RegionDomain.BENCHMARKING.value,
        RegionDomain.ORGANIZATION.value,
        RegionDomain.INTERCOMPANY_AGREEMENTS.value,
    }
    for e in audit.entries:
        if e.domain in year_specific:
            assert e.historical_reuse == HistoricalReuse.REQUIRES_CURRENT_YEAR_SOURCE.value, (
                f"{e.region_id} ({e.domain}) must not permit historical reuse")


def test_statutory_text_is_the_only_reusable_baseline(audit):
    baseline = [e for e in audit.entries
                if e.historical_reuse == HistoricalReuse.REUSABLE_BASELINE.value]
    assert baseline, "expected statutory regions"
    for e in baseline:
        assert e.domain == RegionDomain.STATUTORY_METHODOLOGY.value


def test_narrative_domains_require_verification_not_blind_reuse(audit):
    for e in audit.entries:
        if e.domain in (RegionDomain.BUSINESS_NARRATIVE.value,
                        RegionDomain.FUNCTIONAL_ANALYSIS.value,
                        RegionDomain.COMPETITOR_ANALYSIS.value):
            assert e.historical_reuse == HistoricalReuse.REUSABLE_WITH_VERIFICATION.value


# ============================================================================
# READINESS INTEGRITY
# ============================================================================

def test_segmentation_is_an_exact_partition_of_the_document(audit):
    """No body element may be lost or counted twice by the region segmenter."""
    seg = audit.scorecards["segmentation"]
    assert seg["is_exact_partition"] is True
    assert seg["paragraphs_assigned"] == seg["paragraphs_unique"] == seg["body_paragraphs"]
    assert seg["tables_assigned"] == seg["tables_unique"] == seg["body_tables"]
    assert seg["regions"] == len(audit.regions)

    all_tables = sorted(t for r in audit.regions for t in r.table_ordinals)
    assert all_tables == list(range(seg["body_tables"]))


def test_every_region_has_a_readiness_classification(audit):
    valid = {r.value for r in Readiness}
    assert len(audit.entries) == len(audit.regions)
    for e in audit.entries:
        assert e.readiness in valid
        assert e.readiness != "READY", "bare READY is forbidden; readiness must name its kind"


def test_every_blocked_region_has_a_concrete_reason(audit):
    for e in audit.entries:
        if e.readiness.startswith("BLOCKED"):
            assert e.blocking_reasons, f"{e.region_id} blocked without a reason"
            assert all(len(r) > 20 for r in e.blocking_reasons)


def test_blocked_regions_stay_blocked_and_none_is_auto_mutation_ready(audit):
    """The audit must not promote any region to automatic mutation."""
    counts = audit.scorecards["region_readiness_counts"]
    assert counts.get(Readiness.AUTO_MUTATION_READY.value, 0) == 0
    blocked = [e for e in audit.entries if e.readiness.startswith("BLOCKED")]
    assert blocked, "the current source set cannot support every region"
    for e in blocked:
        assert e.source_coverage in (SourceCoverage.MISSING_SOURCE.value,
                                     SourceCoverage.INSUFFICIENT_EVIDENCE.value)


def test_coverage_levels_are_not_collapsed_into_one_bucket(audit):
    used = {e.source_coverage for e in audit.entries}
    assert len(used) >= 4, f"coverage statuses collapsed: {used}"
    assert SourceCoverage.MISSING_SOURCE.value in used
    assert SourceCoverage.FULLY_SUPPORTED.value in used


def test_scorecard_percentages_are_derived_from_counts(audit):
    for domain, d in audit.scorecards["domain_completeness"].items():
        supported = d["fully"] + d["partial"] + d["not_applicable"]
        expected = round(100.0 * supported / d["total"], 1) if d["total"] else 0.0
        assert d["supported_pct"] == expected, f"{domain} percentage is not a count ratio"
        assert d["total"] == len(d["regions"])


def test_required_artifacts_cover_every_blocked_region(audit):
    covered = {r for a in audit.artifacts for r in a.affected_regions}
    for e in audit.entries:
        if e.readiness.startswith("BLOCKED"):
            assert e.region_id in covered, f"{e.region_id} is blocked but no artifact requests it"


def test_priorities_are_not_inflated(audit):
    for a in audit.artifacts:
        assert a.priority in {p.value for p in Priority}
        if a.priority == Priority.BLOCKING.value:
            assert a.blocking_status is True
            assert any(
                e.readiness.startswith("BLOCKED")
                for e in audit.entries if e.region_id in a.affected_regions)
        else:
            assert a.blocking_status is False


def test_human_review_queue_is_actionable(audit):
    assert audit.review_queue
    for h in audit.review_queue:
        assert h.region_id and h.heading
        assert h.issue and len(h.issue) > 20
        assert h.suggested_human_evidence_request
        assert isinstance(h.blocking, bool)
        assert h.missing_evidence


# ============================================================================
# GROUND TRUTH CANNOT AFFECT READINESS (§17)
# ============================================================================

def test_ground_truth_is_quarantined_during_the_audit():
    GroundTruthQuarantine.arm([PATH_GROUND_TRUTH_FORBIDDEN])
    try:
        with pytest.raises(GroundTruthContaminationError):
            GroundTruthQuarantine.open_document(PATH_GROUND_TRUTH_FORBIDDEN)
    finally:
        GroundTruthQuarantine.disarm()


def test_evaluation_refuses_to_run_before_freeze(audit):
    import dataclasses

    unfrozen = dataclasses.replace(audit, frozen=False)
    with pytest.raises(RuntimeError, match="before readiness was frozen"):
        GapEvaluation.evaluate(unfrozen)


def test_ground_truth_evaluation_does_not_alter_readiness(audit):
    before = [(e.region_id, e.readiness, e.source_coverage) for e in audit.entries]
    before_scorecards = json.dumps(audit.scorecards, sort_keys=True)

    evaluation = GapEvaluation.evaluate(audit)

    after = [(e.region_id, e.readiness, e.source_coverage) for e in audit.entries]
    assert after == before
    assert json.dumps(audit.scorecards, sort_keys=True) == before_scorecards
    assert evaluation["readiness_altered"] is False
    assert evaluation["evaluated_after_freeze"] is True
    assert evaluation["findings"]


def test_readiness_artifact_declares_no_ground_truth_use(artifacts):
    readiness, matrix = artifacts
    assert readiness["ground_truth_used_in_readiness"] is False
    assert matrix["ground_truth_used"] is False
    assert PATH_GROUND_TRUTH_FORBIDDEN.name not in json.dumps(readiness)
    assert PATH_GROUND_TRUTH_FORBIDDEN.name not in json.dumps(matrix)


# ============================================================================
# DETERMINISM & NO FABRICATION
# ============================================================================

def test_repeated_audit_produces_equivalent_output():
    """A second run must produce byte-equivalent artifacts."""
    _r1, readiness1, matrix1, _e1 = run_audit(write=False)
    _r2, readiness2, matrix2, _e2 = run_audit(write=False)
    assert json.dumps(readiness1, sort_keys=True) == json.dumps(readiness2, sort_keys=True)
    assert json.dumps(matrix1, sort_keys=True) == json.dumps(matrix2, sort_keys=True)


def test_requirement_matrix_is_deterministic_in_order(audit):
    ids = [e.region_id for e in audit.entries]
    assert ids == sorted(ids), "region ids must be emitted in stable order"
    assert len(ids) == len(set(ids))


def test_no_fabricated_source_artifacts(audit):
    """Every cited source must be a workbook+sheet that was actually profiled."""
    real = {f"{s['workbook']}!{s['sheet_name']}" for s in audit.sheet_profiles}
    for e in audit.entries:
        for cited in e.current_source_artifacts:
            assert cited in real, f"{e.region_id} cites a source that does not exist: {cited}"


def test_a_region_with_no_source_cites_no_source(audit):
    for e in audit.entries:
        if e.source_coverage == SourceCoverage.MISSING_SOURCE.value:
            assert e.current_source_artifacts == [], (
                f"{e.region_id} is MISSING_SOURCE yet cites {e.current_source_artifacts}")
            assert e.source_roles_found == []


def test_no_invented_cell_ranges(artifacts):
    """A readiness map binds nothing, so it must contain no cell addresses."""
    readiness, matrix = artifacts
    for name, payload in (("readiness", readiness), ("matrix", matrix)):
        for key in ("cell_address", "cell_range", "source_cell_address", "expected_precondition_hash"):
            assert key not in json.dumps(payload), f"{name} contains a binding field '{key}'"

    # No A1-style range literals in any region requirement field.
    range_re = re.compile(r"\b[A-Z]{1,3}\d{1,5}:[A-Z]{1,3}\d{1,5}\b")
    for e in matrix["requirement_matrix"]:
        blob = json.dumps({k: v for k, v in e.items() if k != "evidence"})
        assert not range_re.search(blob), f"{e['region_id']} contains an invented cell range"


def test_no_invented_element_ids(artifacts):
    """No UUID-shaped identifiers may appear: the audit invents no element ids."""
    uuid_re = re.compile(r"\b[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}\b", re.I)
    for payload in artifacts:
        found = uuid_re.findall(json.dumps(payload))
        assert not found, f"invented element ids: {found[:3]}"


def test_available_and_missing_sources_are_reported_separately(artifacts):
    _readiness, matrix = artifacts
    assert "currently_available_sources" in matrix
    assert "required_but_missing_sources" in matrix
    available = matrix["currently_available_sources"]["dataset_roles_present"]
    for role in BENCHMARKING_ROLES:
        assert role.value not in available
    assert matrix["required_but_missing_sources"], "missing sources must be enumerated"


def test_domain_rules_are_evidence_backed(audit):
    """Every classified region records which keyword classified it."""
    for e in audit.entries:
        if e.domain in (RegionDomain.UNKNOWN.value, RegionDomain.STRUCTURAL_CONTAINER.value):
            continue
        assert e.domain_evidence, f"{e.region_id} classified without evidence"
        assert ("heading contains" in e.domain_evidence
                or "inherited from parent section" in e.domain_evidence
                or "classified from contained table" in e.domain_evidence)


def test_domain_rule_table_is_internally_consistent():
    for rule in DOMAIN_RULES:
        assert rule.keywords and all(k == k.lower() for k in rule.keywords)
        assert rule.required_info
        assert rule.granularity
        assert rule.artifact_if_missing in ArtifactType
        assert rule.historical_reuse in HistoricalReuse
