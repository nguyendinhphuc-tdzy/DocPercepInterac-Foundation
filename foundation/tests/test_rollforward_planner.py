"""
Production Roll-Forward Planner (Phase H).

The planner turns authoritative workflow slots into a governed manifest and
mutation plan. These tests pin the properties that make it safe to run against
real client files:

    grounding        every planned value is read from a real source cell
    isolation        only the workflow's own documents are ever opened
    anti-contamination
                     reading Ground Truth is a hard failure, not a warning
    position-free    renumbering tables changes no planning decision
    honesty          absent evidence blocks; ambiguous evidence asks a human;
                     nothing is marked READY to make a fixture pass
    governance       the planner may propose and may not approve or execute
"""
from __future__ import annotations

import sys
import warnings
from pathlib import Path
from typing import List, Tuple

import openpyxl
import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from applications.rollforward.models import (  # noqa: E402
    ExecutionGate,
    ManifestStatus,
    SourceBindingStatus,
)
from applications.rollforward.planner import (  # noqa: E402
    DocumentRef,
    PlannerContaminationError,
    PlannerError,
    PlannerInputs,
    PlanningResult,
    RollForwardPlanner,
    plan_digest,
)
from applications.rollforward.source_intake import GroundTruthGuard  # noqa: E402
from applications.rollforward.table_identity import TableCorrespondenceResolver  # noqa: E402
from applications.rollforward.workflow_intake import FiscalPeriod  # noqa: E402

warnings.filterwarnings("ignore")

DEMO = Path(__file__).resolve().parents[2] / "anonymize client" / "Demo files" / "Demo files"
HISTORICAL = DEMO / "Compare LF" / "HMV-24-Final-Local File for FY2023-EN-R0303KPMG.docx"
TEMPLATE = (DEMO / "Compare LF"
            / "Client-25-Template-Local File for FY20XX-Manufacturer-EN-RddmmKPMG-13062025 "
              "(Decree 20-2025).docx")
FA_RPT = DEMO / "FA&RPTS & Appendix I" / "FA&RPTs" / "HMV-FA&RPT FY2024.xlsx"
APPENDIX_I = (DEMO / "FA&RPTS & Appendix I" / "Appendix I"
              / "HMV-25-Appendix I under D20 for FY2024-Final-W3103.xlsx")
# Evaluation-only oracle. The planner must never open it.
GROUND_TRUTH = DEMO / "Compare LF" / "HMV-26-Final-Local File for FY2024-EN-R2901KPMG.docx"

requires_demo = pytest.mark.skipif(
    not (HISTORICAL.exists() and TEMPLATE.exists() and FA_RPT.exists()),
    reason="demo fixtures are not present")


# ===========================================================================
# SYNTHETIC FIXTURES — a workflow whose evidence genuinely lines up
# ===========================================================================

RPT_HEADERS = ["Transaction", "Related party", "Type of relationship", "Amount (VND)"]
RPT_RECORDS = [
    ["Purchase of raw materials", "Hestra Japan", "Parent company", "112233445566"],
    ["Sale of finished goods", "Hestra Singapore", "Fellow subsidiary", "778899001122"],
    ["Technical service fee", "Hestra Japan", "Parent company", "334455667788"],
]


def _write_table_document(path: Path, headers: List[str], rows: List[List[str]],
                          heading: str = "Related party transactions",
                          decoys: int = 0) -> Path:
    """A document holding one labelled table, optionally behind decoy tables."""
    document = Document()
    for i in range(decoys):
        document.add_heading(f"Unrelated section {i}", level=1)
        decoy = document.add_table(rows=2, cols=2)
        decoy.rows[0].cells[0].text = f"Decoy header {i}A"
        decoy.rows[0].cells[1].text = f"Decoy header {i}B"
        decoy.rows[1].cells[0].text = "value"
        decoy.rows[1].cells[1].text = "value"

    document.add_heading(heading, level=1)
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    for c, header in enumerate(headers):
        table.rows[0].cells[c].text = header
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            table.rows[r].cells[c].text = value
    document.add_paragraph("Trailing narrative paragraph.")
    document.save(str(path))
    return path


def _write_source_workbook(path: Path, headers: List[str], records: List[List[str]],
                           sheet_name: str = "RPTs") -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for record in records:
        sheet.append(record)
    workbook.save(str(path))
    return path


@pytest.fixture
def aligned_workflow(tmp_path) -> PlannerInputs:
    """Template + historical + a source whose columns match, with room to grow."""
    template = _write_table_document(
        tmp_path / "template.docx", RPT_HEADERS, [["", "", "", ""]])
    historical = _write_table_document(
        tmp_path / "historical.docx", RPT_HEADERS, RPT_RECORDS[:1])
    source = _write_source_workbook(tmp_path / "source.xlsx", RPT_HEADERS, RPT_RECORDS)

    return PlannerInputs(
        workflow_id="wf-synth", session_id="sess-synth",
        historical=DocumentRef("doc-hist", historical, historical.name),
        template=DocumentRef("doc-tpl", template, template.name),
        current_sources=(DocumentRef("doc-src", source, source.name),),
        historical_period=FiscalPeriod(2023), current_period=FiscalPeriod(2024),
    )


@pytest.fixture
def production_workflow() -> PlannerInputs:
    return PlannerInputs(
        workflow_id="wf-prod", session_id="sess-prod",
        historical=DocumentRef("doc-hist", HISTORICAL, HISTORICAL.name),
        template=DocumentRef("doc-tpl", TEMPLATE, TEMPLATE.name),
        current_sources=(
            DocumentRef("doc-fa", FA_RPT, FA_RPT.name),
            DocumentRef("doc-app", APPENDIX_I, APPENDIX_I.name),
        ),
        historical_period=FiscalPeriod(2023), current_period=FiscalPeriod(2024),
    )


# ===========================================================================
# 1. PLANNING PRODUCES A GOVERNED PLAN
# ===========================================================================

def test_planner_produces_a_manifest_and_plan(aligned_workflow):
    result = RollForwardPlanner.plan(aligned_workflow)

    assert isinstance(result, PlanningResult)
    assert result.manifest.regions, "no regions were planned"
    assert result.manifest.manifest_id
    assert result.manifest.manifest_version == 1


def test_aligned_evidence_reaches_ready(aligned_workflow):
    """The READY path must be reachable — otherwise nothing is ever executable."""
    result = RollForwardPlanner.plan(aligned_workflow)
    ready = [o for o in result.outcomes if o.disposition == "READY"]

    assert ready, [o.to_dict() for o in result.outcomes]
    assert result.has_executable_work
    assert result.mutation_plan.table_mutations


def test_planned_values_are_read_from_real_source_cells(aligned_workflow):
    """Grounding: every planned cell names the sheet and address it came from."""
    result = RollForwardPlanner.plan(aligned_workflow)
    source_path = aligned_workflow.current_sources[0].path
    workbook = openpyxl.load_workbook(str(source_path), data_only=True)

    checked = 0
    for table in result.mutation_plan.table_mutations:
        for row in table.row_mutations:
            for cell in row.cells:
                assert cell.source_sheet, "a planned cell has no source sheet"
                assert cell.source_cell_address, "a planned cell has no source address"
                actual = workbook[cell.source_sheet][cell.source_cell_address].value
                assert str(actual) == cell.value, (
                    f"{cell.source_sheet}!{cell.source_cell_address} holds {actual!r}, "
                    f"the plan says {cell.value!r}")
                checked += 1
    workbook.close()
    assert checked > 0, "no planned cells to verify"


def test_structural_delta_comes_from_template_and_evidence(aligned_workflow):
    result = RollForwardPlanner.plan(aligned_workflow)
    region = next(r for r in result.manifest.regions
                  if r.execution_gate == ExecutionGate.READY)

    delta = region.structural_delta
    assert delta is not None
    assert delta.target_rows == 1 + len(RPT_RECORDS)     # header + one row per record
    assert delta.insert_count == delta.target_rows - delta.template_rows
    assert delta.observation_source == "template_structure+current_year_evidence"
    assert "source_record_count" in delta.observation_context


def test_every_executable_region_carries_validation_rules(aligned_workflow):
    result = RollForwardPlanner.plan(aligned_workflow)
    for region in result.manifest.regions:
        if region.execution_gate == ExecutionGate.READY:
            kinds = {rule.rule_type.value for rule in region.validation_rules}
            assert "ROW_COUNT_MATCH" in kinds
            assert "COLUMN_COUNT_UNCHANGED" in kinds
            assert "SOURCE_VALUE_PRESENT" in kinds


def test_row_template_is_only_created_when_cloning_is_safe(aligned_workflow):
    result = RollForwardPlanner.plan(aligned_workflow)
    for region in result.manifest.regions:
        if region.row_template is not None:
            assert region.row_template.safe_to_clone is True
            assert region.row_template.row_anchor.startswith(region.region_id)


def test_diffs_describe_the_planned_change(aligned_workflow):
    result = RollForwardPlanner.plan(aligned_workflow)
    assert result.diffs
    diff = result.diffs[0]
    assert diff.before_summary["rows"] < diff.after_summary["rows"]
    assert diff.delta_details[0]["rows_inserted"] > 0


# ===========================================================================
# 2. ANTI-CONTAMINATION — GROUND TRUTH IS NEVER READ
# ===========================================================================

@requires_demo
def test_planner_never_opens_ground_truth(production_workflow):
    """Fails loudly if the planner touches the FY2024 final Local File."""
    opened: List[str] = []
    import docx as docx_module
    import openpyxl as openpyxl_module
    import applications.rollforward.table_identity as table_identity_module

    real_document = docx_module.Document
    real_load = openpyxl_module.load_workbook
    # The profiler binds `Document` at import time, so patching only `docx` would
    # leave every DOCX read invisible to this spy.
    real_profiler_document = table_identity_module.Document

    def spy_document(path=None, *args, **kwargs):
        if path:
            opened.append(str(path))
            if Path(str(path)).name == GROUND_TRUTH.name:
                raise AssertionError("the planner opened Ground Truth")
        return real_document(path, *args, **kwargs)

    def spy_load(filename=None, *args, **kwargs):
        if filename:
            opened.append(str(filename))
            if Path(str(filename)).name == GROUND_TRUTH.name:
                raise AssertionError("the planner opened Ground Truth")
        return real_load(filename, *args, **kwargs)

    docx_module.Document = spy_document
    openpyxl_module.load_workbook = spy_load
    table_identity_module.Document = spy_document
    try:
        result = RollForwardPlanner.plan(production_workflow)
    finally:
        docx_module.Document = real_document
        openpyxl_module.load_workbook = real_load
        table_identity_module.Document = real_profiler_document

    assert opened, "the spy recorded nothing — it is not observing the planner's reads"
    assert not any(GROUND_TRUTH.name in path for path in opened)
    assert result.report()["ground_truth_accessed"] is False
    for record in result.accessed_documents:
        assert GROUND_TRUTH.name not in record


@requires_demo
def test_ground_truth_in_a_slot_is_refused(tmp_path):
    """Even renamed into a slot, Ground Truth is refused — by content hash."""
    disguised = tmp_path / "innocent-looking-template.docx"
    disguised.write_bytes(GROUND_TRUTH.read_bytes())
    GroundTruthGuard.register([GROUND_TRUTH])
    try:
        inputs = PlannerInputs(
            workflow_id="wf", session_id="sess",
            historical=DocumentRef("doc-hist", HISTORICAL, HISTORICAL.name),
            template=DocumentRef("doc-tpl", disguised, disguised.name),
            current_sources=(DocumentRef("doc-fa", FA_RPT, FA_RPT.name),),
            historical_period=FiscalPeriod(2023), current_period=FiscalPeriod(2024),
        )
        with pytest.raises(Exception) as exc:
            RollForwardPlanner.plan(inputs)
        assert "ground truth" in str(exc.value).lower()
    finally:
        GroundTruthGuard.reset()


def test_a_path_outside_the_workflow_is_refused(aligned_workflow, tmp_path):
    from applications.rollforward.planner import InputAccessGuard

    guard = InputAccessGuard(aligned_workflow)
    stranger = tmp_path / "not-an-input.docx"
    stranger.write_bytes(b"x")

    with pytest.raises(PlannerContaminationError, match="not"):
        guard.check(stranger, "curiosity")


def test_the_report_lists_exactly_the_documents_read(aligned_workflow):
    result = RollForwardPlanner.plan(aligned_workflow)
    names = {record.split(" (")[0] for record in result.accessed_documents}

    assert names == {
        aligned_workflow.template.filename,
        aligned_workflow.historical.filename,
        aligned_workflow.current_sources[0].filename,
    }


# ===========================================================================
# 3. POSITION-FREE IDENTITY
# ===========================================================================

def test_correspondence_scoring_cannot_read_position():
    assert TableCorrespondenceResolver.assert_no_positional_input() is True


def test_renumbering_tables_changes_no_planning_decision(tmp_path):
    """Insert unrelated tables before the target; decisions must be identical."""
    plain_template = _write_table_document(
        tmp_path / "t_plain.docx", RPT_HEADERS, [["", "", "", ""]])
    plain_historical = _write_table_document(
        tmp_path / "h_plain.docx", RPT_HEADERS, RPT_RECORDS[:1])
    shifted_template = _write_table_document(
        tmp_path / "t_shift.docx", RPT_HEADERS, [["", "", "", ""]], decoys=3)
    shifted_historical = _write_table_document(
        tmp_path / "h_shift.docx", RPT_HEADERS, RPT_RECORDS[:1], decoys=2)
    source = _write_source_workbook(tmp_path / "src.xlsx", RPT_HEADERS, RPT_RECORDS)
    renamed_source = _write_source_workbook(
        tmp_path / "totally-different-name.xlsx", RPT_HEADERS, RPT_RECORDS)

    def _plan(template, historical, src, doc_ids):
        return RollForwardPlanner.plan(PlannerInputs(
            workflow_id="wf", session_id="sess",
            historical=DocumentRef(doc_ids[0], historical, historical.name),
            template=DocumentRef(doc_ids[1], template, template.name),
            current_sources=(DocumentRef(doc_ids[2], src, src.name),),
            historical_period=FiscalPeriod(2023), current_period=FiscalPeriod(2024),
        ))

    baseline = _plan(plain_template, plain_historical, source, ("h1", "t1", "s1"))
    # Renumbered tables, different document ids, different source filename.
    shifted = _plan(shifted_template, shifted_historical, renamed_source, ("h2", "t2", "s2"))

    def decisions(result):
        return {o.identity_key: (o.disposition, o.binding_verdict)
                for o in result.outcomes}

    baseline_decisions = decisions(baseline)
    shifted_decisions = decisions(shifted)

    # Every table the baseline planned appears in the shifted run with the same
    # identity key and the same decision; the decoys add their own entries.
    for identity_key, decision in baseline_decisions.items():
        assert identity_key in shifted_decisions, "a table lost its identity when renumbered"
        assert shifted_decisions[identity_key] == decision, identity_key


def test_region_ids_follow_content_not_position(tmp_path):
    template = _write_table_document(tmp_path / "a.docx", RPT_HEADERS, [["", "", "", ""]])
    shifted = _write_table_document(tmp_path / "b.docx", RPT_HEADERS, [["", "", "", ""]], decoys=4)
    historical = _write_table_document(tmp_path / "h.docx", RPT_HEADERS, RPT_RECORDS[:1])
    source = _write_source_workbook(tmp_path / "s.xlsx", RPT_HEADERS, RPT_RECORDS)

    def region_ids(template_path):
        result = RollForwardPlanner.plan(PlannerInputs(
            workflow_id="wf", session_id="sess",
            historical=DocumentRef("h", historical, historical.name),
            template=DocumentRef("t", template_path, template_path.name),
            current_sources=(DocumentRef("s", source, source.name),),
            historical_period=FiscalPeriod(2023), current_period=FiscalPeriod(2024)))
        return {o.region_id for o in result.outcomes if o.disposition == "READY"}

    assert region_ids(template) == region_ids(shifted)
    assert region_ids(template), "the ready region disappeared"


# ===========================================================================
# 4. SOURCE BINDING — NEGATIVE CASES
# ===========================================================================

def _single_region_result(tmp_path, *, source_headers=None, records=None,
                          source_present=True, sheet_name="RPTs"):
    template = _write_table_document(tmp_path / "tpl.docx", RPT_HEADERS, [["", "", "", ""]])
    historical = _write_table_document(tmp_path / "hist.docx", RPT_HEADERS, RPT_RECORDS[:1])
    sources: Tuple[DocumentRef, ...] = ()
    if source_present:
        source = _write_source_workbook(
            tmp_path / "src.xlsx", source_headers or RPT_HEADERS,
            records if records is not None else RPT_RECORDS, sheet_name=sheet_name)
        sources = (DocumentRef("doc-src", source, source.name),)

    return RollForwardPlanner.plan(PlannerInputs(
        workflow_id="wf", session_id="sess",
        historical=DocumentRef("doc-hist", historical, historical.name),
        template=DocumentRef("doc-tpl", template, template.name),
        current_sources=sources,
        historical_period=FiscalPeriod(2023), current_period=FiscalPeriod(2024)))


def test_missing_source_blocks_the_region(tmp_path):
    result = _single_region_result(tmp_path, source_present=False)
    dispositions = {o.disposition for o in result.outcomes}

    assert "READY" not in dispositions
    assert "BLOCKED" in dispositions
    assert not result.mutation_plan.table_mutations


def test_a_source_with_no_matching_columns_cannot_bind(tmp_path):
    result = _single_region_result(
        tmp_path, source_headers=["Wholly", "Unrelated", "Column", "Names"])

    for outcome in result.outcomes:
        assert outcome.disposition != "READY"
    for region in result.manifest.regions:
        for binding in region.current_sources:
            assert binding.status != SourceBindingStatus.VERIFIED or binding.cell_range is None


def test_ambiguous_sources_ask_a_human(tmp_path):
    """Two sheets that both carry the role must not be chosen between."""
    template = _write_table_document(tmp_path / "tpl.docx", RPT_HEADERS, [["", "", "", ""]])
    historical = _write_table_document(tmp_path / "hist.docx", RPT_HEADERS, RPT_RECORDS[:1])

    path = tmp_path / "two_sheets.xlsx"
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "RPTs A"
    first.append(RPT_HEADERS)
    for record in RPT_RECORDS:
        first.append(record)
    second = workbook.create_sheet("RPTs B")
    second.append(RPT_HEADERS)
    for record in RPT_RECORDS:
        second.append(record)
    workbook.save(str(path))

    result = RollForwardPlanner.plan(PlannerInputs(
        workflow_id="wf", session_id="sess",
        historical=DocumentRef("doc-hist", historical, historical.name),
        template=DocumentRef("doc-tpl", template, template.name),
        current_sources=(DocumentRef("doc-src", path, path.name),),
        historical_period=FiscalPeriod(2023), current_period=FiscalPeriod(2024)))

    statuses = {b.status for r in result.manifest.regions for b in r.current_sources}
    assert SourceBindingStatus.AMBIGUOUS in statuses
    assert not result.mutation_plan.table_mutations


def test_a_document_source_cannot_back_a_cell_binding(tmp_path):
    """Historical-only / narrative evidence is recorded as rejected, not used."""
    template = _write_table_document(tmp_path / "tpl.docx", RPT_HEADERS, [["", "", "", ""]])
    historical = _write_table_document(tmp_path / "hist.docx", RPT_HEADERS, RPT_RECORDS[:1])
    narrative = _write_table_document(tmp_path / "narrative.docx", RPT_HEADERS, RPT_RECORDS)

    result = RollForwardPlanner.plan(PlannerInputs(
        workflow_id="wf", session_id="sess",
        historical=DocumentRef("doc-hist", historical, historical.name),
        template=DocumentRef("doc-tpl", template, template.name),
        current_sources=(DocumentRef("doc-doc", narrative, narrative.name),),
        historical_period=FiscalPeriod(2023), current_period=FiscalPeriod(2024)))

    assert any(r["reason"] == "NOT_CELL_ADDRESSABLE" for r in result.rejected_evidence)
    assert not result.mutation_plan.table_mutations


def test_template_only_evidence_cannot_satisfy_a_current_year_requirement(aligned_workflow):
    """The template is a structure, never a source: it supplies no bindings."""
    result = RollForwardPlanner.plan(aligned_workflow)
    for region in result.manifest.regions:
        for binding in region.current_sources:
            assert binding.source_doc_id != aligned_workflow.template.document_id
            assert binding.source_doc_name != aligned_workflow.template.filename


def test_no_cell_address_is_hardcoded(aligned_workflow):
    """Every address must be derivable from the profiled header row."""
    result = RollForwardPlanner.plan(aligned_workflow)
    addresses = [cell.source_cell_address
                 for table in result.mutation_plan.table_mutations
                 for row in table.row_mutations for cell in row.cells]

    assert addresses
    source = openpyxl.load_workbook(str(aligned_workflow.current_sources[0].path))
    sheet = source[result.mutation_plan.table_mutations[0].row_mutations[0].cells[0].source_sheet]
    for address in addresses:
        assert sheet[address].value is not None, f"{address} is empty in the source"
    source.close()


@requires_demo
def test_benchmarking_stays_blocked_without_a_benchmarking_source(production_workflow):
    """The production fixtures contain no current-year benchmarking dataset."""
    result = RollForwardPlanner.plan(production_workflow)
    benchmarking = [o for o in result.outcomes
                    if "arm" in o.section_name.lower() or "range" in o.section_name.lower()
                    or "comparable" in o.section_name.lower()]

    assert benchmarking, "no benchmarking-shaped region was planned"
    assert all(o.disposition != "READY" for o in benchmarking), \
        [o.to_dict() for o in benchmarking]


# ===========================================================================
# 5. PRODUCTION FIXTURE ACCEPTANCE
# ===========================================================================

@requires_demo
def test_the_production_fixture_does_not_claim_full_readiness(production_workflow):
    result = RollForwardPlanner.plan(production_workflow)
    summary = result.readiness_summary()

    assert summary["regions_total"] > 0
    assert summary["by_disposition"]["READY"] < summary["regions_total"], \
        "the planner claimed the whole Local File was ready"
    assert result.unresolved_blockers(), "no blockers were reported for incomplete evidence"


@requires_demo
def test_the_production_manifest_never_leaves_planning(production_workflow):
    result = RollForwardPlanner.plan(production_workflow)

    assert result.manifest.status in (ManifestStatus.PLANNED, ManifestStatus.REVIEW_REQUIRED)
    assert result.manifest.approved_by is None
    assert result.manifest.approved_at is None


@requires_demo
def test_planning_does_not_modify_any_input(production_workflow):
    before = {doc.path: (doc.path.stat().st_size, doc.path.stat().st_mtime_ns)
              for doc in production_workflow.all_documents()}

    RollForwardPlanner.plan(production_workflow)

    for path, stamp in before.items():
        assert (path.stat().st_size, path.stat().st_mtime_ns) == stamp, f"{path.name} changed"


# ===========================================================================
# 6. SCHEMA
# ===========================================================================

def test_manifest_carries_the_required_schema(aligned_workflow):
    result = RollForwardPlanner.plan(aligned_workflow)
    manifest = result.manifest

    assert manifest.manifest_id and manifest.manifest_version
    assert manifest.session_id == aligned_workflow.session_id
    assert manifest.historical_document_id == aligned_workflow.historical.document_id
    assert manifest.template_document_id == aligned_workflow.template.document_id
    assert manifest.current_source_document_ids == [
        s.document_id for s in aligned_workflow.current_sources]
    assert manifest.regions

    report = result.report()
    for key in ("inputs", "manifest", "mutation_plan", "readiness_summary",
                "unresolved_blockers", "regions", "rejected_evidence",
                "documents_read", "timings_ms", "governance"):
        assert key in report, key
    assert report["inputs"]["historical_period"] == "FY2023"
    assert report["inputs"]["current_period"] == "FY2024"


def test_mutation_plan_carries_the_required_schema(aligned_workflow):
    result = RollForwardPlanner.plan(aligned_workflow)
    plan = result.mutation_plan

    assert plan.plan_id
    assert plan.manifest_id == result.manifest.manifest_id
    assert plan.manifest_version == result.manifest.manifest_version
    assert plan.target_doc_name == aligned_workflow.template.filename

    for table in plan.table_mutations:
        assert table.target_region_id.startswith("rfr-")
        assert table.table_hash
        assert table.expected_precondition_hash
        assert table.target_row_count > table.initial_row_count
        for row in table.row_mutations:
            for cell in row.cells:
                assert cell.source_doc_name and cell.source_sheet and cell.source_cell_address

    assert len(plan_digest(plan)) == 16


def test_the_plan_digest_is_stable_and_content_derived(aligned_workflow):
    first = RollForwardPlanner.plan(aligned_workflow)
    second = RollForwardPlanner.plan(aligned_workflow)
    assert plan_digest(first.mutation_plan) == plan_digest(second.mutation_plan)


# ===========================================================================
# 7. GOVERNANCE
# ===========================================================================

def test_the_planner_cannot_approve_or_execute(aligned_workflow):
    result = RollForwardPlanner.plan(aligned_workflow)

    assert result.manifest.status is not ManifestStatus.APPROVED
    assert result.manifest.status is not ManifestStatus.EXECUTING
    assert result.manifest.status is not ManifestStatus.COMPLETED
    assert result.manifest.approved_by is None
    assert result.report()["governance"]["approved"] is False


def test_planning_writes_no_output_document(aligned_workflow, tmp_path):
    before = {p.name for p in tmp_path.iterdir()}
    RollForwardPlanner.plan(aligned_workflow)
    assert {p.name for p in tmp_path.iterdir()} == before


def test_periods_must_be_known_and_ordered(aligned_workflow):
    unknown = PlannerInputs(
        workflow_id="wf", session_id="sess",
        historical=aligned_workflow.historical, template=aligned_workflow.template,
        current_sources=aligned_workflow.current_sources,
        historical_period=None, current_period=FiscalPeriod(2024))
    with pytest.raises(PlannerError, match="both fiscal periods"):
        RollForwardPlanner.plan(unknown)

    inverted = PlannerInputs(
        workflow_id="wf", session_id="sess",
        historical=aligned_workflow.historical, template=aligned_workflow.template,
        current_sources=aligned_workflow.current_sources,
        historical_period=FiscalPeriod(2025), current_period=FiscalPeriod(2024))
    with pytest.raises(PlannerError, match="Cannot plan"):
        RollForwardPlanner.plan(inverted)


def test_planning_consults_no_model(aligned_workflow):
    """There is no provider import in the planner, and none is reachable from it."""
    source = (Path(__file__).resolve().parents[1] / "applications" / "rollforward"
              / "planner.py").read_text(encoding="utf-8")

    for forbidden in ("providers", "get_provider", "gemini", "workbench", "openai"):
        assert forbidden not in source.lower(), forbidden
