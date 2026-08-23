"""
Source Intake & Evidence Contract Tests (Phase F)
==================================================
Location: foundation/tests/test_rollforward_source_intake.py

Locks the Phase F contract:

    artifact hashing and duplicate detection
    source package versioning and freezing
    content-derived dataset-role profiling
    benchmarking dataset detection (and non-detection by filename)
    narrative source detection
    stale input detection
    source request satisfaction
    readiness recalculation that can never authorize execution
    ground-truth quarantine
    region hierarchy consistency (the 61 / 81 / 104 resolution)
    deterministic repeated profiling

Synthetic workbooks built inside a test are test fixtures, not claimed client
artifacts; the real client gaps remain unsatisfied.
"""
from pathlib import Path
import json
import sys

import openpyxl
import pytest
from docx import Document

REPO_ROOT = Path(__file__).resolve().parents[2]
for _p in (str(REPO_ROOT), str(REPO_ROOT / "foundation")):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from foundation.applications.rollforward.region_model import (
    APPENDIX_LEVEL_OFFSET,
    HEADING_STYLE_RE,
    TOC_STYLE_RE,
    CanonicalRegionModelBuilder,
    SubregionKind,
    region_count_reconciliation,
)
from foundation.applications.rollforward.source_intake import (
    FORBIDDEN_READINESS,
    NON_SUPPLYING_SCOPES,
    SUPPLYING_SCOPES,
    ArtifactFormat,
    ArtifactStatus,
    DatasetRole,
    EvidenceQuality,
    GroundTruthGuard,
    PackageStatus,
    Priority,
    ReadinessRecalculator,
    RequestStatus,
    RollForwardSourcePackage,
    SourceArtifactRequest,
    SourceIntakeError,
    SourceIntakeProfiler,
    SourceRequestRegister,
    SourceScope,
    compute_file_hash,
)
from foundation.tests.evaluation.rollforward_clean_planner_c2 import (
    PATH_APP1,
    PATH_FARPT,
    PATH_GROUND_TRUTH_FORBIDDEN,
    PATH_HIST,
    PATH_TMPL,
)

BENCHMARK_ROLES = (DatasetRole.COMPARABLE_COMPANIES, DatasetRole.SCREENING_RESULTS,
                   DatasetRole.IQR_RESULTS, DatasetRole.INDEPENDENCE_CODES)


# ============================================================================
# FIXTURES
# ============================================================================

@pytest.fixture(autouse=True)
def _quarantine():
    GroundTruthGuard.reset()
    GroundTruthGuard.register([PATH_GROUND_TRUTH_FORBIDDEN])
    yield
    GroundTruthGuard.reset()


def _workbook(path: Path, sheets):
    wb = openpyxl.Workbook()
    first = True
    for name, rows in sheets:
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        first = False
        for r in rows:
            ws.append(list(r))
    wb.save(str(path))
    wb.close()
    return path


@pytest.fixture
def comparables_workbook(tmp_path):
    """SYNTHETIC test fixture: a genuine comparable-company dataset."""
    rows = [["No", "Company name", "Country", "Ticker", "Tax code", "Business description"]]
    rows += [[i, f"PEER {i} JSC", "Vietnam", f"TK{i}", f"010000000{i}", "Garment manufacturer"]
             for i in range(1, 11)]
    screening = [["Screening criteria", "Eliminated", "Retained"],
                 ["Unavailability of financial data", 245, 195],
                 ["Making loss for three consecutive years", 40, 155]]
    iqr = [["Statistic", "Value"], ["25th percentile", 3.1], ["Median", 4.5],
           ["75th percentile", 6.2], ["Interquartile range", "3.1-6.2"]]
    return _workbook(tmp_path / "client_supplied_2024.xlsx",
                     [("A", rows), ("B", screening), ("C", iqr)])


@pytest.fixture
def real_package():
    pkg = RollForwardSourcePackage(package_id="TEST-PKG")
    for path, scope in ((PATH_HIST, SourceScope.HISTORICAL),
                        (PATH_TMPL, SourceScope.TEMPLATE),
                        (PATH_FARPT, SourceScope.CURRENT_FINANCIAL),
                        (PATH_APP1, SourceScope.CURRENT_TAX)):
        pkg.add(SourceIntakeProfiler.ingest(path, scope))
    return pkg


# ============================================================================
# HASHING & DUPLICATES
# ============================================================================

def test_artifact_is_hash_addressed(tmp_path):
    p = _workbook(tmp_path / "a.xlsx", [("S", [["Item", "Value"], ["Rent", 1]])])
    a = SourceIntakeProfiler.ingest(p)
    assert a.file_hash == compute_file_hash(p)
    assert len(a.file_hash) == 64
    assert a.artifact_id.startswith("src-")
    assert a.file_size == p.stat().st_size
    assert a.is_fresh(p) is True


def test_duplicate_content_is_refused_regardless_of_filename(tmp_path):
    p1 = _workbook(tmp_path / "first.xlsx", [("S", [["Item", "Value"], ["Rent", 1]])])
    p2 = tmp_path / "different_name.xlsx"
    p2.write_bytes(p1.read_bytes())

    pkg = RollForwardSourcePackage(package_id="dup")
    pkg.add(SourceIntakeProfiler.ingest(p1))
    with pytest.raises(SourceIntakeError, match="Duplicate artifact"):
        pkg.add(SourceIntakeProfiler.ingest(p2))
    assert len(pkg.artifacts) == 1


def test_different_content_with_same_filename_is_accepted(tmp_path):
    d1, d2 = tmp_path / "one", tmp_path / "two"
    d1.mkdir(); d2.mkdir()
    a = _workbook(d1 / "same.xlsx", [("S", [["Item", "Value"], ["Rent", 1]])])
    b = _workbook(d2 / "same.xlsx", [("S", [["Item", "Value"], ["Rent", 2]])])
    pkg = RollForwardSourcePackage(package_id="same-name")
    pkg.add(SourceIntakeProfiler.ingest(a))
    pkg.add(SourceIntakeProfiler.ingest(b))
    assert len(pkg.artifacts) == 2


# ============================================================================
# PACKAGE VERSIONING
# ============================================================================

def test_package_versions_on_every_addition(tmp_path):
    pkg = RollForwardSourcePackage(package_id="ver")
    assert pkg.version == 1 and pkg.parent_version is None
    for i in range(3):
        p = _workbook(tmp_path / f"w{i}.xlsx", [("S", [["Item", "Value"], [f"row{i}", i]])])
        pkg.add(SourceIntakeProfiler.ingest(p))
    assert pkg.version == 4
    assert pkg.parent_version == 3
    assert [h["action"] for h in pkg.history] == ["ADD_ARTIFACT"] * 3


def test_package_hash_is_content_identity(tmp_path, real_package):
    before = real_package.package_hash()
    assert before == real_package.package_hash()
    p = _workbook(tmp_path / "extra.xlsx", [("S", [["Item", "Value"], ["x", 1]])])
    real_package.add(SourceIntakeProfiler.ingest(p))
    assert real_package.package_hash() != before


def test_freezing_and_stale_package_cannot_be_frozen(tmp_path):
    p = _workbook(tmp_path / "f.xlsx", [("S", [["Item", "Value"], ["x", 1]])])
    pkg = RollForwardSourcePackage(package_id="frz")
    a = SourceIntakeProfiler.ingest(p)
    pkg.add(a)
    pkg.freeze()
    assert pkg.status == PackageStatus.FROZEN

    p.write_bytes(p.read_bytes() + b"\x00")
    fresh, reasons = pkg.verify_freshness({a.artifact_id: p})
    assert fresh is False and reasons
    assert pkg.status == PackageStatus.STALE
    with pytest.raises(SourceIntakeError, match="STALE package cannot be frozen"):
        pkg.freeze()


# ============================================================================
# STALE INPUT
# ============================================================================

def test_stale_input_detected_on_byte_change(tmp_path):
    p = _workbook(tmp_path / "s.xlsx", [("S", [["Item", "Value"], ["x", 1]])])
    a = SourceIntakeProfiler.ingest(p)
    pkg = RollForwardSourcePackage(package_id="stale")
    pkg.add(a)

    fresh, reasons = pkg.verify_freshness({a.artifact_id: p})
    assert fresh is True and reasons == []

    p.write_bytes(p.read_bytes() + b" ")
    fresh, reasons = pkg.verify_freshness({a.artifact_id: p})
    assert fresh is False
    assert a.status == ArtifactStatus.STALE_INPUT
    assert "content changed" in reasons[0]
    assert a.is_fresh(p) is False


def test_missing_file_is_stale(tmp_path):
    p = _workbook(tmp_path / "gone.xlsx", [("S", [["Item", "Value"], ["x", 1]])])
    a = SourceIntakeProfiler.ingest(p)
    pkg = RollForwardSourcePackage(package_id="gone")
    pkg.add(a)
    p.unlink()
    fresh, reasons = pkg.verify_freshness({a.artifact_id: p})
    assert fresh is False
    assert "no longer exists" in reasons[0]


# ============================================================================
# CONTENT-DERIVED ROLES
# ============================================================================

def test_benchmarking_dataset_is_detected_from_content(comparables_workbook):
    a = SourceIntakeProfiler.ingest(comparables_workbook)
    assert a.status == ArtifactStatus.PROFILED
    assert a.supplies(DatasetRole.COMPARABLE_COMPANIES)
    assert a.supplies(DatasetRole.SCREENING_RESULTS)
    assert a.supplies(DatasetRole.IQR_RESULTS)
    assert a.quality_for(DatasetRole.COMPARABLE_COMPANIES) == EvidenceQuality.VERIFIED


def test_filename_alone_never_creates_a_role(tmp_path):
    """`benchmark.xlsx` with no benchmarking content earns nothing."""
    p = _workbook(tmp_path / "benchmark.xlsx",
                  [("S", [["Item", "Value"], ["Rent", 100], ["Utilities", 50]])])
    a = SourceIntakeProfiler.ingest(p)
    for role in BENCHMARK_ROLES + (DatasetRole.BENCHMARKING_DATA,):
        assert not a.supplies(role), f"{role.value} credited to a file with no such content"
    assert a.dataset_roles == [DatasetRole.UNKNOWN]


def test_renaming_a_file_does_not_change_its_roles(tmp_path, comparables_workbook):
    renamed = tmp_path / "quarterly_expenses.xlsx"
    renamed.write_bytes(comparables_workbook.read_bytes())
    a = SourceIntakeProfiler.ingest(comparables_workbook)
    b = SourceIntakeProfiler.ingest(renamed)
    assert sorted(r.value for r in a.dataset_roles) == sorted(r.value for r in b.dataset_roles)
    assert a.file_hash == b.file_hash


def test_narrative_source_detection(tmp_path):
    doc = Document()
    doc.add_paragraph("Functions performed by the company during the year under review.")
    doc.add_paragraph("Assets used include manufacturing plant and equipment.")
    doc.add_paragraph("Risks assumed include market risk and credit risk.")
    doc.add_paragraph("The functional analysis below reflects current-year operations.")
    p = tmp_path / "client_mi_2024.docx"
    doc.save(str(p))

    a = SourceIntakeProfiler.ingest(p, SourceScope.ADDITIONAL)
    assert a.format == ArtifactFormat.DOCX
    assert a.supplies(DatasetRole.FAR)
    # Phase F.1: FAR is declared NARRATIVE/ANY-shaped by the canonical policy,
    # because functional-analysis evidence legitimately arrives as prose. This
    # fixture is a purpose-written management-information document carrying all
    # three mandatory fields plus the FAR discriminator, so VERIFIED is correct.
    # Structured roles still cannot be verified from prose -- see
    # test_unstructured_prose_cannot_verify_a_structured_role.
    assert a.quality_for(DatasetRole.FAR) == EvidenceQuality.VERIFIED


def test_unstructured_prose_cannot_verify_a_structured_role(tmp_path):
    doc = Document()
    doc.add_paragraph("We discuss the comparable companies, ticker symbols and tax codes "
                      "of the selected set, with quartile and median statistics.")
    p = tmp_path / "narrative_about_benchmarking.docx"
    doc.save(str(p))
    a = SourceIntakeProfiler.ingest(p, SourceScope.ADDITIONAL)
    for role in (DatasetRole.COMPARABLE_COMPANIES, DatasetRole.IQR_RESULTS):
        if a.supplies(role):
            assert a.quality_for(role) == EvidenceQuality.INFERRED, (
                f"{role.value} must never be VERIFIED from prose alone")


def test_unsupported_format_is_rejected_without_inference(tmp_path):
    p = tmp_path / "notes.txt"
    p.write_text("comparable companies quartile median ticker tax code", encoding="utf-8")
    a = SourceIntakeProfiler.ingest(p)
    assert a.status == ArtifactStatus.REJECTED
    assert a.dataset_roles == [DatasetRole.UNKNOWN]
    assert "Nothing was inferred from the filename" in a.notes


def test_deterministic_repeated_profiling(comparables_workbook):
    a = SourceIntakeProfiler.ingest(comparables_workbook)
    b = SourceIntakeProfiler.ingest(comparables_workbook)
    assert a.file_hash == b.file_hash
    assert [r.value for r in a.dataset_roles] == [r.value for r in b.dataset_roles]
    assert [e.to_dict() for e in a.role_evidence] == [e.to_dict() for e in b.role_evidence]
    assert a.profile_summary == b.profile_summary


# ============================================================================
# SCOPE RULE — a prior-year output is not a current-year source
# ============================================================================

def test_historical_and_template_supply_no_current_year_roles(real_package):
    for a in real_package.artifacts:
        if a.source_scope in (SourceScope.HISTORICAL, SourceScope.TEMPLATE):
            assert a.can_supply_current_year is False
            assert a.observed_roles(), "content roles are still recorded as evidence"
            for role in a.dataset_roles:
                assert not a.supplies(role)
                assert a.quality_for(role) == EvidenceQuality.UNKNOWN


def test_package_capability_excludes_non_supplying_scopes(real_package):
    available = {r.value for r in real_package.available_roles()}
    for role in BENCHMARK_ROLES + (DatasetRole.BENCHMARKING_DATA, DatasetRole.FAR):
        assert role.value not in available, (
            f"{role.value} became available; the only artifacts mentioning it are the prior-year "
            f"Local File and the template, which are not current-year sources")
    # What the real current sources genuinely do supply.
    assert DatasetRole.RELATED_PARTY_TRANSACTIONS.value in available
    assert DatasetRole.FINANCIAL_STATEMENTS.value in available


def test_supplying_scope_sets_are_disjoint_and_complete():
    assert SUPPLYING_SCOPES.isdisjoint(NON_SUPPLYING_SCOPES)
    assert SUPPLYING_SCOPES | NON_SUPPLYING_SCOPES == set(SourceScope)


# ============================================================================
# GROUND TRUTH QUARANTINE
# ============================================================================

def test_ground_truth_cannot_be_ingested():
    with pytest.raises(SourceIntakeError, match="evaluation-only Ground Truth"):
        SourceIntakeProfiler.ingest(PATH_GROUND_TRUTH_FORBIDDEN)


def test_ground_truth_is_refused_even_when_renamed(tmp_path):
    copy = tmp_path / "innocent_client_data.docx"
    copy.write_bytes(PATH_GROUND_TRUTH_FORBIDDEN.read_bytes())
    with pytest.raises(SourceIntakeError, match="same content hash"):
        SourceIntakeProfiler.ingest(copy)


def test_evaluation_only_artifact_cannot_join_a_package(tmp_path):
    p = _workbook(tmp_path / "eval.xlsx", [("S", [["Item", "Value"], ["x", 1]])])
    a = SourceIntakeProfiler.ingest(p)
    a.source_scope = SourceScope.EVALUATION_ONLY
    pkg = RollForwardSourcePackage(package_id="eo")
    with pytest.raises(SourceIntakeError, match="EVALUATION_ONLY"):
        pkg.add(a)


# ============================================================================
# SOURCE REQUESTS
# ============================================================================

def _benchmark_request():
    return SourceArtifactRequest(
        artifact_request_id="REQ-TEST",
        artifact_type="BENCHMARKING_REPORT",
        title="Benchmarking dataset",
        affected_regions=["rgn-046", "rgn-051"],
        required_dataset_roles=[DatasetRole.COMPARABLE_COMPANIES,
                                DatasetRole.SCREENING_RESULTS,
                                DatasetRole.IQR_RESULTS],
        required_fields=["comparable set", "screening criteria", "quartiles"],
        priority=Priority.BLOCKING, blocking=True)


def test_request_is_outstanding_against_the_real_package(real_package):
    req = _benchmark_request().evaluate(real_package)
    assert req.current_status == RequestStatus.OUTSTANDING
    assert req.satisfied_by == []
    assert set(req.missing_roles) == set(req.required_dataset_roles)


def test_request_becomes_satisfied_only_by_a_real_supplying_artifact(
        real_package, comparables_workbook):
    req = _benchmark_request().evaluate(real_package)
    assert req.current_status == RequestStatus.OUTSTANDING

    real_package.add(SourceIntakeProfiler.ingest(comparables_workbook, SourceScope.ADDITIONAL))
    req.evaluate(real_package)
    assert req.current_status == RequestStatus.SATISFIED
    assert req.satisfied_by
    assert req.missing_roles == []


def test_partial_satisfaction_is_reported_as_partial(real_package, tmp_path):
    rows = [["No", "Company name", "Country", "Ticker", "Tax code", "Business description"]]
    rows += [[i, f"PEER {i}", "Vietnam", f"T{i}", f"010000{i}", "Manufacturer"] for i in range(1, 6)]
    partial = _workbook(tmp_path / "only_comparables.xlsx", [("A", rows)])
    real_package.add(SourceIntakeProfiler.ingest(partial, SourceScope.ADDITIONAL))

    req = _benchmark_request().evaluate(real_package)
    assert req.current_status == RequestStatus.PARTIALLY_SATISFIED
    assert DatasetRole.COMPARABLE_COMPANIES not in req.missing_roles
    assert req.missing_roles


def test_a_non_supplying_scope_cannot_satisfy_a_request(real_package, comparables_workbook):
    """Even a genuine dataset cannot satisfy a request if filed as HISTORICAL."""
    real_package.add(SourceIntakeProfiler.ingest(comparables_workbook, SourceScope.HISTORICAL))
    req = _benchmark_request().evaluate(real_package)
    assert req.current_status == RequestStatus.OUTSTANDING


def test_register_reports_blocking_outstanding(real_package):
    reg = SourceRequestRegister(requests=[_benchmark_request()]).evaluate_all(real_package)
    assert len(reg.blocking_outstanding()) == 1
    payload = reg.to_dict()
    assert payload["status_counts"]["OUTSTANDING"] == 1
    assert payload["blocking_outstanding"] == 1


# ============================================================================
# READINESS RECALCULATION — PLANNING ONLY
# ============================================================================

REGIONS = [
    {"region_id": "r1", "readiness": "BLOCKED_MISSING_SOURCE",
     "required_source_roles": ["COMPARABLE_COMPANIES", "IQR_RESULTS"]},
    {"region_id": "r2", "readiness": "BLOCKED_MISSING_SOURCE",
     "required_source_roles": ["COMPARABLE_COMPANIES", "CONTRACTUAL_DATA"]},
    {"region_id": "r3", "readiness": "HUMAN_REVIEW_READY",
     "required_source_roles": ["RELATED_PARTY_TRANSACTIONS"]},
    {"region_id": "r4", "readiness": "NOT_APPLICABLE", "required_source_roles": []},
]


def test_readiness_recalculates_to_human_review_not_execution(real_package, comparables_workbook):
    real_package.add(SourceIntakeProfiler.ingest(comparables_workbook, SourceScope.ADDITIONAL))
    transitions, counts = ReadinessRecalculator.recalculate(REGIONS, real_package)

    moved = {t.region_id: t for t in transitions}
    assert "r1" in moved
    assert moved["r1"].recalculated == "HUMAN_REVIEW_READY"
    assert "NOT approved" in moved["r1"].reason
    # r2 still needs contractual data.
    assert "r2" not in moved
    for value in counts:
        assert value not in FORBIDDEN_READINESS


def test_recalculation_can_never_emit_an_execution_state(real_package, comparables_workbook):
    real_package.add(SourceIntakeProfiler.ingest(comparables_workbook, SourceScope.ADDITIONAL))
    _t, counts = ReadinessRecalculator.recalculate(REGIONS, real_package)
    assert set(counts) & set(FORBIDDEN_READINESS) == set()
    assert "AUTO_MUTATION_READY" in FORBIDDEN_READINESS
    assert "APPROVED" in FORBIDDEN_READINESS and "EXECUTING" in FORBIDDEN_READINESS


def test_governance_statement_denies_execution_authority():
    gov = ReadinessRecalculator.approval_still_required()
    assert gov["execution_authorized"] is False
    assert gov["requires_human_approval"] is True
    assert "APPROVED" in gov["forbidden_outputs"]


def test_dry_run_simulation_creates_nothing(tmp_path):
    before = set(tmp_path.iterdir())
    transitions, counts = ReadinessRecalculator.simulate(
        REGIONS, {"COMPARABLE_COMPANIES", "IQR_RESULTS"})
    assert any(t.region_id == "r1" for t in transitions)
    assert set(tmp_path.iterdir()) == before


def test_any_of_role_alias_is_honoured():
    regions = [{"region_id": "n1", "readiness": "BLOCKED_MISSING_SOURCE",
                "required_source_roles": ["NARRATIVE_DATA"]}]
    alias = {"NARRATIVE_DATA": ("FAR", "BUSINESS_NARRATIVE")}
    t_far, _ = ReadinessRecalculator.simulate(regions, {"FAR"}, role_alias=alias)
    t_bus, _ = ReadinessRecalculator.simulate(regions, {"BUSINESS_NARRATIVE"}, role_alias=alias)
    t_none, _ = ReadinessRecalculator.simulate(
        regions, {"RELATED_PARTY_TRANSACTIONS"}, role_alias=alias)
    assert len(t_far) == 1 and len(t_bus) == 1 and len(t_none) == 0


# ============================================================================
# REGION HIERARCHY (§2)
# ============================================================================

@pytest.fixture(scope="module")
def region_model():
    return CanonicalRegionModelBuilder.build(PATH_TMPL, element_count=848)


def test_canonical_hierarchy_is_an_exact_partition(region_model):
    assert region_model.is_exact_partition() is True
    counts = region_model.counts
    assert counts["DOCUMENT"] == 1
    assert counts["SECTION"] < counts["REGION"] < counts["ELEMENT"]


def test_appendix_headings_are_recognised_as_sections(region_model):
    hist = region_model.heading_style_histogram
    appendix = {k: v for k, v in hist.items() if k.lower().startswith("appendix")}
    assert appendix, "the template's Appendix Heading styles must open sections"
    assert sum(appendix.values()) == 17
    appendix_sections = [s for s in region_model.sections if s.is_appendix]
    assert len(appendix_sections) == 17
    for s in appendix_sections:
        assert s.level > APPENDIX_LEVEL_OFFSET


def test_table_of_contents_lines_never_open_a_section(region_model):
    assert region_model.toc_paragraphs_excluded > 0
    for s in region_model.sections:
        assert not TOC_STYLE_RE.match(s.heading_style or "")
    # A contents line looks like a heading; the style is what separates them.
    assert HEADING_STYLE_RE.match("Heading 1")
    assert HEADING_STYLE_RE.match("Appendix Heading 2")
    assert not HEADING_STYLE_RE.match("toc 1")
    assert TOC_STYLE_RE.match("toc 1") and TOC_STYLE_RE.match("toc 2")


def test_region_count_reconciliation_rejects_all_three_historical_claims(region_model):
    rec = region_count_reconciliation(region_model)
    assert rec["answer"].startswith("NO")
    values = {c["value"]: c for c in rec["claims"]}
    assert set(values) == {61, 81, 104}
    assert "UNDERCOUNT" in values[61]["verdict"]
    assert "UNSUPPORTED" in values[81]["verdict"]
    assert "OVERCOUNT" in values[104]["verdict"]
    assert rec["canonical_counts"]["SECTION"] == len(region_model.sections)


def test_the_61_gap_is_exactly_the_appendix_headings(region_model):
    """61 + 17 appendix headings = the canonical 78 sections."""
    rec = region_count_reconciliation(region_model)
    assert rec["main_outline_headings"] + 1 == 61
    assert rec["appendix_headings"] == 17
    assert rec["main_outline_headings"] + rec["appendix_headings"] + 1 == \
        rec["canonical_counts"]["SECTION"]


def test_every_region_belongs_to_a_declared_section(region_model):
    section_ids = {s.section_id for s in region_model.sections}
    for r in region_model.regions:
        assert r.section_id in section_ids
    declared = [rid for s in region_model.sections for rid in s.region_ids]
    assert sorted(declared) == sorted(r.region_id for r in region_model.regions)


def test_each_table_appears_in_exactly_one_region(region_model):
    ordinals = [r.table_ordinal for r in region_model.regions if r.table_ordinal is not None]
    assert sorted(ordinals) == list(range(region_model.body_table_count))
    subs = [s for s in region_model.subregions if s.kind == SubregionKind.TABLE]
    assert len(subs) == region_model.body_table_count


def test_region_model_is_deterministic():
    a = CanonicalRegionModelBuilder.build(PATH_TMPL, element_count=848)
    b = CanonicalRegionModelBuilder.build(PATH_TMPL, element_count=848)
    assert json.dumps(a.to_dict(), sort_keys=True) == json.dumps(b.to_dict(), sort_keys=True)


# ============================================================================
# EMITTED ARTIFACTS
# ============================================================================

def test_emitted_package_and_request_artifacts_are_consistent():
    pkg_path = REPO_ROOT / "docs/evaluation/LocalFile_RollForward_Source_Package_V1.json"
    req_path = REPO_ROOT / "docs/evaluation/LocalFile_RollForward_Source_Request_Matrix_V1.json"
    if not pkg_path.exists() or not req_path.exists():
        pytest.skip("Phase F artifacts not generated")

    pkg = json.loads(pkg_path.read_text(encoding="utf-8"))
    req = json.loads(req_path.read_text(encoding="utf-8"))

    assert pkg["artifact_count"] == len(pkg["artifacts"])
    assert pkg["canonical_region_model"]["counts"]["SECTION"] == 78
    assert pkg["canonical_region_model"]["is_exact_partition"] is True
    for role in (r.value for r in BENCHMARK_ROLES):
        assert role not in pkg["available_roles"]

    assert req["blocking_outstanding"] == req["status_counts"].get("OUTSTANDING", 0)
    assert req["governance"]["execution_authorized"] is False
    for d in req["rebinding_dry_runs"]:
        assert d["is_dry_run"] is True and d["artifact_created"] is False
        for t in d["transitions"]:
            assert t["recalculated"] not in FORBIDDEN_READINESS


def test_ground_truth_never_appears_in_emitted_artifacts():
    for name in ("LocalFile_RollForward_Source_Package_V1.json",
                 "LocalFile_RollForward_Source_Request_Matrix_V1.json"):
        p = REPO_ROOT / "docs/evaluation" / name
        if not p.exists():
            continue
        assert PATH_GROUND_TRUTH_FORBIDDEN.name not in p.read_text(encoding="utf-8")
