"""
Roll-Forward Source Intake & Evidence Contract (Phase F)
=========================================================
Location: foundation/applications/rollforward/source_intake.py

Lets additional client source artifacts be uploaded, profiled, hash-addressed
and folded into a versioned source package -- WITHOUT bypassing governance and
WITHOUT mutating any document.

The contract in one line
------------------------
    New evidence may change READINESS.
    New evidence must never silently trigger EXECUTION.

Concretely:

  * `SourceArtifact` records what a file IS, from its content. A file named
    `benchmark.xlsx` earns BENCHMARKING_DATA only if its cells say so; the
    filename contributes nothing to the role verdict.
  * `RollForwardSourcePackage` is versioned and hash-addressed. Any byte change
    to any member makes the package STALE and forces explicit re-analysis.
  * `SourceRequestRegister` carries the Phase E gaps as machine-readable
    requests, and marks a request SATISFIED only when a real ingested artifact
    supplies every required role.
  * `ReadinessRecalculator` recomputes readiness after intake. Its best possible
    outcome is HUMAN_REVIEW_READY. It cannot emit APPROVED or EXECUTING, and it
    performs no mutation.

Ground Truth is never a source: `GroundTruthGuard` refuses to ingest it.

This module reads files and computes verdicts. It writes no document.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timezone
from enum import Enum
import hashlib
import json
from pathlib import Path
import re
from typing import Any, Dict, List, Optional, Sequence, Set, Tuple

from applications.rollforward.source_capability import (
    DatasetRole as C2DatasetRole,
    SheetDatasetProfile,
    SourceCapabilityProfiler,
    WorkbookCapabilityProfile,
)


# ============================================================================
# 1. TAXONOMIES
# ============================================================================

class ArtifactFormat(str, Enum):
    XLSX = "XLSX"
    DOCX = "DOCX"
    PDF = "PDF"
    CSV = "CSV"
    UNSUPPORTED = "UNSUPPORTED"


class SourceScope(str, Enum):
    """What period / role in the workflow an artifact plays."""
    HISTORICAL = "HISTORICAL"                # prior-year Local File
    TEMPLATE = "TEMPLATE"                    # master template
    CURRENT_FINANCIAL = "CURRENT_FINANCIAL"  # FA&RPT-type workbook
    CURRENT_TAX = "CURRENT_TAX"              # Appendix / statutory workbook
    ADDITIONAL = "ADDITIONAL"                # newly uploaded supporting artifact
    EVALUATION_ONLY = "EVALUATION_ONLY"      # Ground Truth — never a source


class DatasetRole(str, Enum):
    """Dataset roles an artifact may supply. Always content-derived."""
    FINANCIAL_STATEMENTS = "FINANCIAL_STATEMENTS"
    FINANCIAL_ANALYSIS = "FINANCIAL_ANALYSIS"
    RPT = "RPT"
    TAX_SCHEDULE = "TAX_SCHEDULE"
    FIXED_ASSETS = "FIXED_ASSETS"
    INTEREST_EXPENSE = "INTEREST_EXPENSE"
    SEGMENTED_DATA = "SEGMENTED_DATA"
    APPENDIX_DISCLOSURE = "APPENDIX_DISCLOSURE"
    BENCHMARKING_DATA = "BENCHMARKING_DATA"
    COMPARABLE_COMPANIES = "COMPARABLE_COMPANIES"
    IQR_RESULTS = "IQR_RESULTS"
    SCREENING_RESULTS = "SCREENING_RESULTS"
    INDEPENDENCE_CODES = "INDEPENDENCE_CODES"
    FAR = "FAR"
    BUSINESS_NARRATIVE = "BUSINESS_NARRATIVE"
    GROUP_NARRATIVE = "GROUP_NARRATIVE"
    CONTRACTUAL_DATA = "CONTRACTUAL_DATA"
    ORGANIZATIONAL_DATA = "ORGANIZATIONAL_DATA"
    TAXPAYER_PROFILE = "TAXPAYER_PROFILE"
    OWNERSHIP_STRUCTURE = "OWNERSHIP_STRUCTURE"
    FIGURE_SOURCE = "FIGURE_SOURCE"
    UNKNOWN = "UNKNOWN"


class EvidenceQuality(str, Enum):
    """How strongly the artifact's content supports a claimed role."""
    VERIFIED = "VERIFIED"                    # structured records with the expected schema
    STRONGLY_SUPPORTED = "STRONGLY_SUPPORTED"
    INFERRED = "INFERRED"
    UNKNOWN = "UNKNOWN"


class ArtifactStatus(str, Enum):
    REGISTERED = "REGISTERED"
    PROFILED = "PROFILED"
    STALE_INPUT = "STALE_INPUT"
    REJECTED = "REJECTED"
    QUARANTINED = "QUARANTINED"              # Ground Truth or otherwise forbidden


class PackageStatus(str, Enum):
    DRAFT = "DRAFT"
    FROZEN = "FROZEN"
    STALE = "STALE"


class RequestStatus(str, Enum):
    OUTSTANDING = "OUTSTANDING"
    PARTIALLY_SATISFIED = "PARTIALLY_SATISFIED"
    SATISFIED = "SATISFIED"


class Priority(str, Enum):
    BLOCKING = "BLOCKING"
    HIGH = "HIGH"
    MEDIUM = "MEDIUM"
    LOW = "LOW"


EXTENSION_FORMAT = {
    ".xlsx": ArtifactFormat.XLSX, ".xlsm": ArtifactFormat.XLSX,
    ".docx": ArtifactFormat.DOCX, ".pdf": ArtifactFormat.PDF, ".csv": ArtifactFormat.CSV,
}

# Only these scopes may SUPPLY a current-year dataset role.
#
# A prior-year Local File and the master template are Local File *documents*:
# they discuss comparables, quartiles, functions and risks at length, so naive
# content matching credits them with every role in the taxonomy. They are the
# workflow's structural and baseline inputs, never its current-year data
# sources. Treating a prior-year output as a current-year source is exactly the
# contamination Phase D3.1 found; this makes it structurally impossible.
SUPPLYING_SCOPES = {
    SourceScope.CURRENT_FINANCIAL,
    SourceScope.CURRENT_TAX,
    SourceScope.ADDITIONAL,
}

NON_SUPPLYING_SCOPES = {
    SourceScope.HISTORICAL,
    SourceScope.TEMPLATE,
    SourceScope.EVALUATION_ONLY,
}

# Roles that only a genuine structured dataset can satisfy.
STRUCTURED_ROLES = {
    DatasetRole.COMPARABLE_COMPANIES, DatasetRole.IQR_RESULTS,
    DatasetRole.SCREENING_RESULTS, DatasetRole.INDEPENDENCE_CODES,
    DatasetRole.FINANCIAL_STATEMENTS, DatasetRole.RPT, DatasetRole.FIXED_ASSETS,
}


class SourceIntakeError(RuntimeError):
    """Raised when an artifact may not be ingested."""


# ============================================================================
# 2. CONTENT-DERIVED ROLE DETECTION
# ============================================================================

# Signals matched against CONTENT ONLY (headers, label columns, body text).
# The filename is never part of the haystack.
_ROLE_SIGNALS: Tuple[Tuple[DatasetRole, int, Tuple[str, ...]], ...] = (
    (DatasetRole.COMPARABLE_COMPANIES, 3,
     ("company name", "ticker", "tax code", "sic code", "naics", "business description",
      "province", "stock code", "comparable", "country of incorporation")),
    (DatasetRole.IQR_RESULTS, 2,
     ("quartile", "25th percentile", "35th percentile", "75th percentile", "median",
      "interquartile", "lower quartile", "upper quartile", "arm's length range")),
    (DatasetRole.SCREENING_RESULTS, 2,
     ("screening criteria", "eliminated", "retained", "reason for rejection", "rejected",
      "search criteria", "passed", "search strategy", "step")),
    (DatasetRole.INDEPENDENCE_CODES, 2,
     ("independence indicator", "bvd independence", "shareholder with more than",
      "independence code", "no shareholder")),
    (DatasetRole.BENCHMARKING_DATA, 2,
     ("tp catalyst", "orbis", "bureau van dijk", "benchmark", "comparable set", "peer set")),
    (DatasetRole.FAR, 2,
     ("functions performed", "assets used", "risks assumed", "functional analysis",
      "functions/assets/risks", "characterisation", "characterization")),
    (DatasetRole.GROUP_NARRATIVE, 2,
     ("group structure", "ultimate parent", "the group", "group overview", "master file")),
    (DatasetRole.BUSINESS_NARRATIVE, 2,
     ("business strategy", "business overview", "principal activity", "operations overview",
      "business restructuring", "market conditions")),
    (DatasetRole.CONTRACTUAL_DATA, 2,
     ("agreement", "this agreement", "effective date", "the parties", "term of the agreement",
      "signed", "clause")),
    (DatasetRole.ORGANIZATIONAL_DATA, 2,
     ("organisation chart", "organization chart", "reporting line", "headcount",
      "department", "general director", "reports to")),
    (DatasetRole.OWNERSHIP_STRUCTURE, 2,
     ("shareholder", "ownership", "% of ownership", "holding company", "parent company")),
    (DatasetRole.TAXPAYER_PROFILE, 2,
     ("tax code", "enterprise code", "registered address", "fiscal year", "principal activity")),
    (DatasetRole.TAX_SCHEDULE, 2,
     ("corporate income tax", "cit", "taxable income", "tax payable", "deductible expense")),
    (DatasetRole.FIXED_ASSETS, 2,
     ("fixed asset", "depreciation", "accumulated depreciation", "net book value",
      "historical cost")),
    (DatasetRole.RPT, 2,
     ("related party", "related-party", "intercompany", "intra-group", "amount (vnd)",
      "type of relationship", "transaction value")),
    (DatasetRole.INTEREST_EXPENSE, 2,
     ("interest rate", "credit institution", "principal", "maturity date", "loan")),
    (DatasetRole.FINANCIAL_STATEMENTS, 2,
     ("net sales", "revenue", "cost of goods sold", "gross profit", "total assets",
      "profit before tax", "balance sheet")),
    (DatasetRole.FINANCIAL_ANALYSIS, 2,
     ("net cost plus", "operating margin", "return on assets", "profit level indicator",
      "ebit", "margin")),
    (DatasetRole.SEGMENTED_DATA, 2, ("segment", "segmented", "allocation basis", "allocated")),
    (DatasetRole.APPENDIX_DISCLOSURE, 2,
     ("appendix", "form 01", "declaration", "point of reference", "disclosure")),
    (DatasetRole.FIGURE_SOURCE, 2,
     ("chart data", "figure data", "diagram source", "node", "edge", "series")),
)

# Map the C2 workbook profiler's roles onto the Phase F taxonomy so existing
# evidence is reused rather than re-derived.
_C2_TO_F: Dict[C2DatasetRole, DatasetRole] = {
    C2DatasetRole.RELATED_PARTY_TRANSACTIONS: DatasetRole.RPT,
    C2DatasetRole.FINANCIAL_STATEMENTS: DatasetRole.FINANCIAL_STATEMENTS,
    C2DatasetRole.FINANCIAL_ANALYSIS: DatasetRole.FINANCIAL_ANALYSIS,
    C2DatasetRole.FIXED_ASSETS: DatasetRole.FIXED_ASSETS,
    C2DatasetRole.BENCHMARKING_DATA: DatasetRole.BENCHMARKING_DATA,
    C2DatasetRole.COMPARABLE_COMPANIES: DatasetRole.COMPARABLE_COMPANIES,
    C2DatasetRole.IQR_RESULTS: DatasetRole.IQR_RESULTS,
    C2DatasetRole.SCREENING_RESULTS: DatasetRole.SCREENING_RESULTS,
    C2DatasetRole.INTEREST_EXPENSE: DatasetRole.INTEREST_EXPENSE,
    C2DatasetRole.SEGMENTED_DATA: DatasetRole.SEGMENTED_DATA,
    C2DatasetRole.REFERENCE_LIST: DatasetRole.APPENDIX_DISCLOSURE,
    C2DatasetRole.CHECKLIST: DatasetRole.APPENDIX_DISCLOSURE,
    C2DatasetRole.RELATED_PARTIES_REGISTER: DatasetRole.OWNERSHIP_STRUCTURE,
    C2DatasetRole.NARRATIVE_DISCLOSURE: DatasetRole.BUSINESS_NARRATIVE,
}


@dataclass
class RoleEvidence:
    """Why an artifact is credited with a dataset role."""
    role: DatasetRole
    quality: EvidenceQuality
    matched_tokens: List[str]
    record_count: int
    locations: List[str]
    rationale: str

    def to_dict(self) -> Dict[str, Any]:
        return {"role": self.role.value, "quality": self.quality.value,
                "matched_tokens": self.matched_tokens, "record_count": self.record_count,
                "locations": self.locations, "rationale": self.rationale}


# ============================================================================
# 3. SOURCE ARTIFACT
# ============================================================================

def compute_file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()


@dataclass
class SourceArtifact:
    """One hash-addressed source file and what its content proves it can supply."""
    artifact_id: str
    filename: str
    file_hash: str
    file_size: int
    format: ArtifactFormat
    source_scope: SourceScope
    dataset_roles: List[DatasetRole] = field(default_factory=list)
    role_evidence: List[RoleEvidence] = field(default_factory=list)
    created_at: str = field(default_factory=lambda: datetime.now(timezone.utc).isoformat())
    status: ArtifactStatus = ArtifactStatus.REGISTERED
    profile_summary: Dict[str, Any] = field(default_factory=dict)
    notes: str = ""

    @property
    def can_supply_current_year(self) -> bool:
        """A HISTORICAL or TEMPLATE artifact is a baseline, not a data source."""
        return self.source_scope in SUPPLYING_SCOPES

    def supplies(self, role: DatasetRole) -> bool:
        """True only if this artifact may supply a CURRENT-YEAR role."""
        return self.can_supply_current_year and role in self.dataset_roles

    def observed_roles(self) -> List[DatasetRole]:
        """Every role the content matched, regardless of whether it may supply."""
        return list(self.dataset_roles)

    def quality_for(self, role: DatasetRole) -> EvidenceQuality:
        if not self.can_supply_current_year:
            return EvidenceQuality.UNKNOWN
        for e in self.role_evidence:
            if e.role == role:
                return e.quality
        return EvidenceQuality.UNKNOWN

    def is_fresh(self, path: Path) -> bool:
        return path.exists() and compute_file_hash(path) == self.file_hash

    def to_dict(self) -> Dict[str, Any]:
        return {
            "artifact_id": self.artifact_id, "filename": self.filename,
            "file_hash": self.file_hash, "file_size": self.file_size,
            "format": self.format.value, "source_scope": self.source_scope.value,
            "dataset_roles": [r.value for r in self.dataset_roles],
            "role_evidence": [e.to_dict() for e in self.role_evidence],
            "created_at": self.created_at, "status": self.status.value,
            "can_supply_current_year_roles": self.can_supply_current_year,
            "profile_summary": self.profile_summary, "notes": self.notes,
        }


# ============================================================================
# 4. GROUND TRUTH GUARD
# ============================================================================

class GroundTruthGuard:
    """Refuses to ingest the evaluation-only oracle as a source."""

    _forbidden_hashes: Set[str] = set()
    _forbidden_names: Set[str] = set()

    @classmethod
    def register(cls, paths: Sequence[Path]) -> None:
        for p in paths:
            cls._forbidden_names.add(p.name.lower())
            if p.exists():
                cls._forbidden_hashes.add(compute_file_hash(p))

    @classmethod
    def reset(cls) -> None:
        cls._forbidden_hashes.clear()
        cls._forbidden_names.clear()

    @classmethod
    def check(cls, path: Path) -> None:
        if path.name.lower() in cls._forbidden_names:
            raise SourceIntakeError(
                f"'{path.name}' is the evaluation-only Ground Truth artifact and may not be "
                f"ingested as a source. It can satisfy no source request and create no "
                f"dataset role.")
        if path.exists() and compute_file_hash(path) in cls._forbidden_hashes:
            raise SourceIntakeError(
                f"'{path.name}' has the same content hash as the quarantined Ground Truth "
                f"artifact; ingestion refused regardless of filename.")


# ============================================================================
# 5. INTAKE PROFILER
# ============================================================================

class SourceIntakeProfiler:
    """Parses, profiles and grades an artifact from its content alone."""

    MAX_DOCX_PARAGRAPHS = 4000

    @classmethod
    def ingest(
        cls,
        path: Path,
        source_scope: SourceScope = SourceScope.ADDITIONAL,
        artifact_id: Optional[str] = None,
    ) -> SourceArtifact:
        GroundTruthGuard.check(path)
        if not path.exists():
            raise SourceIntakeError(f"Source artifact not found: {path}")

        fmt = EXTENSION_FORMAT.get(path.suffix.lower(), ArtifactFormat.UNSUPPORTED)
        artifact = SourceArtifact(
            artifact_id=artifact_id or f"src-{compute_file_hash(path)[:12]}",
            filename=path.name,
            file_hash=compute_file_hash(path),
            file_size=path.stat().st_size,
            format=fmt,
            source_scope=source_scope,
        )

        if fmt == ArtifactFormat.XLSX:
            roles, evidence, summary = cls._profile_xlsx(path)
        elif fmt == ArtifactFormat.DOCX:
            roles, evidence, summary = cls._profile_docx(path)
        elif fmt == ArtifactFormat.CSV:
            roles, evidence, summary = cls._profile_csv(path)
        else:
            artifact.status = ArtifactStatus.REJECTED
            artifact.notes = (
                f"Unsupported format '{path.suffix}'. Nothing was inferred from the filename.")
            artifact.dataset_roles = [DatasetRole.UNKNOWN]
            return artifact

        artifact.dataset_roles = roles or [DatasetRole.UNKNOWN]
        artifact.role_evidence = evidence
        artifact.profile_summary = summary
        artifact.status = ArtifactStatus.PROFILED
        return artifact

    # ------------------------------------------------------------------

    @classmethod
    def _grade(cls, role: DatasetRole, hits: Sequence[str], records: int,
               structured: bool) -> Tuple[EvidenceQuality, str]:
        """Grades evidence. A structured role needs structured records to be VERIFIED."""
        if role in STRUCTURED_ROLES and not structured:
            return (EvidenceQuality.INFERRED,
                    f"{len(hits)} content signal(s) found, but in unstructured text. "
                    f"{role.value} requires tabular records to be verified.")
        if structured and len(hits) >= 3 and records >= 2:
            return (EvidenceQuality.VERIFIED,
                    f"{len(hits)} schema signals over {records} structured records.")
        if len(hits) >= 3 or (structured and records >= 2):
            return (EvidenceQuality.STRONGLY_SUPPORTED,
                    f"{len(hits)} content signal(s) over {records} record(s).")
        return (EvidenceQuality.INFERRED,
                f"only {len(hits)} content signal(s); insufficient for a verified claim.")

    @classmethod
    def _detect(cls, haystack: str, records: int, structured: bool,
                locations: Sequence[str]) -> Tuple[List[DatasetRole], List[RoleEvidence]]:
        roles: List[DatasetRole] = []
        evidence: List[RoleEvidence] = []
        for role, threshold, tokens in _ROLE_SIGNALS:
            hits = [t for t in tokens if t in haystack]
            if len(hits) < threshold:
                continue
            quality, rationale = cls._grade(role, hits, records, structured)
            roles.append(role)
            evidence.append(RoleEvidence(
                role=role, quality=quality, matched_tokens=hits[:8],
                record_count=records, locations=list(locations)[:6], rationale=rationale))
        return roles, evidence

    @classmethod
    def _profile_xlsx(cls, path: Path) -> Tuple[List[DatasetRole], List[RoleEvidence], Dict[str, Any]]:
        wb: WorkbookCapabilityProfile = SourceCapabilityProfiler.profile_workbook(path)
        roles: List[DatasetRole] = []
        evidence: List[RoleEvidence] = []

        for sheet in wb.sheets:
            hay = " | ".join(h.lower() for h in sheet.headers)
            hay += " || " + " || ".join((c.header or "").lower() for c in sheet.record_schema)
            hay += " || " + " || ".join((c.sample or "").lower() for c in sheet.record_schema)
            structured = sheet.record_count >= 2 and bool(sheet.headers)
            found, ev = cls._detect(hay, sheet.record_count, structured, [sheet.sheet_name])
            for r, e in zip(found, ev):
                if r not in roles:
                    roles.append(r)
                    evidence.append(e)
                else:
                    existing = next(x for x in evidence if x.role == r)
                    if sheet.sheet_name not in existing.locations:
                        existing.locations.append(sheet.sheet_name)

            # Carry across whatever the C2 profiler already proved.
            for c2 in sheet.roles:
                mapped = _C2_TO_F.get(c2)
                if mapped and mapped not in roles:
                    roles.append(mapped)
                    evidence.append(RoleEvidence(
                        role=mapped, quality=EvidenceQuality.STRONGLY_SUPPORTED,
                        matched_tokens=sheet.role_evidence.get(c2.value, [])[:8],
                        record_count=sheet.record_count, locations=[sheet.sheet_name],
                        rationale=f"content profiler credited {c2.value} on this sheet."))

        summary = {
            "kind": "workbook", "sheets": wb.sheet_count,
            "sheet_names": [s.sheet_name for s in wb.sheets],
            "total_records": sum(s.record_count for s in wb.sheets),
            "formula_cells": sum(s.formula_cell_count for s in wb.sheets),
        }
        return roles, evidence, summary

    @classmethod
    def _profile_docx(cls, path: Path) -> Tuple[List[DatasetRole], List[RoleEvidence], Dict[str, Any]]:
        from docx import Document

        doc = Document(str(path))
        paragraphs = [re.sub(r"\s+", " ", p.text).strip()
                      for p in doc.paragraphs[:cls.MAX_DOCX_PARAGRAPHS]]
        paragraphs = [p for p in paragraphs if p]
        hay = " || ".join(p.lower() for p in paragraphs)

        table_headers: List[str] = []
        for t in doc.tables:
            if t.rows:
                table_headers.extend(c.text.strip().lower() for c in t.rows[0].cells)
        hay += " || " + " | ".join(table_headers)

        structured = bool(doc.tables) and any(len(t.rows) >= 3 for t in doc.tables)
        records = sum(max(len(t.rows) - 1, 0) for t in doc.tables)
        locations = [f"{len(doc.tables)} table(s)", f"{len(paragraphs)} paragraph(s)"]
        roles, evidence = cls._detect(hay, records, structured, locations)

        summary = {
            "kind": "document", "paragraphs": len(paragraphs), "tables": len(doc.tables),
            "table_rows": records,
            "has_structured_tables": structured,
        }
        return roles, evidence, summary

    @classmethod
    def _profile_csv(cls, path: Path) -> Tuple[List[DatasetRole], List[RoleEvidence], Dict[str, Any]]:
        import csv

        rows: List[List[str]] = []
        with open(path, newline="", encoding="utf-8", errors="replace") as f:
            for i, row in enumerate(csv.reader(f)):
                if i >= 500:
                    break
                rows.append(row)
        headers = [c.strip().lower() for c in rows[0]] if rows else []
        hay = " | ".join(headers)
        for r in rows[1:60]:
            hay += " || " + " || ".join(c.strip().lower() for c in r[:4])
        records = max(len(rows) - 1, 0)
        roles, evidence = cls._detect(hay, records, records >= 2 and bool(headers), ["csv"])
        return roles, evidence, {"kind": "csv", "rows": len(rows), "headers": headers[:16]}


# ============================================================================
# 6. VERSIONED SOURCE PACKAGE
# ============================================================================

@dataclass
class RollForwardSourcePackage:
    """A versioned, hash-addressed set of everything planning is allowed to read."""
    package_id: str
    version: int = 1
    parent_version: Optional[int] = None
    created_at: str = field(default_factory=lambda: datetime.now(timezone.utc).isoformat())
    status: PackageStatus = PackageStatus.DRAFT
    artifacts: List[SourceArtifact] = field(default_factory=list)
    history: List[Dict[str, Any]] = field(default_factory=list)

    # -- membership ---------------------------------------------------

    def add(self, artifact: SourceArtifact) -> "RollForwardSourcePackage":
        """Adds an artifact and bumps the version. Duplicates are refused."""
        if artifact.source_scope == SourceScope.EVALUATION_ONLY:
            raise SourceIntakeError(
                "An EVALUATION_ONLY artifact may never join a source package.")
        existing = self.by_hash(artifact.file_hash)
        if existing is not None:
            raise SourceIntakeError(
                f"Duplicate artifact: '{artifact.filename}' has the same content hash as "
                f"already-registered '{existing.filename}' ({existing.artifact_id}). "
                f"Content, not filename, decides identity.")
        self.parent_version = self.version
        self.version += 1
        self.artifacts.append(artifact)
        self.status = PackageStatus.DRAFT
        self.history.append({
            "version": self.version, "action": "ADD_ARTIFACT",
            "artifact_id": artifact.artifact_id, "filename": artifact.filename,
            "file_hash": artifact.file_hash,
            "dataset_roles": [r.value for r in artifact.dataset_roles],
            "at": datetime.now(timezone.utc).isoformat(),
        })
        return self

    def by_hash(self, file_hash: str) -> Optional[SourceArtifact]:
        return next((a for a in self.artifacts if a.file_hash == file_hash), None)

    def by_scope(self, scope: SourceScope) -> List[SourceArtifact]:
        return [a for a in self.artifacts if a.source_scope == scope]

    # -- capability ----------------------------------------------------

    def available_roles(self, minimum_quality: EvidenceQuality = EvidenceQuality.INFERRED) -> Set[DatasetRole]:
        order = {EvidenceQuality.UNKNOWN: 0, EvidenceQuality.INFERRED: 1,
                 EvidenceQuality.STRONGLY_SUPPORTED: 2, EvidenceQuality.VERIFIED: 3}
        floor = order[minimum_quality]
        out: Set[DatasetRole] = set()
        for a in self.artifacts:
            if a.status != ArtifactStatus.PROFILED or not a.can_supply_current_year:
                continue
            for e in a.role_evidence:
                if order[e.quality] >= floor:
                    out.add(e.role)
        out.discard(DatasetRole.UNKNOWN)
        return out

    def observed_roles_all_scopes(self) -> Dict[str, List[str]]:
        """What every artifact matched, including non-supplying ones (evidence only)."""
        return {a.filename: [r.value for r in a.dataset_roles] for a in self.artifacts}

    def suppliers_of(self, role: DatasetRole) -> List[SourceArtifact]:
        return [a for a in self.artifacts
                if a.status == ArtifactStatus.PROFILED and a.supplies(role)]

    # -- freshness -----------------------------------------------------

    def verify_freshness(self, resolver: Dict[str, Path]) -> Tuple[bool, List[str]]:
        """Re-hashes every member. Any drift makes the whole package STALE."""
        stale: List[str] = []
        for a in self.artifacts:
            path = resolver.get(a.artifact_id) or resolver.get(a.filename)
            if path is None:
                stale.append(f"{a.filename}: no path supplied for re-hashing")
                continue
            if not path.exists():
                a.status = ArtifactStatus.STALE_INPUT
                stale.append(f"{a.filename}: file no longer exists")
                continue
            current = compute_file_hash(path)
            if current != a.file_hash:
                a.status = ArtifactStatus.STALE_INPUT
                stale.append(
                    f"{a.filename}: content changed "
                    f"(registered {a.file_hash[:12]}…, now {current[:12]}…)")
        if stale:
            self.status = PackageStatus.STALE
            self.history.append({
                "version": self.version, "action": "STALE_INPUT_DETECTED",
                "reasons": stale, "at": datetime.now(timezone.utc).isoformat()})
        return (not stale), stale

    def freeze(self) -> "RollForwardSourcePackage":
        if self.status == PackageStatus.STALE:
            raise SourceIntakeError(
                "A STALE package cannot be frozen. Re-run analysis explicitly after "
                "re-registering the changed artifacts.")
        self.status = PackageStatus.FROZEN
        self.history.append({"version": self.version, "action": "FREEZE",
                             "at": datetime.now(timezone.utc).isoformat()})
        return self

    def package_hash(self) -> str:
        """Deterministic digest of the package's content identity."""
        h = hashlib.sha256()
        for a in sorted(self.artifacts, key=lambda x: x.file_hash):
            h.update(f"{a.file_hash}:{a.source_scope.value}\n".encode("utf-8"))
        return h.hexdigest()

    def to_dict(self) -> Dict[str, Any]:
        return {
            "package_id": self.package_id, "version": self.version,
            "parent_version": self.parent_version, "created_at": self.created_at,
            "status": self.status.value, "package_hash": self.package_hash(),
            "artifact_count": len(self.artifacts),
            "supplying_artifacts": [a.artifact_id for a in self.artifacts
                                    if a.can_supply_current_year],
            "non_supplying_artifacts": [a.artifact_id for a in self.artifacts
                                        if not a.can_supply_current_year],
            "observed_roles_all_scopes": self.observed_roles_all_scopes(),
            "available_roles": sorted(r.value for r in self.available_roles()),
            "available_roles_verified_only": sorted(
                r.value for r in self.available_roles(EvidenceQuality.VERIFIED)),
            "artifacts": [a.to_dict() for a in self.artifacts],
            "history": self.history,
        }


# ============================================================================
# 7. SOURCE REQUEST REGISTER
# ============================================================================

@dataclass
class SourceArtifactRequest:
    """A Phase E gap, expressed so a machine can tell when it is satisfied."""
    artifact_request_id: str
    artifact_type: str
    title: str
    affected_regions: List[str]
    required_dataset_roles: List[DatasetRole]
    required_fields: List[str]
    priority: Priority
    blocking: bool
    current_status: RequestStatus = RequestStatus.OUTSTANDING
    satisfied_by: List[str] = field(default_factory=list)
    missing_roles: List[DatasetRole] = field(default_factory=list)
    minimum_evidence_quality: EvidenceQuality = EvidenceQuality.STRONGLY_SUPPORTED

    def evaluate(self, package: RollForwardSourcePackage) -> "SourceArtifactRequest":
        """Recomputes satisfaction against the package. Never fabricates a match."""
        order = {EvidenceQuality.UNKNOWN: 0, EvidenceQuality.INFERRED: 1,
                 EvidenceQuality.STRONGLY_SUPPORTED: 2, EvidenceQuality.VERIFIED: 3}
        floor = order[self.minimum_evidence_quality]

        satisfied: List[str] = []
        missing: List[DatasetRole] = []
        for role in self.required_dataset_roles:
            suppliers = [a for a in package.suppliers_of(role)
                         if order[a.quality_for(role)] >= floor]
            if suppliers:
                for a in suppliers:
                    if a.artifact_id not in satisfied:
                        satisfied.append(a.artifact_id)
            else:
                missing.append(role)

        self.satisfied_by = satisfied
        self.missing_roles = missing
        if not missing:
            self.current_status = RequestStatus.SATISFIED
        elif satisfied:
            self.current_status = RequestStatus.PARTIALLY_SATISFIED
        else:
            self.current_status = RequestStatus.OUTSTANDING
        return self

    def to_dict(self) -> Dict[str, Any]:
        return {
            "artifact_request_id": self.artifact_request_id,
            "artifact_type": self.artifact_type, "title": self.title,
            "affected_regions": self.affected_regions,
            "required_dataset_roles": [r.value for r in self.required_dataset_roles],
            "required_fields": self.required_fields,
            "priority": self.priority.value, "blocking": self.blocking,
            "current_status": self.current_status.value,
            "satisfied_by": self.satisfied_by,
            "missing_dataset_roles": [r.value for r in self.missing_roles],
            "minimum_evidence_quality": self.minimum_evidence_quality.value,
        }


@dataclass
class SourceRequestRegister:
    """The machine-readable set of outstanding source requests."""
    requests: List[SourceArtifactRequest] = field(default_factory=list)

    def evaluate_all(self, package: RollForwardSourcePackage) -> "SourceRequestRegister":
        for r in self.requests:
            r.evaluate(package)
        return self

    def outstanding(self) -> List[SourceArtifactRequest]:
        return [r for r in self.requests if r.current_status != RequestStatus.SATISFIED]

    def blocking_outstanding(self) -> List[SourceArtifactRequest]:
        return [r for r in self.outstanding() if r.blocking]

    def to_dict(self) -> Dict[str, Any]:
        counts: Dict[str, int] = {}
        for r in self.requests:
            counts[r.current_status.value] = counts.get(r.current_status.value, 0) + 1
        return {
            "total_requests": len(self.requests),
            "status_counts": counts,
            "blocking_outstanding": len(self.blocking_outstanding()),
            "requests": [r.to_dict() for r in self.requests],
        }


# ============================================================================
# 8. READINESS RECALCULATION (PLANNING ONLY)
# ============================================================================

class RecalculatedReadiness(str, Enum):
    """The only readiness values intake may produce. No APPROVED, no EXECUTING."""
    AUTO_NOOP_READY = "AUTO_NOOP_READY"
    HUMAN_REVIEW_READY = "HUMAN_REVIEW_READY"
    BLOCKED_MISSING_SOURCE = "BLOCKED_MISSING_SOURCE"
    BLOCKED_INSUFFICIENT_EVIDENCE = "BLOCKED_INSUFFICIENT_EVIDENCE"
    NOT_APPLICABLE = "NOT_APPLICABLE"


# Readiness values this module is structurally forbidden to emit.
FORBIDDEN_READINESS = ("APPROVED", "EXECUTING", "AUTO_MUTATION_READY", "COMPLETED",
                       "FINAL_VALIDATED")


@dataclass
class ReadinessTransition:
    region_id: str
    previous: str
    recalculated: str
    reason: str
    newly_available_roles: List[str] = field(default_factory=list)

    def to_dict(self) -> Dict[str, Any]:
        return dict(self.__dict__)


class ReadinessRecalculator:
    """Recomputes readiness after intake. Planning only; it mutates nothing.

    The best outcome any region can reach here is HUMAN_REVIEW_READY. A region
    never becomes AUTO_MUTATION_READY through intake, because that would also
    require an implemented mutation strategy, which this phase does not add.
    """

    @classmethod
    def recalculate(
        cls,
        regions: Sequence[Dict[str, Any]],
        package: RollForwardSourcePackage,
        role_alias: Optional[Dict[str, Sequence[str]]] = None,
        minimum_quality: EvidenceQuality = EvidenceQuality.STRONGLY_SUPPORTED,
    ) -> Tuple[List[ReadinessTransition], Dict[str, int]]:
        """`regions` are Phase E requirement entries as dicts."""
        available = {r.value for r in package.available_roles(minimum_quality)}
        return cls.simulate(regions, available, role_alias=role_alias)

    @classmethod
    def simulate(
        cls,
        regions: Sequence[Dict[str, Any]],
        available_roles: Set[str],
        role_alias: Optional[Dict[str, Sequence[str]]] = None,
    ) -> Tuple[List[ReadinessTransition], Dict[str, int]]:
        """Recomputes readiness against an arbitrary set of available roles.

        Used both for a real package and for an explicitly-labelled dry run that
        asks "what would change if these roles arrived?". A dry run creates no
        file and adds nothing to any package -- it fabricates no source.

        `role_alias` maps a requirement's role name onto one or more roles that
        would satisfy it (any-of). An empty tuple means the requirement needs no
        current source at all.
        """
        alias = role_alias or {}
        transitions: List[ReadinessTransition] = []
        counts: Dict[str, int] = {}

        def satisfied(role: str) -> bool:
            candidates = alias.get(role, (role,))
            if not candidates:
                return True          # requires no current source
            return any(c in available_roles for c in candidates)

        for region in regions:
            previous = region.get("readiness", "")
            required = list(region.get("required_source_roles", []))
            met = [r for r in required if satisfied(r)]
            unmet = [r for r in required if not satisfied(r)]

            if previous == RecalculatedReadiness.NOT_APPLICABLE.value:
                recalculated, reason = previous, "structural container or empty region"
            elif previous.startswith("BLOCKED") and required and not unmet:
                recalculated = RecalculatedReadiness.HUMAN_REVIEW_READY.value
                reason = (
                    f"Evidence supplies every required role {required}. The region becomes "
                    f"eligible for human review; it is NOT approved and NOT executable.")
            elif previous.startswith("BLOCKED") and met:
                recalculated = previous
                reason = f"Partially supplied ({met} of {required}); still blocked on {unmet}."
            else:
                recalculated, reason = previous, "no change: no newly available required role"

            assert recalculated not in FORBIDDEN_READINESS, (
                f"intake attempted to emit forbidden readiness '{recalculated}'")

            counts[recalculated] = counts.get(recalculated, 0) + 1
            if recalculated != previous:
                transitions.append(ReadinessTransition(
                    region_id=region.get("region_id", "?"), previous=previous,
                    recalculated=recalculated, reason=reason, newly_available_roles=met))

        return transitions, counts

    @staticmethod
    def approval_still_required() -> Dict[str, Any]:
        """The governance statement intake always returns alongside a recalculation."""
        return {
            "execution_authorized": False,
            "requires_human_approval": True,
            "statement": (
                "Recalculated readiness changes what a human MAY review. It does not approve "
                "anything, does not transition any manifest to APPROVED or EXECUTING, and "
                "triggers no mutation. Human approval remains the only execution authorization."),
            "forbidden_outputs": list(FORBIDDEN_READINESS),
        }
