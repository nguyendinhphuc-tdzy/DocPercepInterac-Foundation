"""
Full Document Roll-Forward Orchestrator (Phase D3)
===================================================
Location: foundation/applications/rollforward/orchestrator.py

Coordinates the already-built deterministic Roll-Forward components into one
governed, transactional, full-document execution pipeline:

    Phase B profile -> Phase C source binding -> RollForwardManifest
        -> HUMAN APPROVAL -> MutationPlan
        -> D1 StructuralWritebackEngine
        -> D2 DataReconciliationEngine
        -> D3 FullDocumentValidator
        -> Ground Truth evaluation (evaluation-only)
        -> FINAL_VALIDATED publication

This module owns coordination, gating, transaction boundaries, lineage and
reporting. It owns no perception, no mutation mechanics, no reconciliation
mathematics, and no state-transition rules -- every one of those is delegated
to the existing component that already implements it.

Non-negotiable invariants enforced here:
  * The Agent never approves; only a human `user` actor can.
  * Approval is version-locked to the exact manifest version executed.
  * Source and template freshness are verified before any mutation.
  * BLOCKED / UNKNOWN / MANUAL_REVIEW regions are excluded, never resolved.
  * The original template is never mutated in place.
  * A partially-mutated artifact can never become FINAL_VALIDATED.
  * Ground Truth is compared only after execution and never feeds back into it.
"""
from __future__ import annotations

from dataclasses import dataclass, field as dc_field
from datetime import datetime, timezone
from enum import Enum
import hashlib
import json
import os
from pathlib import Path
import shutil
from typing import Any, Dict, List, Optional, Sequence, Set, Tuple
import uuid

from docx import Document
from pydantic import BaseModel, Field

from applications.rollforward.data_reconciliation import (
    DataReconciliationEngine,
    ManifestReconciliationSummary,
    ReconciliationStatus,
    SourceFreshnessTracker,
)
from applications.rollforward.full_validation import (
    DocumentBaseline,
    ExpectedMutation,
    FullDocumentValidationReport,
    FullDocumentValidator,
    table_structure_signature,
    compute_file_sha256,
)
from applications.rollforward.models import (
    ExecutionGate,
    ManifestStatus,
    RegionClassification,
    RollForwardManifest,
    RollForwardRegion,
    SourceBindingStatus,
)
from applications.rollforward.state_machine import (
    RollForwardStateError,
    RollForwardStateMachine,
)
from applications.rollforward.structural_writeback import (
    ExecutionOutcome,
    MutationExecutionResult,
    MutationPlan,
    StructuralWritebackEngine,
    TableMutationSpec,
)

# Actor identifiers that can never stand in for a human approver.
RESERVED_NON_HUMAN_APPROVERS = {
    "", "system", "agent", "claude", "assistant", "foundation",
    "system-automated", "agent-automated", "auto", "none", "null",
}

# Reconciliation outcomes that are hard, unrecoverable D3 failures (Phase D3 §12).
HARD_RECONCILIATION_FAILURES = {
    ReconciliationStatus.MISMATCH,
    ReconciliationStatus.TYPE_MISMATCH,
    ReconciliationStatus.MISSING_SOURCE,
    ReconciliationStatus.MISSING_OUTPUT,
    ReconciliationStatus.TRANSFORMATION_MISMATCH,
    ReconciliationStatus.STALE_INPUT,
}

# Reconciliation outcomes that block publication but are not corruption.
MANUAL_REVIEW_RECONCILIATION = {
    ReconciliationStatus.MANUAL_REVIEW,
    ReconciliationStatus.FORMAT_MISMATCH,
    ReconciliationStatus.BLOCKED,
}


# ============================================================================
# 1. EXECUTION OUTCOME TAXONOMY
# ============================================================================

class ExecutionStatus(str, Enum):
    """Terminal outcome of a full-document roll-forward execution."""
    COMPLETED = "COMPLETED"
    FAILED = "FAILED"
    REQUIRES_MANUAL_REVIEW = "REQUIRES_MANUAL_REVIEW"
    NOOP = "NOOP"


class PublicationState(str, Enum):
    """Whether the generated artifact was published as the validated output."""
    FINAL_VALIDATED = "FINAL_VALIDATED"
    NOT_PUBLISHED = "NOT_PUBLISHED"


class FailureCode(str, Enum):
    """Explicit, structured reason an execution did not complete."""
    APPROVAL_MISSING = "APPROVAL_MISSING"
    APPROVAL_NOT_BY_USER = "APPROVAL_NOT_BY_USER"
    APPROVAL_VERSION_MISMATCH = "APPROVAL_VERSION_MISMATCH"
    PLAN_VERSION_MISMATCH = "PLAN_VERSION_MISMATCH"
    STALE_INPUT = "STALE_INPUT"
    STALE_TEMPLATE = "STALE_TEMPLATE"
    NO_EXECUTABLE_REGIONS = "NO_EXECUTABLE_REGIONS"
    TEMPLATE_NOT_FOUND = "TEMPLATE_NOT_FOUND"
    STRUCTURAL_MUTATION_FAILED = "STRUCTURAL_MUTATION_FAILED"
    DATA_RECONCILIATION_FAILED = "DATA_RECONCILIATION_FAILED"
    DATA_RECONCILIATION_MANUAL_REVIEW = "DATA_RECONCILIATION_MANUAL_REVIEW"
    FULL_VALIDATION_FAILED = "FULL_VALIDATION_FAILED"
    STATE_TRANSITION_REJECTED = "STATE_TRANSITION_REJECTED"
    INTERNAL_EXECUTION_ERROR = "INTERNAL_EXECUTION_ERROR"


class RegionExecutionStatus(str, Enum):
    """Per-region disposition within one execution."""
    EXECUTED = "EXECUTED"
    EXCLUDED = "EXCLUDED"


class ExclusionReason(str, Enum):
    """Why a region was not executed. Blocked regions stay blocked."""
    NOT_IN_APPROVED_PLAN = "NOT_IN_APPROVED_PLAN"
    REGION_NOT_IN_MANIFEST = "REGION_NOT_IN_MANIFEST"
    EXECUTION_GATE_BLOCKED = "EXECUTION_GATE_BLOCKED"
    CLASSIFICATION_UNKNOWN = "CLASSIFICATION_UNKNOWN"
    CLASSIFICATION_MANUAL_REVIEW = "CLASSIFICATION_MANUAL_REVIEW"
    SOURCE_BINDING_AMBIGUOUS = "SOURCE_BINDING_AMBIGUOUS"
    SOURCE_BINDING_STALE = "SOURCE_BINDING_STALE"
    SOURCE_BINDING_MISSING = "SOURCE_BINDING_MISSING"
    SOURCE_BINDING_UNVERIFIED = "SOURCE_BINDING_UNVERIFIED"
    REVIEW_STATE_REJECTED = "REVIEW_STATE_REJECTED"
    PLAN_TARGET_MISMATCH = "PLAN_TARGET_MISMATCH"
    UNSUPPORTED_OPERATION = "UNSUPPORTED_OPERATION"


class GroundTruthFindingStatus(str, Enum):
    """Honest, four-way evidence grading against the Ground Truth oracle."""
    VERIFIED = "VERIFIED"
    STRONGLY_SUPPORTED = "STRONGLY_SUPPORTED"
    INFERRED = "INFERRED"
    CONTRADICTED = "CONTRADICTED"


# ============================================================================
# 2. EXECUTION RECORDS & REPORT
# ============================================================================

class RegionExecutionRecord(BaseModel):
    """Disposition of one manifest region within one execution."""
    region_id: str
    section_name: str = ""
    classification: RegionClassification = RegionClassification.UNKNOWN
    execution_gate: ExecutionGate = ExecutionGate.BLOCKED
    status: RegionExecutionStatus = RegionExecutionStatus.EXCLUDED
    table_index: Optional[int] = None
    exclusion_reason: Optional[ExclusionReason] = None
    reason_detail: str = ""
    review_state: str = "PENDING"


class StructuralChangeRecord(BaseModel):
    """Structural delta actually applied to one target table."""
    region_id: str
    table_index: int
    operation: str
    rows_before: int
    rows_after: int
    rows_inserted: int
    columns_before: int
    columns_after: int
    precondition_fingerprint: str = ""
    postcondition_fingerprint: str = ""


class ReconciliationDigest(BaseModel):
    """Value-free digest of the D2 reconciliation summary (privacy policy §16)."""
    overall_status: str
    total_tables: int = 0
    total_cells: int = 0
    matched_cells: int = 0
    mismatched_cells: int = 0
    missing_cells: int = 0
    type_mismatches: int = 0
    format_mismatches: int = 0
    manual_review_items: int = 0
    blocked_items: int = 0
    source_freshness_verified: bool = False
    per_table: List[Dict[str, Any]] = Field(default_factory=list)


class GroundTruthFinding(BaseModel):
    """One evaluation-only observation against the Ground Truth oracle."""
    subject: str
    status: GroundTruthFindingStatus
    detail: str
    evidence: Dict[str, Any] = Field(default_factory=dict)


class GroundTruthEvaluationReport(BaseModel):
    """Evaluation-only Ground Truth comparison. Never alters execution state."""
    evaluated: bool = False
    ground_truth_document: str = ""
    ground_truth_sha256: str = ""
    findings: List[GroundTruthFinding] = Field(default_factory=list)
    note: str = (
        "Evaluation-only. Ground Truth was read after execution completed and was "
        "never used to generate values, copy structures, or modify the generated output."
    )

    def counts(self) -> Dict[str, int]:
        out = {s.value: 0 for s in GroundTruthFindingStatus}
        for f in self.findings:
            out[f.status.value] += 1
        return out


class ExecutionLineage(BaseModel):
    """Execution-level lineage chain, redacted of document plaintext."""
    schema_version: str = "1.0.0"
    nodes: List[Dict[str, Any]] = Field(default_factory=list)
    edges: List[Dict[str, Any]] = Field(default_factory=list)

    def chain_kinds(self) -> List[str]:
        return [n["type"] for n in self.nodes]


class RollForwardExecutionReport(BaseModel):
    """Complete, immutable record of one full-document roll-forward execution."""
    execution_id: str
    manifest_id: str
    manifest_version: int
    approved_manifest_version: Optional[int] = None
    approver: Optional[str] = None
    approved_at: Optional[str] = None
    mutation_plan_id: str = ""
    mutation_plan_digest: str = ""

    template_document: str = ""
    template_hash: str = ""
    source_hashes: Dict[str, str] = Field(default_factory=dict)

    started_at: str = ""
    ended_at: str = ""

    executed_regions: List[RegionExecutionRecord] = Field(default_factory=list)
    excluded_regions: List[RegionExecutionRecord] = Field(default_factory=list)
    structural_changes: List[StructuralChangeRecord] = Field(default_factory=list)

    reconciliation: Optional[ReconciliationDigest] = None
    validation_summary: Dict[str, Any] = Field(default_factory=dict)
    ground_truth_evaluation: Optional[GroundTruthEvaluationReport] = None
    lineage: Optional[ExecutionLineage] = None

    rollback_occurred: bool = False
    staging_discarded: bool = False
    template_preserved: bool = True
    idempotent_noop: bool = False

    output_path: Optional[str] = None
    output_hash: Optional[str] = None

    status: ExecutionStatus = ExecutionStatus.FAILED
    publication_state: PublicationState = PublicationState.NOT_PUBLISHED
    failure_code: Optional[FailureCode] = None
    failure_detail: str = ""

    execution_manifest_status: ManifestStatus = ManifestStatus.APPROVED
    state_transitions: List[str] = Field(default_factory=list)
    unresolved_blocked_summary: Dict[str, int] = Field(default_factory=dict)

    def to_json_dict(self) -> Dict[str, Any]:
        """Serializable report. Contains no document plaintext cell values."""
        return json.loads(self.model_dump_json())


# ============================================================================
# 3. EXECUTION REQUEST
# ============================================================================

@dataclass
class ExecutionRequest:
    """Everything the orchestrator needs; nothing it is allowed to invent.

    A dataclass rather than a pydantic model on purpose: the manifest and the
    mutation plan are handed over by reference, exactly as the planning layer
    built them, with no re-validation, no coercion, and no silent rewriting of
    the governed artifact the user approved.
    """
    manifest: RollForwardManifest
    mutation_plan: MutationPlan
    template_path: Path
    output_path: Path
    source_paths: List[Path] = dc_field(default_factory=list)
    expected_source_hashes: Dict[str, str] = dc_field(default_factory=dict)
    expected_template_hash: str = ""
    expected_approver: Optional[str] = None
    ground_truth_path: Optional[Path] = None
    execution_id: str = dc_field(default_factory=lambda: f"exec-{uuid.uuid4().hex[:12]}")
    actor_id: str = "system-orchestrator"


# ============================================================================
# 4. APPROVED SCOPE RESOLUTION
# ============================================================================

class ApprovedScope(BaseModel):
    """The mutation subset that approval and gating actually authorize."""
    executable_specs: List[TableMutationSpec] = Field(default_factory=list)
    executed_records: List[RegionExecutionRecord] = Field(default_factory=list)
    excluded_records: List[RegionExecutionRecord] = Field(default_factory=list)

    @property
    def executable_region_ids(self) -> Set[str]:
        return {s.target_region_id for s in self.executable_specs}


class MutationScopeResolver:
    """Partitions manifest regions into executable and excluded sets.

    This resolver only *filters*. It never resolves, infers, re-classifies, or
    auto-approves a blocked region, and it never invents a table index: the
    table index must come from the approved MutationPlan and must agree with
    the region's own Phase B/C row-template anchor.
    """

    @classmethod
    def resolve(cls, manifest: RollForwardManifest, mutation_plan: MutationPlan) -> ApprovedScope:
        scope = ApprovedScope()
        regions_by_id: Dict[str, RollForwardRegion] = {r.region_id: r for r in manifest.regions}
        planned_region_ids = {s.target_region_id for s in mutation_plan.table_mutations}

        for spec in mutation_plan.table_mutations:
            region = regions_by_id.get(spec.target_region_id)
            if region is None:
                scope.excluded_records.append(
                    RegionExecutionRecord(
                        region_id=spec.target_region_id,
                        exclusion_reason=ExclusionReason.REGION_NOT_IN_MANIFEST,
                        table_index=spec.table_index,
                        reason_detail=(
                            "MutationPlan targets a region that does not exist in the approved "
                            "manifest; the plan may not introduce new targets."
                        ),
                    )
                )
                continue

            reason, detail = cls._blocking_reason(region, spec)
            record = RegionExecutionRecord(
                region_id=region.region_id,
                section_name=region.section_name,
                classification=region.classification,
                execution_gate=region.execution_gate,
                table_index=spec.table_index,
                review_state=region.review_state,
            )
            if reason is not None:
                record.status = RegionExecutionStatus.EXCLUDED
                record.exclusion_reason = reason
                record.reason_detail = detail
                scope.excluded_records.append(record)
                continue

            record.status = RegionExecutionStatus.EXECUTED
            record.reason_detail = "READY and covered by the approved MutationPlan."
            scope.executed_records.append(record)
            scope.executable_specs.append(spec)

        # Every manifest region outside the approved plan is reported, never executed.
        for region in manifest.regions:
            if region.region_id in planned_region_ids:
                continue
            reason, detail = cls._blocking_reason(region, None)
            scope.excluded_records.append(
                RegionExecutionRecord(
                    region_id=region.region_id,
                    section_name=region.section_name,
                    classification=region.classification,
                    execution_gate=region.execution_gate,
                    table_index=cls.table_index_from_region(region),
                    review_state=region.review_state,
                    status=RegionExecutionStatus.EXCLUDED,
                    exclusion_reason=reason or ExclusionReason.NOT_IN_APPROVED_PLAN,
                    reason_detail=detail
                    or "Region is not covered by the approved MutationPlan and was not executed.",
                )
            )

        return scope

    @classmethod
    def _blocking_reason(
        cls, region: RollForwardRegion, spec: Optional[TableMutationSpec]
    ) -> Tuple[Optional[ExclusionReason], str]:
        """Returns the first governance reason this region may not execute."""
        if region.classification == RegionClassification.UNKNOWN:
            return (
                ExclusionReason.CLASSIFICATION_UNKNOWN,
                "Region classification is UNKNOWN; strictly blocked from auto-execution.",
            )
        if region.classification == RegionClassification.MANUAL_REVIEW:
            return (
                ExclusionReason.CLASSIFICATION_MANUAL_REVIEW,
                "Region requires a human tax reviewer; D3 does not resolve manual-review regions.",
            )
        if region.review_state == "REJECTED":
            return (ExclusionReason.REVIEW_STATE_REJECTED, "Region was explicitly rejected during review.")

        for src in region.current_sources:
            if src.status == SourceBindingStatus.AMBIGUOUS:
                return (
                    ExclusionReason.SOURCE_BINDING_AMBIGUOUS,
                    f"Source binding '{src.source_doc_name}' is AMBIGUOUS.",
                )
            if src.status == SourceBindingStatus.STALE:
                return (
                    ExclusionReason.SOURCE_BINDING_STALE,
                    f"Source binding '{src.source_doc_name}' is STALE.",
                )
            if src.status == SourceBindingStatus.MISSING:
                return (
                    ExclusionReason.SOURCE_BINDING_MISSING,
                    f"Source binding '{src.source_doc_name}' is MISSING.",
                )
            if src.status == SourceBindingStatus.UNVERIFIED:
                return (
                    ExclusionReason.SOURCE_BINDING_UNVERIFIED,
                    f"Source binding '{src.source_doc_name}' was never verified.",
                )

        if region.execution_gate != ExecutionGate.READY:
            return (
                ExclusionReason.EXECUTION_GATE_BLOCKED,
                f"Region execution gate is {region.execution_gate.value}.",
            )

        if spec is not None:
            anchor_table = cls.table_index_from_region(region)
            if anchor_table is not None and anchor_table != spec.table_index:
                return (
                    ExclusionReason.PLAN_TARGET_MISMATCH,
                    f"MutationPlan targets table {spec.table_index} but the manifest region's "
                    f"row-template anchor resolves to table {anchor_table}.",
                )
            if spec.operation != "INSERT_ROWS":
                return (
                    ExclusionReason.UNSUPPORTED_OPERATION,
                    f"Operation '{spec.operation}' is not supported by the D1 writeback engine.",
                )
            if spec.target_row_count < spec.initial_row_count:
                return (
                    ExclusionReason.UNSUPPORTED_OPERATION,
                    f"Plan implies row deletion ({spec.initial_row_count} -> {spec.target_row_count}); "
                    "the D1 writeback engine supports row insertion only.",
                )

        return (None, "")

    @staticmethod
    def table_index_from_region(region: RollForwardRegion) -> Optional[int]:
        """Reads the table index out of the region's Phase B row-template anchor.

        Anchor format (perception/anchor_builder.py): ``table:<idx>:<hash>_row:<n>``.
        """
        if region.row_template is None:
            return None
        anchor = region.row_template.row_anchor or ""
        if not anchor.startswith("table:"):
            return None
        try:
            return int(anchor.split(":", 2)[1])
        except (IndexError, ValueError):
            return None


# ============================================================================
# 5. GROUND TRUTH EVALUATOR (EVALUATION-ONLY)
# ============================================================================

class GroundTruthEvaluator:
    """Compares a *finished* artifact against the oracle. Read-only, after the fact.

    Ground Truth never generates a value, never contributes a structure, and
    never changes execution state or publication. It exists to report honestly
    how close the deterministic pipeline landed.
    """

    @classmethod
    def evaluate(
        cls,
        generated_path: Path,
        ground_truth_path: Path,
        expected_mutations: Sequence[ExpectedMutation],
    ) -> GroundTruthEvaluationReport:
        report = GroundTruthEvaluationReport(
            evaluated=True,
            ground_truth_document=ground_truth_path.name,
            ground_truth_sha256=compute_file_sha256(ground_truth_path),
        )

        gen_doc = Document(str(generated_path))
        gt_doc = Document(str(ground_truth_path))

        gen_sigs = [table_structure_signature(i, t) for i, t in enumerate(gen_doc.tables)]
        gt_sigs = [table_structure_signature(i, t) for i, t in enumerate(gt_doc.tables)]
        gt_by_header: Dict[str, List[Any]] = {}
        for sig in gt_sigs:
            gt_by_header.setdefault(sig.header_signature, []).append(sig)
        gt_by_first_cell: Dict[str, List[Any]] = {}
        for i, t in enumerate(gt_doc.tables):
            first = t.rows[0].cells[0].text.strip().lower() if t.rows else ""
            gt_by_first_cell.setdefault(first, []).append(gt_sigs[i])

        report.findings.append(
            GroundTruthFinding(
                subject="document.table_count",
                status=(
                    GroundTruthFindingStatus.VERIFIED
                    if len(gen_sigs) == len(gt_sigs)
                    else GroundTruthFindingStatus.CONTRADICTED
                ),
                detail=(
                    f"Generated document has {len(gen_sigs)} tables; "
                    f"Ground Truth has {len(gt_sigs)}."
                ),
                evidence={"generated": len(gen_sigs), "ground_truth": len(gt_sigs)},
            )
        )

        for em in expected_mutations:
            if em.table_index >= len(gen_sigs):
                report.findings.append(
                    GroundTruthFinding(
                        subject=f"table[{em.table_index}].rows",
                        status=GroundTruthFindingStatus.INFERRED,
                        detail="Target table absent from the generated document; nothing to compare.",
                        evidence={"region_id": em.region_id},
                    )
                )
                continue

            gen_sig = gen_sigs[em.table_index]
            matches = gt_by_header.get(gen_sig.header_signature, [])
            if matches:
                gt_rows = [m.row_count for m in matches]
                status = (
                    GroundTruthFindingStatus.VERIFIED
                    if gen_sig.row_count in gt_rows
                    else GroundTruthFindingStatus.CONTRADICTED
                )
                report.findings.append(
                    GroundTruthFinding(
                        subject=f"table[{em.table_index}].rows",
                        status=status,
                        detail=(
                            f"Header-identical Ground Truth table(s) at index {[m.table_index for m in matches]} "
                            f"have {gt_rows} rows; generated table has {gen_sig.row_count}."
                        ),
                        evidence={
                            "region_id": em.region_id,
                            "generated_rows": gen_sig.row_count,
                            "ground_truth_rows": gt_rows,
                        },
                    )
                )
                continue

            first_cell = (
                gen_doc.tables[em.table_index].rows[0].cells[0].text.strip().lower()
                if gen_doc.tables[em.table_index].rows
                else ""
            )
            partial = gt_by_first_cell.get(first_cell, [])
            if partial:
                report.findings.append(
                    GroundTruthFinding(
                        subject=f"table[{em.table_index}].rows",
                        status=GroundTruthFindingStatus.STRONGLY_SUPPORTED,
                        detail=(
                            f"No header-identical Ground Truth table exists, but table(s) "
                            f"{[m.table_index for m in partial]} share the leading header cell; "
                            f"Ground Truth row counts {[m.row_count for m in partial]} vs "
                            f"generated {gen_sig.row_count}. The FY2024 Ground Truth column set "
                            "evolved relative to the FY20XX template."
                        ),
                        evidence={
                            "region_id": em.region_id,
                            "generated_rows": gen_sig.row_count,
                            "ground_truth_rows": [m.row_count for m in partial],
                        },
                    )
                )
            else:
                report.findings.append(
                    GroundTruthFinding(
                        subject=f"table[{em.table_index}].rows",
                        status=GroundTruthFindingStatus.INFERRED,
                        detail=(
                            "No Ground Truth table could be deterministically correlated to this "
                            "target table; no claim is made either way."
                        ),
                        evidence={"region_id": em.region_id, "generated_rows": gen_sig.row_count},
                    )
                )

        return report


# ============================================================================
# 6. LINEAGE (PRIVACY-REDACTED)
# ============================================================================

class LineagePrivacyPolicy:
    """Document plaintext never leaves the reconciliation layer (Phase D3 §16)."""

    @staticmethod
    def digest(value: Any) -> str:
        if value is None:
            return "sha256:none"
        return "sha256:" + hashlib.sha256(str(value).encode("utf-8")).hexdigest()[:16]


class ExecutionLineageBuilder:
    """Extends the D2 cell lineage into a full execution-level chain."""

    @classmethod
    def build(
        cls,
        report: RollForwardExecutionReport,
        reconciliation: Optional[ManifestReconciliationSummary],
    ) -> ExecutionLineage:
        lineage = ExecutionLineage()
        add_node = lineage.nodes.append
        add_edge = lineage.edges.append

        for name, sha in sorted(report.source_hashes.items()):
            src_id = f"source_document:{name}"
            add_node({
                "id": src_id, "type": "SOURCE_DOCUMENT",
                "metadata": {"document_name": name, "sha256": sha},
            })
            add_edge({"from": src_id, "to": "source_binding:approved", "relation": "BOUND_BY"})

        binding_id = "source_binding:approved"
        add_node({
            "id": binding_id, "type": "SOURCE_BINDING",
            "metadata": {
                "verified_bindings": len(report.executed_regions),
                "binding_engine": "Phase C DeterministicSourceBindingEngine",
            },
        })

        manifest_id = f"manifest:{report.manifest_id}:v{report.manifest_version}"
        add_node({
            "id": manifest_id, "type": "MANIFEST_VERSION",
            "metadata": {
                "manifest_id": report.manifest_id,
                "manifest_version": report.manifest_version,
                "approved_manifest_version": report.approved_manifest_version,
                "approver": report.approver,
                "approved_at": report.approved_at,
            },
        })
        add_edge({"from": binding_id, "to": manifest_id, "relation": "DECLARED_IN"})

        plan_id = f"mutation_plan:{report.mutation_plan_id}"
        add_node({
            "id": plan_id, "type": "MUTATION_PLAN",
            "metadata": {
                "plan_id": report.mutation_plan_id,
                "plan_digest": report.mutation_plan_digest,
                "table_mutations": len(report.structural_changes),
            },
        })
        add_edge({"from": manifest_id, "to": plan_id, "relation": "AUTHORIZES"})

        exec_id = f"execution:{report.execution_id}"
        add_node({
            "id": exec_id, "type": "EXECUTION",
            "metadata": {
                "execution_id": report.execution_id,
                "started_at": report.started_at,
                "template_hash": report.template_hash,
            },
        })
        add_edge({"from": plan_id, "to": exec_id, "relation": "EXECUTED_AS"})

        for change in report.structural_changes:
            region_node = f"target_region:{change.region_id}:table{change.table_index}"
            add_node({
                "id": region_node, "type": "TARGET_REGION",
                "metadata": {
                    "region_id": change.region_id,
                    "table_index": change.table_index,
                    "rows_before": change.rows_before,
                    "rows_after": change.rows_after,
                    "rows_inserted": change.rows_inserted,
                },
            })
            add_edge({"from": exec_id, "to": region_node, "relation": "MUTATED"})

        if reconciliation is not None:
            for t_sum in reconciliation.table_summaries:
                region_node = f"target_region:{t_sum.target_region_id}:table{t_sum.table_index}"
                for c_rec in t_sum.cell_records:
                    cell_node = (
                        f"target_cell:T{t_sum.table_index}"
                        f"R{c_rec.target.row_idx}C{c_rec.target.col_idx}"
                    )
                    add_node({
                        "id": cell_node, "type": "TARGET_CELL",
                        "metadata": {
                            "table_index": t_sum.table_index,
                            "row_idx": c_rec.target.row_idx,
                            "col_idx": c_rec.target.col_idx,
                            "value_digest": LineagePrivacyPolicy.digest(c_rec.output_display_value),
                            "semantic_match": c_rec.semantic_match,
                            "display_match": c_rec.display_match,
                            "status": c_rec.status.value,
                            "source_document": c_rec.source.document_name if c_rec.source else None,
                            "source_sheet": c_rec.source.sheet_name if c_rec.source else None,
                            "source_cell": c_rec.source.cell_address if c_rec.source else None,
                        },
                    })
                    add_edge({"from": region_node, "to": cell_node, "relation": "CONTAINS_CELL"})

        doc_node = f"generated_document:{report.output_hash or 'unpublished'}"
        add_node({
            "id": doc_node, "type": "GENERATED_DOCUMENT",
            "metadata": {
                "output_path": report.output_path,
                "output_hash": report.output_hash,
                "publication_state": report.publication_state.value,
            },
        })
        add_edge({"from": exec_id, "to": doc_node, "relation": "PRODUCED"})

        val_node = f"validation_result:{report.execution_id}"
        add_node({
            "id": val_node, "type": "VALIDATION_RESULT",
            "metadata": {
                "status": report.status.value,
                "validation_summary": report.validation_summary,
                "reconciliation_status": (
                    report.reconciliation.overall_status if report.reconciliation else None
                ),
            },
        })
        add_edge({"from": doc_node, "to": val_node, "relation": "VALIDATED_BY"})

        return lineage


# ============================================================================
# 7. IDEMPOTENCE LEDGER
# ============================================================================

class ExecutionLedger(BaseModel):
    """Sidecar record of the last successfully published execution."""
    execution_id: str
    manifest_id: str
    manifest_version: int
    mutation_plan_id: str
    mutation_plan_digest: str
    output_hash: str
    published_at: str
    status: ExecutionStatus = ExecutionStatus.COMPLETED

    @staticmethod
    def path_for(output_path: Path) -> Path:
        return output_path.with_name(output_path.name + ".d3ledger.json")

    @classmethod
    def load(cls, output_path: Path) -> Optional["ExecutionLedger"]:
        ledger_path = cls.path_for(output_path)
        if not ledger_path.exists():
            return None
        try:
            return cls.model_validate_json(ledger_path.read_text(encoding="utf-8"))
        except Exception:  # noqa: BLE001 - a corrupt ledger simply means "no prior execution"
            return None

    def save(self, output_path: Path) -> None:
        self.path_for(output_path).write_text(
            json.dumps(json.loads(self.model_dump_json()), indent=2), encoding="utf-8"
        )


# ============================================================================
# 8. THE ORCHESTRATOR
# ============================================================================

class RollForwardOrchestrator:
    """Governed, transactional, full-document roll-forward execution pipeline."""

    @classmethod
    def execute(cls, request: ExecutionRequest) -> RollForwardExecutionReport:
        """Runs the full approved workflow and returns an immutable report."""
        manifest = request.manifest
        plan = request.mutation_plan
        plan_digest = _mutation_plan_digest(plan)

        report = RollForwardExecutionReport(
            execution_id=request.execution_id,
            manifest_id=manifest.manifest_id,
            manifest_version=manifest.manifest_version,
            approved_manifest_version=manifest.approved_manifest_version,
            approver=manifest.approved_by,
            approved_at=manifest.approved_at,
            mutation_plan_id=plan.plan_id,
            mutation_plan_digest=plan_digest,
            template_document=request.template_path.name,
            started_at=_now(),
            output_path=str(request.output_path),
        )

        # ---- 1. APPROVAL GATING (§5) -----------------------------------
        gate_failure = cls._verify_approval(manifest, plan, request)
        if gate_failure is not None:
            return cls._fail(report, *gate_failure)

        # ---- 2. TEMPLATE FRESHNESS (§7) --------------------------------
        if not request.template_path.exists():
            return cls._fail(
                report, FailureCode.TEMPLATE_NOT_FOUND,
                f"Target template not found: {request.template_path}",
            )
        template_hash = compute_file_sha256(request.template_path)
        report.template_hash = template_hash
        if request.expected_template_hash and template_hash != request.expected_template_hash:
            return cls._fail(
                report, FailureCode.STALE_TEMPLATE,
                f"Target template changed since the manifest was created "
                f"(expected {request.expected_template_hash[:12]}..., got {template_hash[:12]}...). "
                "Re-run analysis/replanning explicitly; D3 will not silently rebuild the manifest.",
            )

        # ---- 3. SOURCE FRESHNESS (§6) ----------------------------------
        report.source_hashes = dict(request.expected_source_hashes)
        if request.source_paths:
            fresh, stale_reasons = SourceFreshnessTracker.verify_freshness(
                request.expected_source_hashes, request.source_paths
            )
            if not fresh:
                return cls._fail(
                    report, FailureCode.STALE_INPUT,
                    "Bound current-source artifacts changed since planning: "
                    + "; ".join(stale_reasons)
                    + ". Aborting before mutation; automatic re-planning is not permitted.",
                )

        # ---- 4. EXECUTABLE SCOPE (§8, §22) -----------------------------
        scope = MutationScopeResolver.resolve(manifest, plan)
        report.executed_regions = scope.executed_records
        report.excluded_regions = scope.excluded_records
        report.unresolved_blocked_summary = _summarize_exclusions(scope.excluded_records)

        if not scope.executable_specs:
            return cls._fail(
                report, FailureCode.NO_EXECUTABLE_REGIONS,
                "No region in the approved MutationPlan is both READY and plan-covered. "
                "Blocked regions remain blocked and were not executed.",
            )

        # ---- 5. IDEMPOTENCE (§21) --------------------------------------
        noop = cls._check_idempotence(request, plan_digest, scope, report)
        if noop is not None:
            return noop

        # ---- 6. SCOPED EXECUTION MANIFEST & STATE MACHINE (§4) ---------
        exec_manifest = cls._scoped_execution_manifest(manifest, scope)
        try:
            RollForwardStateMachine.start_execution(
                exec_manifest,
                reason=f"D3 orchestrated execution {request.execution_id} "
                       f"over {len(scope.executable_specs)} approved region(s).",
            )
        except RollForwardStateError as ex:
            return cls._fail(report, FailureCode.STATE_TRANSITION_REJECTED, str(ex))
        report.state_transitions.append(f"{ManifestStatus.APPROVED.value}->{ManifestStatus.EXECUTING.value}")
        report.execution_manifest_status = exec_manifest.status

        scoped_plan = plan.model_copy(deep=True)
        scoped_plan.table_mutations = list(scope.executable_specs)

        # ---- 7. TRANSACTIONAL STAGING (§9) -----------------------------
        request.output_path.parent.mkdir(parents=True, exist_ok=True)
        staging_dir = request.output_path.parent / f".d3_staging_{request.execution_id}"
        staging_dir.mkdir(parents=True, exist_ok=True)
        staged_output = staging_dir / request.output_path.name

        target_indices = [s.table_index for s in scope.executable_specs]
        baseline = DocumentBaseline.capture(request.template_path, target_indices)
        expected_mutations = [
            ExpectedMutation(
                region_id=s.target_region_id,
                table_index=s.table_index,
                baseline_row_count=(
                    baseline.table_signatures[s.table_index].row_count
                    if s.table_index in baseline.table_signatures
                    else s.initial_row_count
                ),
                expected_row_count=s.target_row_count,
                expected_inserted_rows=len(s.row_mutations),
                operation=s.operation,
            )
            for s in scope.executable_specs
        ]

        reconciliation_summary: Optional[ManifestReconciliationSummary] = None

        try:
            # ---- 8. D1 STRUCTURAL WRITEBACK ----------------------------
            mutation_result: MutationExecutionResult = StructuralWritebackEngine.execute(
                manifest=exec_manifest,
                mutation_plan=scoped_plan,
                doc_path=request.template_path,
                output_path=staged_output,
                actor_id=request.actor_id,
            )
            if not mutation_result.success or mutation_result.outcome != ExecutionOutcome.APPLIED:
                return cls._abort(
                    report, staging_dir, exec_manifest,
                    FailureCode.STRUCTURAL_MUTATION_FAILED,
                    f"D1 structural writeback returned {mutation_result.outcome.value}: "
                    f"{mutation_result.error_message or 'no detail'}",
                )

            report.structural_changes = cls._structural_changes(
                baseline, staged_output, scope.executable_specs
            )

            # ---- 9. D2 DATA RECONCILIATION (§12) -----------------------
            reconciliation_summary = DataReconciliationEngine.reconcile_document_output(
                manifest=exec_manifest,
                mutation_plan=scoped_plan,
                doc_output_path=staged_output,
                source_paths=request.source_paths,
                source_hashes=request.expected_source_hashes,
            )
            report.reconciliation = _digest_reconciliation(reconciliation_summary)

            if reconciliation_summary.overall_status in HARD_RECONCILIATION_FAILURES:
                return cls._abort(
                    report, staging_dir, exec_manifest,
                    FailureCode.DATA_RECONCILIATION_FAILED,
                    f"D2 reconciliation returned {reconciliation_summary.overall_status.value} "
                    f"({reconciliation_summary.mismatched_cells} mismatched, "
                    f"{reconciliation_summary.missing_cells} missing, "
                    f"{reconciliation_summary.type_mismatches} type mismatches).",
                )

            # ---- 10. FULL-DOCUMENT VALIDATION (§10, §11, §13, §14) -----
            validation: FullDocumentValidationReport = FullDocumentValidator.validate(
                baseline=baseline,
                output_path=staged_output,
                expected_mutations=expected_mutations,
            )
            report.validation_summary = validation.summary()

            if not validation.is_valid:
                return cls._abort(
                    report, staging_dir, exec_manifest,
                    FailureCode.FULL_VALIDATION_FAILED,
                    "Full-document validation failed: " + "; ".join(validation.hard_failures[:5]),
                )

            # Reconciliation items needing a human block publication but are not corruption.
            if reconciliation_summary.overall_status in MANUAL_REVIEW_RECONCILIATION:
                cls._transition(report, exec_manifest, ManifestStatus.VALIDATED,
                                "Structural and full-document validation passed.")
                cls._transition(report, exec_manifest, ManifestStatus.REQUIRES_MANUAL_REVIEW,
                                "Unresolved reconciliation items require a human tax reviewer.")
                shutil.rmtree(staging_dir, ignore_errors=True)
                report.staging_discarded = True
                report.rollback_occurred = True
                report.status = ExecutionStatus.REQUIRES_MANUAL_REVIEW
                report.publication_state = PublicationState.NOT_PUBLISHED
                report.failure_code = FailureCode.DATA_RECONCILIATION_MANUAL_REVIEW
                report.failure_detail = (
                    f"D2 reconciliation returned {reconciliation_summary.overall_status.value}; "
                    "artifact withheld from publication pending human review."
                )
                report.output_hash = None
                report.ended_at = _now()
                report.lineage = ExecutionLineageBuilder.build(report, reconciliation_summary)
                report.execution_manifest_status = exec_manifest.status
                return report

            # ---- 11. STATE: VALIDATED -> COMPLETED ---------------------
            cls._transition(report, exec_manifest, ManifestStatus.VALIDATED,
                            "All hard validation gates passed.")
            cls._transition(report, exec_manifest, ManifestStatus.COMPLETED,
                            "Full-document roll-forward completed.")

            # ---- 12. ATOMIC PUBLICATION (§9, §15, §18) -----------------
            if request.output_path.exists():
                request.output_path.unlink()
            os.replace(str(staged_output), str(request.output_path))
            shutil.rmtree(staging_dir, ignore_errors=True)

            report.output_hash = compute_file_sha256(request.output_path)
            report.status = ExecutionStatus.COMPLETED
            report.publication_state = PublicationState.FINAL_VALIDATED
            report.execution_manifest_status = exec_manifest.status

            ExecutionLedger(
                execution_id=request.execution_id,
                manifest_id=manifest.manifest_id,
                manifest_version=manifest.manifest_version,
                mutation_plan_id=plan.plan_id,
                mutation_plan_digest=plan_digest,
                output_hash=report.output_hash,
                published_at=_now(),
            ).save(request.output_path)

        except Exception as ex:  # noqa: BLE001 - any escape must roll back, never publish
            return cls._abort(
                report, staging_dir, exec_manifest,
                FailureCode.INTERNAL_EXECUTION_ERROR,
                f"Unhandled execution error: {type(ex).__name__}: {ex}",
            )
        finally:
            report.template_preserved = (
                request.template_path.exists()
                and compute_file_sha256(request.template_path) == template_hash
            )

        # ---- 13. GROUND TRUTH EVALUATION (§19) -------------------------
        if request.ground_truth_path and request.ground_truth_path.exists():
            status_before, publication_before = report.status, report.publication_state
            report.ground_truth_evaluation = GroundTruthEvaluator.evaluate(
                generated_path=request.output_path,
                ground_truth_path=request.ground_truth_path,
                expected_mutations=expected_mutations,
            )
            # Ground Truth is evaluation-only: execution state is provably untouched.
            report.status, report.publication_state = status_before, publication_before

        # ---- 14. LINEAGE (§17) -----------------------------------------
        report.ended_at = _now()
        report.lineage = ExecutionLineageBuilder.build(report, reconciliation_summary)
        return report

    # ------------------------------------------------------------------
    # Gating helpers
    # ------------------------------------------------------------------

    @classmethod
    def _verify_approval(
        cls,
        manifest: RollForwardManifest,
        plan: MutationPlan,
        request: ExecutionRequest,
    ) -> Optional[Tuple[FailureCode, str]]:
        """Rejects missing, system/agent, stale, and version-mismatched approvals."""
        if manifest.status != ManifestStatus.APPROVED:
            return (
                FailureCode.APPROVAL_MISSING,
                f"Manifest status is {manifest.status.value}; execution requires APPROVED.",
            )
        if not manifest.approved_by:
            return (FailureCode.APPROVAL_MISSING, "Manifest carries no approver identity.")
        if not manifest.approved_at:
            return (FailureCode.APPROVAL_MISSING, "Manifest carries no approval timestamp.")
        if manifest.approved_by.strip().lower() in RESERVED_NON_HUMAN_APPROVERS:
            return (
                FailureCode.APPROVAL_NOT_BY_USER,
                f"Approver '{manifest.approved_by}' is a system/agent identity. "
                "Only an explicit human user may approve a roll-forward manifest.",
            )

        approval_logs = [
            log for log in manifest.history if log.to_state == ManifestStatus.APPROVED
        ]
        if approval_logs and approval_logs[-1].actor != "user":
            return (
                FailureCode.APPROVAL_NOT_BY_USER,
                f"The recorded APPROVED transition was performed by actor "
                f"'{approval_logs[-1].actor}', not 'user'.",
            )

        if manifest.approved_manifest_version != manifest.manifest_version:
            return (
                FailureCode.APPROVAL_VERSION_MISMATCH,
                f"Stale approval: approved version {manifest.approved_manifest_version} "
                f"!= current manifest version {manifest.manifest_version}.",
            )
        if plan.manifest_version != manifest.manifest_version:
            return (
                FailureCode.PLAN_VERSION_MISMATCH,
                f"MutationPlan targets manifest version {plan.manifest_version} "
                f"but the approved manifest is version {manifest.manifest_version}.",
            )
        if plan.manifest_id != manifest.manifest_id:
            return (
                FailureCode.PLAN_VERSION_MISMATCH,
                f"MutationPlan belongs to manifest {plan.manifest_id}, not {manifest.manifest_id}.",
            )
        if request.expected_approver and manifest.approved_by != request.expected_approver:
            return (
                FailureCode.APPROVAL_NOT_BY_USER,
                f"Approver '{manifest.approved_by}' does not match the expected approver "
                f"'{request.expected_approver}'.",
            )
        return None

    @classmethod
    def _check_idempotence(
        cls,
        request: ExecutionRequest,
        plan_digest: str,
        scope: ApprovedScope,
        report: RollForwardExecutionReport,
    ) -> Optional[RollForwardExecutionReport]:
        """Detects an already-applied, already-published identical execution."""
        ledger = ExecutionLedger.load(request.output_path)
        if ledger is None or not request.output_path.exists():
            return None
        if (
            ledger.manifest_id != request.manifest.manifest_id
            or ledger.manifest_version != request.manifest.manifest_version
            or ledger.mutation_plan_digest != plan_digest
        ):
            return None

        current_hash = compute_file_sha256(request.output_path)
        if current_hash != ledger.output_hash:
            return None

        # Structural confirmation that the postcondition really is already present.
        doc = Document(str(request.output_path))
        for spec in scope.executable_specs:
            if spec.table_index >= len(doc.tables):
                return None
            if len(doc.tables[spec.table_index].rows) != spec.target_row_count:
                return None

        report.status = ExecutionStatus.NOOP
        report.publication_state = PublicationState.FINAL_VALIDATED
        report.idempotent_noop = True
        report.output_hash = current_hash
        report.ended_at = _now()
        report.execution_manifest_status = ManifestStatus.COMPLETED
        report.failure_detail = (
            f"Execution {ledger.execution_id} already applied this exact MutationPlan "
            f"(digest {plan_digest[:12]}...) and published output {current_hash[:12]}.... "
            "No duplicate mutation performed."
        )
        report.lineage = ExecutionLineageBuilder.build(report, None)
        return report

    # ------------------------------------------------------------------
    # Scope, transitions, failure handling
    # ------------------------------------------------------------------

    @staticmethod
    def _scoped_execution_manifest(
        manifest: RollForwardManifest, scope: ApprovedScope
    ) -> RollForwardManifest:
        """A deep copy narrowed to exactly the approved, READY regions.

        Narrowing only *removes* regions -- it never relaxes a gate. The caller's
        manifest object is never mutated, so the governed artifact the user
        approved is left byte-identical while the execution lifecycle is tracked
        on this projection using the one and only RollForwardStateMachine.
        """
        scoped = manifest.model_copy(deep=True)
        executable_ids = scope.executable_region_ids
        scoped.regions = [r for r in scoped.regions if r.region_id in executable_ids]
        scoped.figures = []
        return scoped

    @staticmethod
    def _transition(
        report: RollForwardExecutionReport,
        exec_manifest: RollForwardManifest,
        target: ManifestStatus,
        reason: str,
    ) -> None:
        before = exec_manifest.status
        RollForwardStateMachine.transition(exec_manifest, target, actor="system", reason=reason)
        report.state_transitions.append(f"{before.value}->{target.value}")
        report.execution_manifest_status = exec_manifest.status

    @staticmethod
    def _structural_changes(
        baseline: DocumentBaseline,
        output_path: Path,
        specs: Sequence[TableMutationSpec],
    ) -> List[StructuralChangeRecord]:
        doc = Document(str(output_path))
        records: List[StructuralChangeRecord] = []
        for spec in specs:
            base_sig = baseline.table_signatures.get(spec.table_index)
            if spec.table_index >= len(doc.tables):
                continue
            out_sig = table_structure_signature(spec.table_index, doc.tables[spec.table_index])
            rows_before = base_sig.row_count if base_sig else spec.initial_row_count
            records.append(
                StructuralChangeRecord(
                    region_id=spec.target_region_id,
                    table_index=spec.table_index,
                    operation=spec.operation,
                    rows_before=rows_before,
                    rows_after=out_sig.row_count,
                    rows_inserted=out_sig.row_count - rows_before,
                    columns_before=base_sig.column_count if base_sig else out_sig.column_count,
                    columns_after=out_sig.column_count,
                    precondition_fingerprint=base_sig.semantic_fingerprint if base_sig else "",
                    postcondition_fingerprint=out_sig.semantic_fingerprint,
                )
            )
        return records

    @classmethod
    def _abort(
        cls,
        report: RollForwardExecutionReport,
        staging_dir: Path,
        exec_manifest: Optional[RollForwardManifest],
        code: FailureCode,
        detail: str,
    ) -> RollForwardExecutionReport:
        """Discards the staging transaction; nothing partial may ever be published."""
        shutil.rmtree(staging_dir, ignore_errors=True)
        report.staging_discarded = True
        report.rollback_occurred = True
        report.output_hash = None
        if exec_manifest is not None and exec_manifest.status == ManifestStatus.EXECUTING:
            RollForwardStateMachine.mark_failed(exec_manifest, error_message=detail)
            report.state_transitions.append(
                f"{ManifestStatus.EXECUTING.value}->{ManifestStatus.FAILED.value}"
            )
            report.execution_manifest_status = exec_manifest.status
        return cls._fail(report, code, detail)

    @staticmethod
    def _fail(
        report: RollForwardExecutionReport, code: FailureCode, detail: str
    ) -> RollForwardExecutionReport:
        report.status = ExecutionStatus.FAILED
        report.publication_state = PublicationState.NOT_PUBLISHED
        report.failure_code = code
        report.failure_detail = detail
        report.ended_at = _now()
        if report.lineage is None:
            report.lineage = ExecutionLineageBuilder.build(report, None)
        return report


# ============================================================================
# 9. MODULE HELPERS
# ============================================================================

def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _mutation_plan_digest(plan: MutationPlan) -> str:
    """Content digest of a MutationPlan.

    Deliberately excludes `plan_id` and `created_at`: those identify *this*
    planning run, not the mutation it describes. Idempotence must key on what
    would actually be written to the document, so that re-planning the same
    approved change is recognised as already applied rather than re-executed.
    """
    payload = plan.model_dump(mode="json", exclude={"plan_id", "created_at"})
    return hashlib.sha256(
        json.dumps(payload, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()


def _summarize_exclusions(records: Sequence[RegionExecutionRecord]) -> Dict[str, int]:
    counts: Dict[str, int] = {}
    for r in records:
        key = r.exclusion_reason.value if r.exclusion_reason else "UNSPECIFIED"
        counts[key] = counts.get(key, 0) + 1
    return counts


def _digest_reconciliation(summary: ManifestReconciliationSummary) -> ReconciliationDigest:
    """Aggregates D2 output without carrying any document plaintext forward."""
    return ReconciliationDigest(
        overall_status=summary.overall_status.value,
        total_tables=summary.total_tables,
        total_cells=summary.total_cells,
        matched_cells=summary.matched_cells,
        mismatched_cells=summary.mismatched_cells,
        missing_cells=summary.missing_cells,
        type_mismatches=summary.type_mismatches,
        format_mismatches=summary.format_mismatches,
        manual_review_items=summary.manual_review_items,
        blocked_items=summary.blocked_items,
        source_freshness_verified=summary.source_freshness_verified,
        per_table=[
            {
                "table_index": t.table_index,
                "region_id": t.target_region_id,
                "status": t.status.value,
                "total_cells": t.total_cells,
                "matched_cells": t.matched_cells,
                "mismatched_cells": t.mismatched_cells,
                "inserted_rows": t.inserted_rows,
                "target_row_count": t.target_row_count,
            }
            for t in summary.table_summaries
        ],
    )
