"""
Full Document Validation Gate (Phase D3)
=========================================
Location: foundation/applications/rollforward/full_validation.py

Validates an entire generated Local File artifact -- not just the approved
target tables -- against a deterministic baseline captured from the source
template before any mutation was applied.

Design principles:
1. Success is NEVER "the document opens in Word". Every registered hard rule
   must pass before the artifact may be published.
2. Every difference between baseline and output must be explicitly declared as
   an ExpectedMutation. Anything undeclared is a hard failure.
3. Only deterministic Foundation evidence is used (perception pipeline, OOXML
   package inspection, python-docx structure). No heuristics, no tolerances,
   no invented validations.
4. This module never mutates a document and never publishes anything; it only
   observes and reports.
"""
from __future__ import annotations

from enum import Enum
import hashlib
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Set, Tuple
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table, _Row
from pydantic import BaseModel, Field

from applications.rollforward.structural_writeback import FingerprintService
from perception.anchor_builder import assign_anchors
from perception.element_classifier import classify_blocks
from perception.parser import extract_geometry, extract_media_manifest

_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

# OOXML parts that must exist for the package to be a well-formed WordprocessingML document.
REQUIRED_PACKAGE_PARTS = ("[Content_Types].xml", "word/document.xml", "_rels/.rels")


# ============================================================================
# 1. CHECK TAXONOMY
# ============================================================================

class ValidationCategory(str, Enum):
    """The nine validation families required by the Phase D3 contract."""
    PACKAGE_PARSE = "PACKAGE_PARSE"                      # A
    PERCEPTION = "PERCEPTION"                            # B
    ELEMENT_INVENTORY = "ELEMENT_INVENTORY"              # C
    TARGET_REGIONS = "TARGET_REGIONS"                    # D
    NON_TARGET_INTEGRITY = "NON_TARGET_INTEGRITY"        # E
    TABLE_STRUCTURE = "TABLE_STRUCTURE"                  # F
    MEDIA_RELATIONSHIPS = "MEDIA_RELATIONSHIPS"          # G
    HEADERS_FOOTERS = "HEADERS_FOOTERS"                  # H
    DOCUMENT_STRUCTURE = "DOCUMENT_STRUCTURE"            # I


class CheckSeverity(str, Enum):
    """A BLOCKER failure forbids publication; a WARNING is reported only."""
    BLOCKER = "BLOCKER"
    WARNING = "WARNING"


class CheckStatus(str, Enum):
    PASSED = "PASSED"
    FAILED = "FAILED"
    SKIPPED = "SKIPPED"


class ValidationCheck(BaseModel):
    """One registered, deterministic validation rule outcome."""
    check_id: str
    category: ValidationCategory
    severity: CheckSeverity = CheckSeverity.BLOCKER
    status: CheckStatus = CheckStatus.PASSED
    message: str = ""
    evidence: Dict[str, Any] = Field(default_factory=dict)


class ExpectedMutation(BaseModel):
    """The only differences a generated artifact is permitted to exhibit."""
    region_id: str
    table_index: int = Field(ge=0)
    baseline_row_count: int = Field(ge=0)
    expected_row_count: int = Field(ge=0)
    expected_inserted_rows: int = Field(ge=0)
    operation: str = "INSERT_ROWS"


# ============================================================================
# 2. DETERMINISTIC DOCUMENT BASELINE
# ============================================================================

class TableStructureSignature(BaseModel):
    """Structural fingerprint of a single DOCX table."""
    table_index: int
    row_count: int
    column_count: int
    grid_widths: List[int] = Field(default_factory=list)
    merge_signature: str = ""
    header_signature: str = ""
    semantic_fingerprint: str = ""


class DocumentBaseline(BaseModel):
    """Immutable pre-execution observation of the source template.

    Everything the FullDocumentValidator later compares against is captured
    here, before a single byte of the template is touched.
    """
    document_path: str
    document_sha256: str
    package_parts: List[str] = Field(default_factory=list)
    declared_relationships: List[str] = Field(default_factory=list)
    orphan_relationship_refs: List[str] = Field(default_factory=list)
    media_ids: List[str] = Field(default_factory=list)
    block_count: int = 0
    anchor_count: int = 0
    element_count: int = 0
    paragraph_count: int = 0
    section_count: int = 0
    table_count: int = 0
    body_order_signature: str = ""
    heading_signature: str = ""
    heading_count: int = 0
    header_footer_signature: str = ""
    table_signatures: Dict[int, TableStructureSignature] = Field(default_factory=dict)
    target_table_indices: List[int] = Field(default_factory=list)
    target_region_fingerprints: Dict[int, str] = Field(default_factory=dict)
    non_target_fingerprint: str = ""

    @classmethod
    def capture(cls, doc_path: Path, target_table_indices: Sequence[int]) -> "DocumentBaseline":
        """Captures the complete deterministic baseline of a DOCX artifact."""
        targets: Set[int] = set(int(i) for i in target_table_indices)
        doc = Document(str(doc_path))

        parts, declared_rels, orphan_refs = _inspect_package(doc_path)
        blocks = extract_geometry(str(doc_path))
        anchors = assign_anchors(blocks, "docx")
        elements = classify_blocks(blocks, "docx", anchors)

        table_signatures: Dict[int, TableStructureSignature] = {
            t_idx: table_structure_signature(t_idx, table) for t_idx, table in enumerate(doc.tables)
        }
        heading_sig, heading_count = _heading_signature(doc)

        return cls(
            document_path=str(doc_path),
            document_sha256=compute_file_sha256(doc_path),
            package_parts=parts,
            declared_relationships=declared_rels,
            orphan_relationship_refs=orphan_refs,
            media_ids=sorted(a.media_id for a in extract_media_manifest(str(doc_path), "docx")),
            block_count=len(blocks),
            anchor_count=len(anchors),
            element_count=len(elements),
            paragraph_count=len(doc.paragraphs),
            section_count=len(doc.sections),
            table_count=len(doc.tables),
            body_order_signature=_body_order_signature(doc),
            heading_signature=heading_sig,
            heading_count=heading_count,
            header_footer_signature=_header_footer_signature(blocks),
            table_signatures=table_signatures,
            target_table_indices=sorted(targets),
            target_region_fingerprints={
                t_idx: table_signatures[t_idx].semantic_fingerprint
                for t_idx in sorted(targets)
                if t_idx in table_signatures
            },
            non_target_fingerprint=FingerprintService.compute_document_non_target_fingerprint(
                doc, targets
            ),
        )


class FullDocumentValidationReport(BaseModel):
    """Complete outcome of the D3 full-document validation gate."""
    is_valid: bool = True
    checks: List[ValidationCheck] = Field(default_factory=list)
    hard_failures: List[str] = Field(default_factory=list)
    warnings: List[str] = Field(default_factory=list)
    baseline_sha256: str = ""
    output_sha256: str = ""
    output_baseline: Optional[DocumentBaseline] = None

    @property
    def passed_count(self) -> int:
        return sum(1 for c in self.checks if c.status == CheckStatus.PASSED)

    @property
    def failed_count(self) -> int:
        return sum(1 for c in self.checks if c.status == CheckStatus.FAILED)

    def summary(self) -> Dict[str, Any]:
        by_category: Dict[str, Dict[str, int]] = {}
        for c in self.checks:
            bucket = by_category.setdefault(c.category.value, {"passed": 0, "failed": 0, "skipped": 0})
            bucket[c.status.value.lower()] += 1
        return {
            "is_valid": self.is_valid,
            "total_checks": len(self.checks),
            "passed": self.passed_count,
            "failed": self.failed_count,
            "hard_failures": list(self.hard_failures),
            "warnings": list(self.warnings),
            "by_category": by_category,
        }


# ============================================================================
# 3. FULL DOCUMENT VALIDATOR
# ============================================================================

class FullDocumentValidator:
    """Validates an entire generated artifact against its pre-execution baseline."""

    @classmethod
    def validate(
        cls,
        baseline: DocumentBaseline,
        output_path: Path,
        expected_mutations: Sequence[ExpectedMutation],
    ) -> FullDocumentValidationReport:
        """Runs every registered hard rule over the generated document."""
        report = FullDocumentValidationReport(baseline_sha256=baseline.document_sha256)
        expected_by_table: Dict[int, ExpectedMutation] = {m.table_index: m for m in expected_mutations}

        # --- A. PACKAGE / PARSE ------------------------------------------
        if not output_path.exists():
            cls._fail(
                report, "PKG-001", ValidationCategory.PACKAGE_PARSE,
                "Generated artifact does not exist.", {"path": str(output_path)},
            )
            return report

        report.output_sha256 = compute_file_sha256(output_path)

        try:
            with zipfile.ZipFile(str(output_path)) as zf:
                bad = zf.testzip()
            if bad is not None:
                raise zipfile.BadZipFile(f"Corrupt package member: {bad}")
        except (zipfile.BadZipFile, OSError) as ex:
            cls._fail(
                report, "PKG-002", ValidationCategory.PACKAGE_PARSE,
                f"OOXML package is not structurally valid: {ex}", {"path": str(output_path)},
            )
            return report
        cls._pass(report, "PKG-002", ValidationCategory.PACKAGE_PARSE, "OOXML package integrity verified.")

        out_parts, out_rels, out_orphans = _inspect_package(output_path)
        missing_required = [p for p in REQUIRED_PACKAGE_PARTS if p not in out_parts]
        if missing_required:
            cls._fail(
                report, "PKG-003", ValidationCategory.PACKAGE_PARSE,
                f"Required OOXML parts missing: {missing_required}", {"missing": missing_required},
            )
            return report
        cls._pass(report, "PKG-003", ValidationCategory.PACKAGE_PARSE, "All required OOXML parts present.")

        lost_parts = sorted(set(baseline.package_parts) - set(out_parts))
        if lost_parts:
            cls._fail(
                report, "PKG-004", ValidationCategory.PACKAGE_PARSE,
                f"Package parts lost during mutation: {lost_parts}", {"lost_parts": lost_parts},
            )
        else:
            cls._pass(
                report, "PKG-004", ValidationCategory.PACKAGE_PARSE,
                f"All {len(baseline.package_parts)} baseline package parts retained.",
            )

        try:
            doc_out = Document(str(output_path))
        except Exception as ex:  # noqa: BLE001 - any parse failure is a hard gate
            cls._fail(
                report, "PKG-005", ValidationCategory.PACKAGE_PARSE,
                f"Generated artifact failed to open: {ex}", {},
            )
            return report
        cls._pass(report, "PKG-005", ValidationCategory.PACKAGE_PARSE, "Generated artifact opens successfully.")

        # --- B. PERCEPTION ------------------------------------------------
        try:
            blocks = extract_geometry(str(output_path))
            anchors = assign_anchors(blocks, "docx")
            elements = classify_blocks(blocks, "docx", anchors)
        except Exception as ex:  # noqa: BLE001 - perception must never regress
            cls._fail(
                report, "PER-001", ValidationCategory.PERCEPTION,
                f"Foundation perception raised an unexpected failure: {ex}", {},
            )
            return report

        cls._pass(
            report, "PER-001", ValidationCategory.PERCEPTION,
            f"Full Foundation perception succeeded: {len(blocks)} blocks, "
            f"{len(anchors)} anchors, {len(elements)} elements.",
            {"blocks": len(blocks), "anchors": len(anchors), "elements": len(elements)},
        )

        if len(anchors) == len(blocks):
            cls._pass(
                report, "PER-002", ValidationCategory.PERCEPTION,
                "Every perceived geometry block resolved to an anchor.",
            )
        else:
            cls._fail(
                report, "PER-002", ValidationCategory.PERCEPTION,
                f"Anchor assignment incomplete: {len(anchors)} anchors for {len(blocks)} blocks.",
                {"blocks": len(blocks), "anchors": len(anchors)},
            )

        perceived_tables = {b.get("table_index") for b in blocks if b.get("table_index") is not None}
        missing_targets = sorted(t for t in expected_by_table if t not in perceived_tables)
        if missing_targets:
            cls._fail(
                report, "PER-003", ValidationCategory.PERCEPTION,
                f"Approved target tables absent from perceived geometry: {missing_targets}",
                {"missing_tables": missing_targets},
            )
        else:
            cls._pass(
                report, "PER-003", ValidationCategory.PERCEPTION,
                f"All {len(expected_by_table)} approved target tables remain addressable after mutation.",
            )

        # --- C. ELEMENT INVENTORY ----------------------------------------
        element_delta = len(elements) - baseline.element_count
        if len(elements) < baseline.element_count:
            cls._fail(
                report, "INV-001", ValidationCategory.ELEMENT_INVENTORY,
                f"Element loss detected: baseline {baseline.element_count} -> output {len(elements)}.",
                {"baseline": baseline.element_count, "output": len(elements), "delta": element_delta},
            )
        else:
            cls._pass(
                report, "INV-001", ValidationCategory.ELEMENT_INVENTORY,
                f"No element loss: baseline {baseline.element_count} -> output {len(elements)} (+{element_delta}).",
                {"baseline": baseline.element_count, "output": len(elements), "delta": element_delta},
            )

        out_paragraph_count = len(doc_out.paragraphs)
        if out_paragraph_count != baseline.paragraph_count:
            cls._fail(
                report, "INV-002", ValidationCategory.ELEMENT_INVENTORY,
                f"Body paragraph count changed: {baseline.paragraph_count} -> {out_paragraph_count}. "
                "Row insertion must not alter body-level paragraphs.",
                {"baseline": baseline.paragraph_count, "output": out_paragraph_count},
            )
        else:
            cls._pass(
                report, "INV-002", ValidationCategory.ELEMENT_INVENTORY,
                f"Body paragraph inventory unchanged ({out_paragraph_count}).",
            )

        # --- D. TARGET REGIONS -------------------------------------------
        out_signatures = {t_idx: table_structure_signature(t_idx, t) for t_idx, t in enumerate(doc_out.tables)}

        for t_idx, expected in sorted(expected_by_table.items()):
            base_sig = baseline.table_signatures.get(t_idx)
            out_sig = out_signatures.get(t_idx)
            if base_sig is None or out_sig is None:
                cls._fail(
                    report, f"TGT-{t_idx:03d}-000", ValidationCategory.TARGET_REGIONS,
                    f"Approved target table {t_idx} not found in baseline or output.",
                    {"region_id": expected.region_id},
                )
                continue

            if out_sig.row_count == expected.expected_row_count:
                cls._pass(
                    report, f"TGT-{t_idx:03d}-ROWS", ValidationCategory.TARGET_REGIONS,
                    f"Table {t_idx} ({expected.region_id}) row count "
                    f"{base_sig.row_count} -> {out_sig.row_count} as approved.",
                    {"region_id": expected.region_id, "expected": expected.expected_row_count},
                )
            else:
                cls._fail(
                    report, f"TGT-{t_idx:03d}-ROWS", ValidationCategory.TARGET_REGIONS,
                    f"Table {t_idx} ({expected.region_id}) row count is {out_sig.row_count}, "
                    f"approved plan declared {expected.expected_row_count}.",
                    {"region_id": expected.region_id, "actual": out_sig.row_count},
                )

            if out_sig.column_count == base_sig.column_count:
                cls._pass(
                    report, f"TGT-{t_idx:03d}-COLS", ValidationCategory.TARGET_REGIONS,
                    f"Table {t_idx} column count unchanged ({out_sig.column_count}).",
                )
            else:
                cls._fail(
                    report, f"TGT-{t_idx:03d}-COLS", ValidationCategory.TARGET_REGIONS,
                    f"Table {t_idx} column count changed: {base_sig.column_count} -> {out_sig.column_count}.",
                    {"region_id": expected.region_id},
                )

            if len(set(out_sig.grid_widths)) <= 1:
                cls._pass(
                    report, f"TGT-{t_idx:03d}-GRID", ValidationCategory.TARGET_REGIONS,
                    f"Table {t_idx} grid width consistent across all {out_sig.row_count} rows.",
                )
            else:
                cls._fail(
                    report, f"TGT-{t_idx:03d}-GRID", ValidationCategory.TARGET_REGIONS,
                    f"Table {t_idx} has inconsistent grid widths: {sorted(set(out_sig.grid_widths))}.",
                    {"region_id": expected.region_id, "grid_widths": out_sig.grid_widths},
                )

            if out_sig.header_signature == base_sig.header_signature:
                cls._pass(
                    report, f"TGT-{t_idx:03d}-HDR", ValidationCategory.TARGET_REGIONS,
                    f"Table {t_idx} header row preserved verbatim.",
                )
            else:
                cls._fail(
                    report, f"TGT-{t_idx:03d}-HDR", ValidationCategory.TARGET_REGIONS,
                    f"Table {t_idx} header row changed during mutation.",
                    {"region_id": expected.region_id},
                )

        # --- E. NON-TARGET INTEGRITY --------------------------------------
        target_set = set(baseline.target_table_indices) | set(expected_by_table)
        out_non_target_fp = FingerprintService.compute_document_non_target_fingerprint(doc_out, target_set)
        if out_non_target_fp == baseline.non_target_fingerprint:
            cls._pass(
                report, "NTI-001", ValidationCategory.NON_TARGET_INTEGRITY,
                "Zero semantic drift across all non-target regions.",
                {"fingerprint": out_non_target_fp},
            )
        else:
            drifted = [
                t_idx
                for t_idx, sig in baseline.table_signatures.items()
                if t_idx not in target_set
                and out_signatures.get(t_idx) is not None
                and out_signatures[t_idx].semantic_fingerprint != sig.semantic_fingerprint
            ]
            cls._fail(
                report, "NTI-001", ValidationCategory.NON_TARGET_INTEGRITY,
                "Unexpected semantic drift detected in non-target document regions.",
                {
                    "baseline_fingerprint": baseline.non_target_fingerprint,
                    "output_fingerprint": out_non_target_fp,
                    "drifted_tables": drifted,
                },
            )

        undeclared_tables = [
            t_idx
            for t_idx, sig in baseline.table_signatures.items()
            if t_idx not in expected_by_table
            and out_signatures.get(t_idx) is not None
            and out_signatures[t_idx].semantic_fingerprint != sig.semantic_fingerprint
        ]
        if undeclared_tables:
            cls._fail(
                report, "NTI-002", ValidationCategory.NON_TARGET_INTEGRITY,
                f"Tables changed without an approved ExpectedMutation: {undeclared_tables}",
                {"undeclared_tables": undeclared_tables},
            )
        else:
            cls._pass(
                report, "NTI-002", ValidationCategory.NON_TARGET_INTEGRITY,
                "Every observed table change is covered by an approved ExpectedMutation.",
            )

        # --- F. TABLE STRUCTURE -------------------------------------------
        if len(doc_out.tables) == baseline.table_count:
            cls._pass(
                report, "TBL-001", ValidationCategory.TABLE_STRUCTURE,
                f"Table count unchanged ({baseline.table_count}).",
            )
        else:
            cls._fail(
                report, "TBL-001", ValidationCategory.TABLE_STRUCTURE,
                f"Table count changed: {baseline.table_count} -> {len(doc_out.tables)}.",
                {"baseline": baseline.table_count, "output": len(doc_out.tables)},
            )

        col_changed = [
            t_idx
            for t_idx, sig in baseline.table_signatures.items()
            if out_signatures.get(t_idx) is not None
            and out_signatures[t_idx].column_count != sig.column_count
        ]
        if col_changed:
            cls._fail(
                report, "TBL-002", ValidationCategory.TABLE_STRUCTURE,
                f"Column count changed in tables {col_changed}.", {"tables": col_changed},
            )
        else:
            cls._pass(
                report, "TBL-002", ValidationCategory.TABLE_STRUCTURE,
                "Column counts preserved across every table in the document.",
            )

        merge_changed = [
            t_idx
            for t_idx, sig in baseline.table_signatures.items()
            if t_idx not in expected_by_table
            and out_signatures.get(t_idx) is not None
            and out_signatures[t_idx].merge_signature != sig.merge_signature
        ]
        if merge_changed:
            cls._fail(
                report, "TBL-003", ValidationCategory.TABLE_STRUCTURE,
                f"Merge topology changed in non-target tables {merge_changed}.", {"tables": merge_changed},
            )
        else:
            cls._pass(
                report, "TBL-003", ValidationCategory.TABLE_STRUCTURE,
                "Merge topology preserved in every non-target table.",
            )

        orphan_merges = _orphan_vmerge_cells(doc_out)
        if orphan_merges:
            cls._fail(
                report, "TBL-004", ValidationCategory.TABLE_STRUCTURE,
                f"Orphaned vMerge continuation cells introduced: {orphan_merges[:10]}",
                {"orphans": orphan_merges},
            )
        else:
            cls._pass(
                report, "TBL-004", ValidationCategory.TABLE_STRUCTURE,
                "No orphaned vMerge continuation cells anywhere in the document.",
            )

        # --- G. MEDIA / RELATIONSHIPS -------------------------------------
        out_media = sorted(a.media_id for a in extract_media_manifest(str(output_path), "docx"))
        lost_media = sorted(set(baseline.media_ids) - set(out_media))
        if lost_media:
            cls._fail(
                report, "MED-001", ValidationCategory.MEDIA_RELATIONSHIPS,
                f"Embedded media lost during mutation: {lost_media}", {"lost_media": lost_media},
            )
        else:
            cls._pass(
                report, "MED-001", ValidationCategory.MEDIA_RELATIONSHIPS,
                f"All {len(baseline.media_ids)} embedded media assets retained.",
                {"media_count": len(out_media)},
            )

        new_orphans = sorted(set(out_orphans) - set(baseline.orphan_relationship_refs))
        if new_orphans:
            cls._fail(
                report, "MED-002", ValidationCategory.MEDIA_RELATIONSHIPS,
                f"Orphaned relationship references introduced: {new_orphans}", {"orphans": new_orphans},
            )
        else:
            cls._pass(
                report, "MED-002", ValidationCategory.MEDIA_RELATIONSHIPS,
                "No orphaned relationship references introduced.",
                {"pre_existing_orphans": len(baseline.orphan_relationship_refs)},
            )

        lost_rels = sorted(set(baseline.declared_relationships) - set(out_rels))
        if lost_rels:
            cls._fail(
                report, "MED-003", ValidationCategory.MEDIA_RELATIONSHIPS,
                f"Declared relationships lost: {lost_rels}", {"lost_relationships": lost_rels},
            )
        else:
            cls._pass(
                report, "MED-003", ValidationCategory.MEDIA_RELATIONSHIPS,
                f"All {len(baseline.declared_relationships)} declared relationships retained.",
            )

        # --- H. HEADERS / FOOTERS ------------------------------------------
        out_hf = _header_footer_signature(blocks)
        if out_hf == baseline.header_footer_signature:
            cls._pass(
                report, "HDF-001", ValidationCategory.HEADERS_FOOTERS,
                "Headers and footers preserved exactly (not an approved target).",
            )
        else:
            cls._fail(
                report, "HDF-001", ValidationCategory.HEADERS_FOOTERS,
                "Header/footer content changed but was never declared as a target region.",
                {"baseline": baseline.header_footer_signature, "output": out_hf},
            )

        # --- I. DOCUMENT STRUCTURE -----------------------------------------
        if len(doc_out.sections) == baseline.section_count:
            cls._pass(
                report, "DOC-001", ValidationCategory.DOCUMENT_STRUCTURE,
                f"Section count unchanged ({baseline.section_count}).",
            )
        else:
            cls._fail(
                report, "DOC-001", ValidationCategory.DOCUMENT_STRUCTURE,
                f"Section count changed: {baseline.section_count} -> {len(doc_out.sections)}.", {},
            )

        out_heading_sig, out_heading_count = _heading_signature(doc_out)
        if out_heading_sig == baseline.heading_signature:
            cls._pass(
                report, "DOC-002", ValidationCategory.DOCUMENT_STRUCTURE,
                f"Heading sequence preserved ({out_heading_count} headings, order and text identical).",
            )
        else:
            cls._fail(
                report, "DOC-002", ValidationCategory.DOCUMENT_STRUCTURE,
                f"Heading sequence changed: {baseline.heading_count} -> {out_heading_count} headings "
                "or heading text/order drifted.",
                {"baseline_headings": baseline.heading_count, "output_headings": out_heading_count},
            )

        out_body_sig = _body_order_signature(doc_out)
        if out_body_sig == baseline.body_order_signature:
            cls._pass(
                report, "DOC-003", ValidationCategory.DOCUMENT_STRUCTURE,
                "Body reading order preserved: paragraph/table interleaving is byte-identical in sequence.",
            )
        else:
            cls._fail(
                report, "DOC-003", ValidationCategory.DOCUMENT_STRUCTURE,
                "Body reading order changed: paragraph/table ordering drifted.",
                {"baseline": baseline.body_order_signature, "output": out_body_sig},
            )

        return report

    # ------------------------------------------------------------------
    # Check recording helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _pass(
        report: FullDocumentValidationReport,
        check_id: str,
        category: ValidationCategory,
        message: str,
        evidence: Optional[Dict[str, Any]] = None,
    ) -> None:
        report.checks.append(
            ValidationCheck(
                check_id=check_id,
                category=category,
                severity=CheckSeverity.BLOCKER,
                status=CheckStatus.PASSED,
                message=message,
                evidence=evidence or {},
            )
        )

    @staticmethod
    def _fail(
        report: FullDocumentValidationReport,
        check_id: str,
        category: ValidationCategory,
        message: str,
        evidence: Optional[Dict[str, Any]] = None,
    ) -> None:
        report.checks.append(
            ValidationCheck(
                check_id=check_id,
                category=category,
                severity=CheckSeverity.BLOCKER,
                status=CheckStatus.FAILED,
                message=message,
                evidence=evidence or {},
            )
        )
        report.is_valid = False
        report.hard_failures.append(f"[{check_id}] {message}")


# ============================================================================
# 4. DETERMINISTIC OBSERVATION PRIMITIVES
# ============================================================================

def compute_file_sha256(path: Path) -> str:
    """SHA256 hex digest of a file's bytes."""
    hasher = hashlib.sha256()
    with open(path, "rb") as f:
        while chunk := f.read(65536):
            hasher.update(chunk)
    return hasher.hexdigest()


def _row_grid_width(row: _Row) -> int:
    """Effective column span of one table row, honouring w:gridSpan."""
    width = 0
    for tc in row._tr.findall(qn("w:tc")):
        tcPr = tc.find(qn("w:tcPr"))
        gs = tcPr.find(qn("w:gridSpan")) if tcPr is not None else None
        width += int(gs.get(qn("w:val")) or "1") if gs is not None else 1
    return width


def table_structure_signature(table_index: int, table: Table) -> TableStructureSignature:
    """Structural + semantic fingerprint of one table."""
    merge_hasher = hashlib.sha256()
    header_hasher = hashlib.sha256()
    grid_widths: List[int] = []

    for r_idx, row in enumerate(table.rows):
        grid_widths.append(_row_grid_width(row))
        for c_idx, tc in enumerate(row._tr.findall(qn("w:tc"))):
            tcPr = tc.find(qn("w:tcPr"))
            gs = tcPr.find(qn("w:gridSpan")) if tcPr is not None else None
            vm = tcPr.find(qn("w:vMerge")) if tcPr is not None else None
            gs_val = gs.get(qn("w:val")) if gs is not None else "1"
            vm_val = (vm.get(qn("w:val")) or "continue") if vm is not None else "none"
            merge_hasher.update(f"r{r_idx}c{c_idx}:{gs_val}:{vm_val}|".encode("utf-8"))

    if table.rows:
        for cell in table.rows[0].cells:
            header_hasher.update((cell.text.strip() + "|").encode("utf-8"))

    return TableStructureSignature(
        table_index=table_index,
        row_count=len(table.rows),
        column_count=len(table.columns),
        grid_widths=grid_widths,
        merge_signature=merge_hasher.hexdigest()[:16],
        header_signature=header_hasher.hexdigest()[:16],
        semantic_fingerprint=FingerprintService.compute_table_semantic_fingerprint(table),
    )


def _orphan_vmerge_cells(doc: Document) -> List[str]:
    """Finds vMerge='continue' cells with no preceding 'restart' in the same column."""
    orphans: List[str] = []
    for t_idx, table in enumerate(doc.tables):
        open_columns: Set[int] = set()
        for r_idx, row in enumerate(table.rows):
            for c_idx, tc in enumerate(row._tr.findall(qn("w:tc"))):
                tcPr = tc.find(qn("w:tcPr"))
                vm = tcPr.find(qn("w:vMerge")) if tcPr is not None else None
                if vm is None:
                    open_columns.discard(c_idx)
                    continue
                val = vm.get(qn("w:val")) or "continue"
                if val == "restart":
                    open_columns.add(c_idx)
                elif c_idx not in open_columns:
                    orphans.append(f"T{t_idx}R{r_idx}C{c_idx}")
    return orphans


def _body_order_signature(doc: Document) -> str:
    """Hash of the ordered sequence of body-level element kinds (reading order)."""
    hasher = hashlib.sha256()
    for child in doc.element.body:
        tag = child.tag.split("}")[-1]
        if tag in ("p", "tbl", "sectPr"):
            hasher.update((tag + ";").encode("utf-8"))
    return hasher.hexdigest()[:16]


def _heading_signature(doc: Document) -> Tuple[str, int]:
    """Hash of the ordered (style, text) sequence of every heading paragraph."""
    hasher = hashlib.sha256()
    count = 0
    for p in doc.paragraphs:
        style_name = p.style.name if p.style else ""
        if "Heading" not in style_name and not style_name.lower().startswith("title"):
            continue
        text = p.text.strip()
        if not text:
            continue
        count += 1
        hasher.update(f"{style_name}:{text}|".encode("utf-8"))
    return hasher.hexdigest()[:16], count


def _header_footer_signature(blocks: Sequence[Dict[str, Any]]) -> str:
    """Hash of every header/footer text block perceived by Foundation."""
    hasher = hashlib.sha256()
    for b in blocks:
        if b.get("kind") in ("header", "footer"):
            hasher.update(f"{b.get('kind')}:{(b.get('text') or '').strip()}|".encode("utf-8"))
    return hasher.hexdigest()[:16]


def _inspect_package(doc_path: Path) -> Tuple[List[str], List[str], List[str]]:
    """Reads the OOXML package directly: parts, declared relationships, orphan refs.

    An "orphan ref" is an ``r:id``/``r:embed`` attribute in a part whose target
    relationship id is not declared in that part's ``_rels`` file.
    """
    parts: List[str] = []
    declared: List[str] = []
    referenced: Set[str] = set()
    declared_by_part: Dict[str, Set[str]] = {}

    with zipfile.ZipFile(str(doc_path)) as zf:
        parts = sorted(zf.namelist())

        for name in parts:
            if not name.endswith(".rels"):
                continue
            try:
                root = ET.fromstring(zf.read(name))
            except ET.ParseError:
                continue
            owner = _rels_owner(name)
            ids = set()
            for rel in root.findall(f"{{{_REL_NS}}}Relationship"):
                r_id = rel.get("Id") or ""
                ids.add(r_id)
                declared.append(f"{owner}::{r_id}::{rel.get('Type', '').rsplit('/', 1)[-1]}")
            declared_by_part[owner] = ids

        for name in parts:
            if not name.endswith(".xml") or name.endswith(".rels"):
                continue
            try:
                root = ET.fromstring(zf.read(name))
            except ET.ParseError:
                continue
            owner_ids = declared_by_part.get(name, set())
            for node in root.iter():
                for attr, value in node.attrib.items():
                    if not attr.startswith(f"{{{_R_NS}}}"):
                        continue
                    if not value or not value.startswith("rId"):
                        continue
                    if value not in owner_ids:
                        referenced.add(f"{name}::{value}")

    return parts, sorted(declared), sorted(referenced)


def _rels_owner(rels_name: str) -> str:
    """Maps ``word/_rels/document.xml.rels`` back to ``word/document.xml``."""
    directory, _, filename = rels_name.rpartition("/")
    base = directory[: -len("/_rels")] if directory.endswith("_rels") else directory
    base = base.rstrip("/")
    owner = filename[: -len(".rels")] if filename.endswith(".rels") else filename
    return f"{base}/{owner}" if base else owner
