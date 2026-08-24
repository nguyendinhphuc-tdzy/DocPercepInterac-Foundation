"""
Production Roll-Forward Planner (Phase H)
=========================================
Location: foundation/applications/rollforward/planner.py

Turns an authoritative workflow package into a governed plan:

    workflow slots
        -> template region model + position-free table identity
        -> historical correlation (structural/semantic, never ordinal)
        -> current-year source capability profiles
        -> verified source bindings
        -> structural deltas, row templates, validation rules
        -> RollForwardManifest + MutationPlan + RollForwardDiff[]

and stops there. The planner may propose; only a human may approve, and only
RollForwardOrchestrator may execute.

What this module refuses to do
------------------------------
    * read Ground Truth — every path is checked against the declared inputs AND
      against GroundTruthGuard before it is opened (`PlannerContaminationError`)
    * identify a table by its position — correlation runs through
      TableCorrespondenceResolver, whose scoring cannot see `ordinal`
    * hardcode a region id, a sheet name or a cell address — every address is
      discovered from the source workbook's own profiled header row
    * invent a value — every planned cell value is read out of a real source cell
    * ask a model anything — there is no provider import in this file
    * mark anything READY to make a fixture pass — a region is executable only
      when a verified binding and a structurally safe prototype row both exist

Where a value comes from
------------------------
For each template table, the planner derives a target schema from the table's own
header, asks the source workbooks which of them carry the required dataset roles,
then maps template columns onto source columns **by header text**. A planned cell
value is read from `sheet[column_letter][header_row + 1 + record_index]` — an
address computed from the profiled sheet, never a literal.
"""
from __future__ import annotations

import hashlib
import time
import uuid
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from applications.rollforward.models import (
    ExecutionGate,
    FigureBinding,
    HistoricalReference,
    ManifestStatus,
    RegionClassification,
    RollForwardDiff,
    RollForwardManifest,
    RollForwardRegion,
    RowTemplate,
    SourceBinding,
    SourceBindingStatus,
    SourceType,
    StructuralDelta,
    ValidationRule,
    ValidationRuleType,
    ValidationSeverity,
    DiffChangeType,
)
from applications.rollforward.semantic_binding import (
    BindingVerdict,
    SemanticBindingValidator,
    TargetRegionSchema,
    TargetSchemaDeriver,
)
from applications.rollforward.source_capability import (
    SheetDatasetProfile,
    SourceCapabilityProfiler,
    WorkbookCapabilityProfile,
)
from applications.rollforward.source_intake import GroundTruthGuard
from applications.rollforward.state_machine import RollForwardStateMachine
from applications.rollforward.structural_writeback import (
    CellMutationSpec,
    FingerprintService,
    MutationPlan,
    RowMutationSpec,
    TableMutationSpec,
)
from applications.rollforward.table_identity import (
    CorrespondenceConfidence,
    TableCorrespondence,
    TableCorrespondenceResolver,
    TableIdentityProfiler,
    TableIdentitySignature,
)
from applications.rollforward.workflow_intake import FiscalPeriod, RollForwardPeriods


class PlannerError(RuntimeError):
    """Raised when planning cannot proceed on the inputs it was given."""


class PlannerContaminationError(PlannerError):
    """Raised when the planner is asked to read something that is not an input.

    Ground Truth is evaluation-only. A planner that reads it is not planning, it
    is copying an answer — so reaching for a path outside the declared inputs is
    a hard failure, never a warning.
    """


# ============================================================================
# 1. INPUTS
# ============================================================================

@dataclass(frozen=True)
class DocumentRef:
    """One workflow document: its id, its file, and nothing inferred."""
    document_id: str
    path: Path
    filename: str

    @property
    def suffix(self) -> str:
        return self.path.suffix.lower()


@dataclass(frozen=True)
class PlannerInputs:
    """Exactly the authoritative workflow slots. Nothing else is in scope."""
    workflow_id: str
    session_id: str
    historical: DocumentRef
    template: DocumentRef
    current_sources: Tuple[DocumentRef, ...]
    historical_period: Optional[FiscalPeriod] = None
    current_period: Optional[FiscalPeriod] = None

    def all_documents(self) -> Tuple[DocumentRef, ...]:
        return (self.historical, self.template) + tuple(self.current_sources)


class InputAccessGuard:
    """Every file the planner opens must be a declared workflow input.

    Two independent checks, because one of them can be defeated by a bug and the
    other by a rename:

        1. the resolved path must be one of the declared inputs
        2. GroundTruthGuard must not recognise its content

    Every allowed access is recorded, so the planning report can state exactly
    which files were read.
    """

    def __init__(self, inputs: PlannerInputs):
        self._allowed = {doc.path.resolve(): doc for doc in inputs.all_documents()}
        self.accessed: List[str] = []

    def check(self, path: Path, purpose: str) -> Path:
        resolved = Path(path).resolve()
        if resolved not in self._allowed:
            raise PlannerContaminationError(
                f"The planner tried to read '{resolved.name}' for {purpose}, which is not "
                f"one of the workflow's input documents. Planning reads the historical "
                f"Local File, the master template and the current-year sources — nothing else.")
        # Content-level refusal: a Ground Truth file renamed into a slot is still
        # Ground Truth, and this catches it by hash.
        GroundTruthGuard.check(resolved)
        record = f"{resolved.name} ({purpose})"
        if record not in self.accessed:
            self.accessed.append(record)
        return resolved


# ============================================================================
# 2. RESULT
# ============================================================================

@dataclass
class PlannedRegionOutcome:
    """Why one region ended up READY, HUMAN_REVIEW or BLOCKED."""
    region_id: str
    section_name: str
    disposition: str                     # READY | HUMAN_REVIEW | BLOCKED
    identity_key: str
    correspondence: str
    binding_verdict: str
    reasons: List[str] = field(default_factory=list)
    source_labels: List[str] = field(default_factory=list)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "region_id": self.region_id,
            "section_name": self.section_name,
            "disposition": self.disposition,
            "identity_key": self.identity_key,
            "correspondence": self.correspondence,
            "binding_verdict": self.binding_verdict,
            "reasons": list(self.reasons),
            "source_labels": list(self.source_labels),
        }


@dataclass
class PlanningResult:
    """Everything one planning run produced, and how it got there."""
    manifest: RollForwardManifest
    mutation_plan: MutationPlan
    diffs: List[RollForwardDiff]
    outcomes: List[PlannedRegionOutcome]
    rejected_evidence: List[Dict[str, Any]]
    accessed_documents: List[str]
    timings_ms: Dict[str, float]
    inputs: PlannerInputs

    @property
    def ready_regions(self) -> List[RollForwardRegion]:
        return [r for r in self.manifest.regions
                if r.execution_gate == ExecutionGate.READY]

    @property
    def has_executable_work(self) -> bool:
        return bool(self.mutation_plan.table_mutations)

    def readiness_summary(self) -> Dict[str, Any]:
        """Region counts by disposition, plus what each blocked region needs."""
        counts = {"READY": 0, "HUMAN_REVIEW": 0, "BLOCKED": 0}
        for outcome in self.outcomes:
            counts[outcome.disposition] = counts.get(outcome.disposition, 0) + 1
        return {
            "regions_total": len(self.outcomes),
            "by_disposition": counts,
            "executable": counts["READY"] > 0 and self.has_executable_work,
            "statement": (
                "A region is READY only when a verified current-year binding and a "
                "structurally safe prototype row both exist."),
        }

    def unresolved_blockers(self) -> List[Dict[str, Any]]:
        """Every region that cannot proceed, and the reason it cannot."""
        return [
            {
                "region_id": outcome.region_id,
                "section_name": outcome.section_name,
                "disposition": outcome.disposition,
                "binding_verdict": outcome.binding_verdict,
                "correspondence": outcome.correspondence,
                "reasons": list(outcome.reasons),
            }
            for outcome in self.outcomes if outcome.disposition != "READY"
        ]

    def report(self) -> Dict[str, Any]:
        """Machine-readable planning report. Carries no document plaintext."""
        by_disposition: Dict[str, int] = {}
        for outcome in self.outcomes:
            by_disposition[outcome.disposition] = by_disposition.get(outcome.disposition, 0) + 1

        return {
            "schema_version": "1.0.0",
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "workflow_id": self.inputs.workflow_id,
            "session_id": self.inputs.session_id,
            "inputs": {
                "historical_document_id": self.inputs.historical.document_id,
                "historical_document": self.inputs.historical.filename,
                "template_document_id": self.inputs.template.document_id,
                "template_document": self.inputs.template.filename,
                "current_source_document_ids": [s.document_id for s in self.inputs.current_sources],
                "current_source_documents": [s.filename for s in self.inputs.current_sources],
                "historical_period": (self.inputs.historical_period.label
                                      if self.inputs.historical_period else None),
                "current_period": (self.inputs.current_period.label
                                   if self.inputs.current_period else None),
            },
            "manifest": {
                "manifest_id": self.manifest.manifest_id,
                "manifest_version": self.manifest.manifest_version,
                "status": self.manifest.status.value,
                "region_count": len(self.manifest.regions),
                "ready_regions": len(self.ready_regions),
            },
            "mutation_plan": {
                "plan_id": self.mutation_plan.plan_id,
                "tables": len(self.mutation_plan.table_mutations),
                "rows": sum(len(t.row_mutations) for t in self.mutation_plan.table_mutations),
                "cells": sum(len(r.cells) for t in self.mutation_plan.table_mutations
                             for r in t.row_mutations),
                "digest": plan_digest(self.mutation_plan),
            },
            "dispositions": by_disposition,
            "readiness_summary": self.readiness_summary(),
            "unresolved_blockers": self.unresolved_blockers(),
            "regions": [o.to_dict() for o in self.outcomes],
            "rejected_evidence": list(self.rejected_evidence),
            "documents_read": list(self.accessed_documents),
            "ground_truth_accessed": False,
            "timings_ms": dict(self.timings_ms),
            "governance": {
                "approved": bool(self.manifest.approved_by),
                "executable": self.has_executable_work,
                "statement": (
                    "Planning only. The manifest is not approved, nothing was mutated, and "
                    "execution requires the human approval endpoint."),
            },
        }


def plan_digest(plan: MutationPlan) -> str:
    """Stable digest of a plan's addressable content."""
    hasher = hashlib.sha256()
    for table in plan.table_mutations:
        hasher.update(f"{table.target_region_id}|{table.table_hash}|{table.operation}".encode())
        for row in table.row_mutations:
            for cell in row.cells:
                hasher.update(
                    f"{row.row_idx}:{cell.col_idx}:{cell.source_doc_name}:"
                    f"{cell.source_sheet}:{cell.source_cell_address}".encode())
    return hasher.hexdigest()[:16]


# ============================================================================
# 3. THE PLANNER
# ============================================================================

class RollForwardPlanner:
    """Deterministic production planning over authoritative workflow inputs."""

    # A prototype row may be cloned only when the table's own structure says it
    # is safe: a header plus at least one data row, and no merge topology in the
    # body that cloning would corrupt.
    MIN_ROWS_FOR_ROW_TEMPLATE = 2

    @classmethod
    def plan(cls, inputs: PlannerInputs) -> PlanningResult:
        timings: Dict[str, float] = {}
        guard = InputAccessGuard(inputs)

        started = time.perf_counter()
        cls._validate_periods(inputs)
        timings["period_validation_ms"] = _elapsed(started)

        started = time.perf_counter()
        template_path = guard.check(inputs.template.path, "template structure")
        template_sigs = TableIdentityProfiler.profile_document(
            template_path, document_id=inputs.template.document_id)
        timings["template_profile_ms"] = _elapsed(started)

        started = time.perf_counter()
        historical_path = guard.check(inputs.historical.path, "historical correlation")
        historical_sigs = TableIdentityProfiler.profile_document(
            historical_path, document_id=inputs.historical.document_id)
        # Position-free: the resolver's scoring function cannot read `ordinal`.
        correspondences = TableCorrespondenceResolver.correspond(template_sigs, historical_sigs)
        timings["historical_correlation_ms"] = _elapsed(started)

        started = time.perf_counter()
        workbooks, source_paths, rejected = cls._profile_sources(inputs, guard)
        timings["source_profile_ms"] = _elapsed(started)

        started = time.perf_counter()
        regions: List[RollForwardRegion] = []
        outcomes: List[PlannedRegionOutcome] = []
        table_specs: List[TableMutationSpec] = []
        diffs: List[RollForwardDiff] = []

        correspondence_by_key = {c.left.identity_key: c for c in correspondences}

        for sig in template_sigs:
            region, outcome, spec, diff = cls._plan_region(
                sig=sig,
                correspondence=correspondence_by_key.get(sig.identity_key),
                workbooks=workbooks,
                source_paths=source_paths,
                inputs=inputs,
                template_path=template_path,
            )
            regions.append(region)
            outcomes.append(outcome)
            if spec is not None:
                table_specs.append(spec)
            if diff is not None:
                diffs.append(diff)
        timings["region_planning_ms"] = _elapsed(started)

        started = time.perf_counter()
        manifest = cls._assemble_manifest(inputs, regions)
        mutation_plan = MutationPlan(
            manifest_id=manifest.manifest_id,
            manifest_version=manifest.manifest_version,
            target_doc_name=inputs.template.filename,
            table_mutations=table_specs,
        )
        timings["plan_assembly_ms"] = _elapsed(started)

        return PlanningResult(
            manifest=manifest,
            mutation_plan=mutation_plan,
            diffs=diffs,
            outcomes=outcomes,
            rejected_evidence=rejected,
            accessed_documents=list(guard.accessed),
            timings_ms=timings,
            inputs=inputs,
        )

    # ------------------------------------------------------------------
    # 3.1 PERIODS
    # ------------------------------------------------------------------

    @staticmethod
    def _validate_periods(inputs: PlannerInputs) -> None:
        periods = RollForwardPeriods(historical=inputs.historical_period,
                                     current=inputs.current_period)
        if periods.relationship.value == "UNKNOWN":
            raise PlannerError(
                "Planning needs both fiscal periods: they are read from the documents' own "
                "content and neither may be assumed.")
        if not periods.satisfies_invariant:
            raise PlannerError(f"Cannot plan a roll-forward: {periods.describe()}")

    # ------------------------------------------------------------------
    # 3.2 SOURCES
    # ------------------------------------------------------------------

    @classmethod
    def _profile_sources(cls, inputs: PlannerInputs, guard: InputAccessGuard
                         ) -> Tuple[List[WorkbookCapabilityProfile], Dict[str, Path],
                                    List[Dict[str, Any]]]:
        """Profile every current-year source; record what could not contribute."""
        workbooks: List[WorkbookCapabilityProfile] = []
        source_paths: Dict[str, Path] = {}
        rejected: List[Dict[str, Any]] = []

        for source in inputs.current_sources:
            path = guard.check(source.path, "current-year source evidence")
            if source.suffix not in (".xlsx", ".xlsm"):
                # A narrative source carries no addressable cells, so it cannot
                # back a cell-level binding. Recorded, not silently dropped.
                rejected.append({
                    "document": source.filename,
                    "document_id": source.document_id,
                    "reason": "NOT_CELL_ADDRESSABLE",
                    "detail": ("Only workbook sources expose addressable cells; a document "
                               "source cannot back a verified cell binding."),
                })
                continue
            profile = SourceCapabilityProfiler.profile_workbook(path)
            workbooks.append(profile)
            source_paths[profile.document_name] = path

        return workbooks, source_paths, rejected

    # ------------------------------------------------------------------
    # 3.3 ONE REGION
    # ------------------------------------------------------------------

    @classmethod
    def _plan_region(
        cls,
        sig: TableIdentitySignature,
        correspondence: Optional[TableCorrespondence],
        workbooks: Sequence[WorkbookCapabilityProfile],
        source_paths: Dict[str, Path],
        inputs: PlannerInputs,
        template_path: Path,
    ) -> Tuple[RollForwardRegion, PlannedRegionOutcome,
               Optional[TableMutationSpec], Optional[RollForwardDiff]]:
        # Region ids are derived from the table's own position-free identity, so
        # the same table keeps its id when the document is renumbered.
        region_id = f"rfr-{sig.identity_key}"
        section_name = sig.nearest_heading or sig.caption_before or "Untitled region"

        schema = TargetSchemaDeriver.derive(region_id, sig)
        validation = SemanticBindingValidator.validate(schema, list(workbooks))

        historical_reference = cls._historical_reference(correspondence, inputs)
        bindings, binding_notes = cls._build_bindings(sig, schema, workbooks, inputs, validation)

        classification, review_state, reasons = cls._classify(
            sig, correspondence, validation, bindings)

        region = RollForwardRegion(
            region_id=region_id,
            section_name=section_name,
            target_document_id=inputs.template.document_id,
            classification=classification,
            historical_reference=historical_reference,
            current_sources=bindings,
            review_state=review_state,
            notes="; ".join(reasons) or None,
        )

        spec: Optional[TableMutationSpec] = None
        diff: Optional[RollForwardDiff] = None

        if region.execution_gate == ExecutionGate.READY:
            delta, row_template, spec, diff = cls._plan_mutation(
                sig, schema, bindings, workbooks, source_paths, template_path, region_id)
            region.structural_delta = delta
            region.row_template = row_template
            region.validation_rules = cls._validation_rules(sig, delta)
            if spec is None:
                # A region can be bound and still have nothing safe to write.
                # It stays in the manifest as reviewable, never as executable.
                region.classification = RegionClassification.MANUAL_REVIEW
                region.review_state = "MANUAL_REQUIRED"
                reasons.append("No structurally safe row growth was derivable from the evidence.")
                region.notes = "; ".join(reasons)

        disposition = cls._disposition(region, bindings, spec)

        outcome = PlannedRegionOutcome(
            region_id=region_id,
            section_name=section_name,
            disposition=disposition,
            identity_key=sig.identity_key,
            correspondence=(correspondence.confidence.value if correspondence
                            else CorrespondenceConfidence.UNCORRELATED.value),
            binding_verdict=validation.verdict.value,
            reasons=reasons + binding_notes,
            source_labels=[f"{b.source_doc_name}!{b.sheet_name}" for b in bindings if b.sheet_name],
        )
        return region, outcome, spec, diff

    @staticmethod
    def _disposition(region: RollForwardRegion, bindings: List[SourceBinding],
                     spec: Optional[TableMutationSpec]) -> str:
        """READY / HUMAN_REVIEW / BLOCKED, from evidence rather than optimism.

        BLOCKED means "no current-year artifact supports this at all" — the user
        must supply something before the region can move. HUMAN_REVIEW means the
        evidence exists but a person has to resolve it (ambiguous source, weak
        historical correspondence, unverified binding).
        """
        if region.execution_gate == ExecutionGate.READY and spec is not None:
            return "READY"
        statuses = {b.status for b in bindings}
        if statuses and statuses <= {SourceBindingStatus.MISSING}:
            return "BLOCKED"
        return "HUMAN_REVIEW"

    # ------------------------------------------------------------------
    # 3.4 BINDINGS — discovered addresses only
    # ------------------------------------------------------------------

    @classmethod
    def _build_bindings(
        cls,
        sig: TableIdentitySignature,
        schema: TargetRegionSchema,
        workbooks: Sequence[WorkbookCapabilityProfile],
        inputs: PlannerInputs,
        validation,
    ) -> Tuple[List[SourceBinding], List[str]]:
        notes: List[str] = []
        candidates = cls._candidate_sheets(schema, workbooks)

        if not candidates:
            return ([SourceBinding(
                source_doc_id="",
                source_doc_name="(none)",
                source_type=SourceType.XLSX,
                status=SourceBindingStatus.MISSING,
                match_basis=["required_role_absent"],
                reason=("No current-year source carries the dataset role this region "
                        f"requires ({[r.value for r in schema.required_source_roles]})."),
            )], notes)

        if len(candidates) > 1:
            labels = ", ".join(f"{wb.document_name}!{sheet.sheet_name}"
                               for wb, sheet in candidates)
            notes.append(f"Multiple candidate sources ({labels}); a human must choose.")
            return ([SourceBinding(
                source_doc_id=cls._document_id_for(inputs, wb.document_name),
                source_doc_name=wb.document_name,
                source_type=SourceType.XLSX,
                sheet_name=sheet.sheet_name,
                status=SourceBindingStatus.AMBIGUOUS,
                match_basis=["role_match", "multiple_candidates"],
                reason=("More than one current-year sheet can supply this region; the "
                        "planner will not choose between them."),
            ) for wb, sheet in candidates], notes)

        workbook, sheet = candidates[0]
        column_map = cls._map_columns(sig, sheet)
        if not column_map:
            notes.append(
                f"{workbook.document_name}!{sheet.sheet_name} carries the role but none of "
                f"its columns match this table's headers.")
            return ([SourceBinding(
                source_doc_id=cls._document_id_for(inputs, workbook.document_name),
                source_doc_name=workbook.document_name,
                source_type=SourceType.XLSX,
                sheet_name=sheet.sheet_name,
                status=SourceBindingStatus.MISSING,
                match_basis=["role_match", "no_column_correspondence"],
                reason="No column in the source sheet corresponds to this table's columns.",
            )], notes)

        verified = validation.verdict == BindingVerdict.VERIFIED
        return ([SourceBinding(
            source_doc_id=cls._document_id_for(inputs, workbook.document_name),
            source_doc_name=workbook.document_name,
            source_type=SourceType.XLSX,
            sheet_name=sheet.sheet_name,
            cell_range=cls._data_range(sheet, column_map),
            status=(SourceBindingStatus.VERIFIED if verified
                    else SourceBindingStatus.UNVERIFIED),
            match_basis=["role_match", "header_correspondence"],
            reason=(f"{sheet.record_count} records under header row "
                    f"{sheet.header_row_index}; columns matched by header text."),
            provenance={
                "sheet": sheet.sheet_name,
                "header_row_index": sheet.header_row_index,
                "matched_columns": {str(k): v for k, v in column_map.items()},
                "record_count": sheet.record_count,
            },
        )], notes)

    @staticmethod
    def _candidate_sheets(schema: TargetRegionSchema,
                          workbooks: Sequence[WorkbookCapabilityProfile]
                          ) -> List[Tuple[WorkbookCapabilityProfile, SheetDatasetProfile]]:
        found = []
        for workbook in workbooks:
            for role in schema.required_source_roles:
                if not workbook.has_role(role):
                    continue
                for sheet in workbook.sheets_with_role(role):
                    if (workbook, sheet) not in found:
                        found.append((workbook, sheet))
        return found

    @staticmethod
    def _map_columns(sig: TableIdentitySignature,
                     sheet: SheetDatasetProfile) -> Dict[int, str]:
        """Template column index -> source column letter, matched by header text.

        Matching is on normalised header text, so a renamed file or a re-ordered
        sheet changes nothing. A column with no counterpart is simply not mapped,
        and an unmapped column is never written.
        """
        source_headers = {
            _norm_header(col.header): col.letter
            for col in sheet.record_schema if col.header
        }
        mapping: Dict[int, str] = {}
        for index, header in enumerate(sig.header_cells):
            key = _norm_header(header)
            if not key:
                continue
            if key in source_headers:
                mapping[index] = source_headers[key]
                continue
            # Containment is the weakest accepted signal, and only when unambiguous.
            hits = [letter for text, letter in source_headers.items()
                    if key and (key in text or text in key)]
            if len(hits) == 1:
                mapping[index] = hits[0]
        return mapping

    @staticmethod
    def _data_range(sheet: SheetDatasetProfile, column_map: Dict[int, str]) -> Optional[str]:
        if not column_map or sheet.header_row_index is None:
            return None
        letters = sorted(column_map.values())
        first_row = sheet.header_row_index + 1
        last_row = first_row + max(sheet.record_count - 1, 0)
        return f"{letters[0]}{first_row}:{letters[-1]}{last_row}"

    @staticmethod
    def _document_id_for(inputs: PlannerInputs, filename: str) -> str:
        for source in inputs.current_sources:
            if source.filename == filename:
                return source.document_id
        return ""

    # ------------------------------------------------------------------
    # 3.5 CLASSIFICATION
    # ------------------------------------------------------------------

    @staticmethod
    def _classify(sig, correspondence, validation, bindings
                  ) -> Tuple[RegionClassification, str, List[str]]:
        reasons: List[str] = []

        if correspondence is None or correspondence.confidence == CorrespondenceConfidence.UNCORRELATED:
            reasons.append(
                "No corresponding table in the historical Local File; the region has no "
                "prior-year basis to roll forward from.")
            return RegionClassification.MANUAL_REVIEW, "MANUAL_REQUIRED", reasons

        if correspondence.confidence == CorrespondenceConfidence.WEAK:
            reasons.append(
                f"Historical correspondence is weak (score {correspondence.score:.2f}, "
                f"matched on {', '.join(correspondence.matched_factors) or 'partial signals'}).")

        statuses = {b.status for b in bindings}
        if SourceBindingStatus.AMBIGUOUS in statuses:
            reasons.append("Source binding is ambiguous.")
            return RegionClassification.MANUAL_REVIEW, "MANUAL_REQUIRED", reasons
        if statuses == {SourceBindingStatus.MISSING}:
            reasons.append("No current-year evidence supports this region.")
            return RegionClassification.STATIC, "PENDING", reasons
        if SourceBindingStatus.UNVERIFIED in statuses:
            reasons.append(f"Binding not verified: {validation.verdict.value}.")
            return RegionClassification.MANUAL_REVIEW, "MANUAL_REQUIRED", reasons

        if correspondence.confidence == CorrespondenceConfidence.WEAK:
            return RegionClassification.MANUAL_REVIEW, "MANUAL_REQUIRED", reasons

        reasons.append(
            f"Verified current-year binding and {correspondence.confidence.value.lower()} "
            f"historical correspondence.")
        return RegionClassification.UPDATE, "PENDING", reasons

    # ------------------------------------------------------------------
    # 3.6 STRUCTURAL DELTA, ROW TEMPLATE, MUTATIONS
    # ------------------------------------------------------------------

    @classmethod
    def _plan_mutation(
        cls,
        sig: TableIdentitySignature,
        schema: TargetRegionSchema,
        bindings: List[SourceBinding],
        workbooks: Sequence[WorkbookCapabilityProfile],
        source_paths: Dict[str, Path],
        template_path: Path,
        region_id: str,
    ) -> Tuple[Optional[StructuralDelta], Optional[RowTemplate],
               Optional[TableMutationSpec], Optional[RollForwardDiff]]:
        binding = bindings[0]
        sheet = cls._sheet_for(binding, workbooks)
        if sheet is None or sheet.header_row_index is None:
            return None, None, None, None

        column_map = {int(k): v for k, v in
                      (binding.provenance.get("matched_columns") or {}).items()}
        if not column_map:
            return None, None, None, None

        template_rows = sig.row_count
        if template_rows < cls.MIN_ROWS_FOR_ROW_TEMPLATE:
            return None, None, None, None
        if not cls._prototype_row_is_clonable(template_path, sig):
            # Cloning a prototype row through merged topology is not structurally
            # safe, and the planner does not attempt it.
            return None, None, None, None

        header_rows = 1
        template_data_rows = template_rows - header_rows
        target_rows = header_rows + sheet.record_count
        insert_count = max(target_rows - template_rows, 0)
        if insert_count == 0:
            return None, None, None, None

        delta = StructuralDelta(
            template_rows=template_rows,
            target_rows=target_rows,
            insert_count=insert_count,
            delete_count=0,
            column_delta=0,
            merge_topology_changed=False,
            row_template_anchor=f"{region_id}:row:{template_rows - 1}",
            observation_source="template_structure+current_year_evidence",
            observation_context={
                "template_data_rows": template_data_rows,
                "source_record_count": sheet.record_count,
                "source_sheet": sheet.sheet_name,
            },
        )

        row_template = RowTemplate(
            template_row_idx=template_rows - 1,
            row_anchor=delta.row_template_anchor,
            column_schemas=[c.to_dict() for c in sig.column_schemas],
            safe_to_clone=True,
        )

        rows, sourced_cells = cls._read_row_values(
            sheet=sheet, column_map=column_map, binding=binding,
            source_paths=source_paths, start_row=template_rows, insert_count=insert_count)
        if not rows:
            return delta, row_template, None, None

        spec = TableMutationSpec(
            target_region_id=region_id,
            table_index=sig.ordinal,          # intra-document locator for the writer only
            table_hash=sig.table_hash or sig.identity_key,
            operation="INSERT_ROWS",
            source_row_template_idx=template_rows - 1,
            initial_row_count=template_rows,
            target_row_count=target_rows,
            insert_count=insert_count,
            expected_precondition_hash=cls._precondition_hash(template_path, sig),
            expected_postcondition_hash="",
            row_mutations=rows,
        )

        diff = RollForwardDiff(
            region_id=region_id,
            change_type=DiffChangeType.ROW_ADDED,
            before_summary={"rows": template_rows, "columns": sig.column_count},
            after_summary={"rows": target_rows, "columns": sig.column_count},
            delta_details=[{
                "rows_inserted": insert_count,
                "cells_written": sourced_cells,
                "source": f"{binding.source_doc_name}!{sheet.sheet_name}",
            }],
        )
        return delta, row_template, spec, diff

    @classmethod
    def _read_row_values(cls, sheet: SheetDatasetProfile, column_map: Dict[int, str],
                         binding: SourceBinding, source_paths: Dict[str, Path],
                         start_row: int, insert_count: int
                         ) -> Tuple[List[RowMutationSpec], int]:
        """Read the planned values out of the source workbook's real cells."""
        import openpyxl

        workbook_path = source_paths.get(binding.source_doc_name)
        if workbook_path is None:
            return [], 0

        rows: List[RowMutationSpec] = []
        cells_written = 0
        workbook = openpyxl.load_workbook(str(workbook_path), data_only=True, read_only=True)
        try:
            worksheet = workbook[sheet.sheet_name]
            first_data_row = (sheet.header_row_index or 0) + 1
            for offset in range(insert_count):
                excel_row = first_data_row + offset
                cells: List[CellMutationSpec] = []
                for col_idx, letter in sorted(column_map.items()):
                    address = f"{letter}{excel_row}"
                    value = worksheet[address].value
                    if value is None:
                        continue
                    cells.append(CellMutationSpec(
                        col_idx=col_idx,
                        source_doc_name=binding.source_doc_name,
                        source_sheet=sheet.sheet_name,
                        source_cell_address=address,
                        value=str(value),
                    ))
                if cells:
                    rows.append(RowMutationSpec(row_idx=start_row + offset, cells=cells))
                    cells_written += len(cells)
        finally:
            workbook.close()
        return rows, cells_written

    @staticmethod
    def _sheet_for(binding: SourceBinding,
                   workbooks: Sequence[WorkbookCapabilityProfile]) -> Optional[SheetDatasetProfile]:
        for workbook in workbooks:
            if workbook.document_name != binding.source_doc_name:
                continue
            for sheet in workbook.sheets:
                if sheet.sheet_name == binding.sheet_name:
                    return sheet
        return None

    @staticmethod
    def _prototype_row_is_clonable(template_path: Path, sig: TableIdentitySignature) -> bool:
        """Is the last row of this table safe to use as a clone prototype?

        Safe means the row is a plain grid row: one distinct cell per column, and
        no vertical-merge continuation. Cloning a row that participates in a merge
        would duplicate half of a merged span and corrupt the table, so a region
        whose prototype fails this check is left for a human rather than planned.
        """
        from docx import Document

        document = Document(str(template_path))
        if sig.ordinal >= len(document.tables):
            return False
        table = document.tables[sig.ordinal]
        if len(table.rows) < RollForwardPlanner.MIN_ROWS_FOR_ROW_TEMPLATE:
            return False

        prototype = table.rows[-1]
        distinct_cells = {id(cell._tc) for cell in prototype.cells}
        if len(distinct_cells) != len(table.columns):
            return False                      # horizontal merge / gridSpan
        for cell in prototype.cells:
            if cell._tc.tcPr is not None and cell._tc.tcPr.find(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}vMerge"
            ) is not None:
                return False                  # vertical merge continuation
        return True

    @staticmethod
    def _precondition_hash(template_path: Path, sig: TableIdentitySignature) -> str:
        from docx import Document

        document = Document(str(template_path))
        if sig.ordinal >= len(document.tables):
            return ""
        return FingerprintService.compute_table_semantic_fingerprint(document.tables[sig.ordinal])

    # ------------------------------------------------------------------
    # 3.7 VALIDATION RULES + MANIFEST
    # ------------------------------------------------------------------

    @staticmethod
    def _validation_rules(sig: TableIdentitySignature,
                          delta: Optional[StructuralDelta]) -> List[ValidationRule]:
        rules = [
            ValidationRule(
                rule_type=ValidationRuleType.COLUMN_COUNT_UNCHANGED,
                severity=ValidationSeverity.BLOCKER,
                parameters={"expected_columns": sig.column_count},
                description="A roll-forward adds rows; it never changes the column count.",
            ),
            ValidationRule(
                rule_type=ValidationRuleType.MERGE_TOPOLOGY_PRESERVED,
                severity=ValidationSeverity.BLOCKER,
                parameters={"merge_signature": sig.merge_signature},
                description="Merged-cell topology must survive row cloning unchanged.",
            ),
            ValidationRule(
                rule_type=ValidationRuleType.SOURCE_VALUE_PRESENT,
                severity=ValidationSeverity.BLOCKER,
                parameters={},
                description="Every written cell must resolve to a real source cell.",
            ),
        ]
        if delta is not None:
            rules.insert(0, ValidationRule(
                rule_type=ValidationRuleType.ROW_COUNT_MATCH,
                severity=ValidationSeverity.BLOCKER,
                parameters={"expected_rows": delta.target_rows},
                description="The written table must end with exactly the planned row count.",
            ))
        return rules

    @classmethod
    def _assemble_manifest(cls, inputs: PlannerInputs,
                           regions: List[RollForwardRegion]) -> RollForwardManifest:
        manifest = RollForwardManifest(
            schema_version="1.0.0",
            manifest_version=1,
            session_id=inputs.session_id,
            historical_document_id=inputs.historical.document_id,
            template_document_id=inputs.template.document_id,
            current_source_document_ids=[s.document_id for s in inputs.current_sources],
            status=ManifestStatus.DISCOVERED,
            regions=regions,
            figures=cls._figure_bindings(),
        )
        RollForwardStateMachine.transition(
            manifest, ManifestStatus.PLANNED, actor="system",
            reason="Deterministic production planning complete.")
        if manifest.has_unresolved_reviews():
            RollForwardStateMachine.transition(
                manifest, ManifestStatus.REVIEW_REQUIRED, actor="system",
                reason="Regions require human review before approval.")
        # The planner stops here. APPROVED, EXECUTING and COMPLETED are reachable
        # only through the human approval endpoint and the orchestrator.
        return manifest

    @staticmethod
    def _figure_bindings() -> List[FigureBinding]:
        """No figure is bound without real evidence for it.

        A figure binding needs a current-year source for the image or chart it
        replaces. No such source is addressable from a workbook profile, so the
        planner produces none rather than binding a figure to nothing.
        """
        return []

    @staticmethod
    def _historical_reference(correspondence: Optional[TableCorrespondence],
                              inputs: PlannerInputs) -> Optional[HistoricalReference]:
        if correspondence is None or correspondence.right is None:
            return None
        return HistoricalReference(
            doc_id=inputs.historical.document_id,
            doc_name=inputs.historical.filename,
            table_index=correspondence.right.ordinal,
            value_snippet=None,
        )


def _elapsed(started: float) -> float:
    return round((time.perf_counter() - started) * 1000, 2)


def _norm_header(text: str) -> str:
    return " ".join(str(text or "").lower().split())


__all__ = [
    "DocumentRef",
    "PlannerInputs",
    "PlannerError",
    "PlannerContaminationError",
    "InputAccessGuard",
    "PlannedRegionOutcome",
    "PlanningResult",
    "RollForwardPlanner",
    "plan_digest",
]
